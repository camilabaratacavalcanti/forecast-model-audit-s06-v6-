"""
Leitura física de um workbook "descritivo das variáveis".

Responsabilidade única: transformar a aba do workbook em linhas
brutas, sem interpretar o conteúdo.

    - o cabeçalho é LOCALIZADO (a primeira linha que contém todas as
      colunas obrigatórias), não assumido numa linha fixa — o yield
      tem cabeçalho na linha 1, os demais na linha 2;
    - a janela de dados é derivada do CONTEÚDO: toda linha abaixo do
      cabeçalho é lida, e linhas sem nenhum valor são descartadas
      como vazias (não viram entidade) — nunca uma faixa fixa de
      linhas (ex.: LAST_DATA_ROW=154, que no MaxHT v9 lia 32 linhas
      vazias);
    - conteúdo fora das colunas do cabeçalho é erro explícito
      (conteúdo residual), assim como uma linha com conteúdo e sem
      `name`;
    - nenhum valor é alterado: sem strip, sem conversão de tipo.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from pathlib import Path


REQUIRED_COLUMNS = (
    "Type",
    "name",
    "description",
    "unit",
    "value",
    "version",
    "variable_type",
    "frequency",
    "scope_type",
    "scope_value",
    "source_reference",
    "status",
    "value_type",
    "allowed_values",
    "declared_result_states",
    "expression",
    "fonte",
)

# Colunas livres conhecidas: lidas e preservadas na linha bruta, sem
# significado contratual.
OPTIONAL_COLUMNS = ("OBS",)

# Quantas linhas do topo são examinadas à procura do cabeçalho.
HEADER_SEARCH_ROWS = 10


class WorkbookFormatError(ValueError):
    """O workbook não respeita o formato físico esperado."""


@dataclass(frozen=True)
class WorkbookRow:
    row: int
    values: dict

    def get(self, column: str):
        return self.values.get(column)


@dataclass(frozen=True)
class WorkbookData:
    path: Path
    file_name: str
    sha256: str
    sheet: str
    header_row: int
    columns: tuple[str, ...]
    rows: tuple[WorkbookRow, ...]
    empty_rows: tuple[int, ...] = field(default=())
    last_sheet_row: int = 0


def sha256_of(path: str | Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _is_empty(value) -> bool:
    return value is None or (isinstance(value, str) and value == "")


def read_workbook(path: str | Path, sheet: str) -> WorkbookData:
    import openpyxl

    path = Path(path)
    workbook = openpyxl.load_workbook(path, data_only=True)

    if sheet not in workbook.sheetnames:
        raise WorkbookFormatError(
            f"{path.name}: aba '{sheet}' não encontrada "
            f"(abas: {workbook.sheetnames})."
        )

    ws = workbook[sheet]

    header_row = None
    header: dict[int, str] = {}

    for row_number in range(1, min(HEADER_SEARCH_ROWS, ws.max_row) + 1):
        cells = {
            column: ws.cell(row=row_number, column=column).value
            for column in range(1, ws.max_column + 1)
        }
        names = {v for v in cells.values() if isinstance(v, str)}

        if set(REQUIRED_COLUMNS) <= names:
            header_row = row_number
            header = {
                column: value
                for column, value in cells.items()
                if not _is_empty(value)
            }
            break

    if header_row is None:
        found = [
            ws.cell(row=r, column=c).value
            for r in range(1, min(HEADER_SEARCH_ROWS, ws.max_row) + 1)
            for c in range(1, ws.max_column + 1)
            if isinstance(ws.cell(row=r, column=c).value, str)
        ]
        missing = sorted(set(REQUIRED_COLUMNS) - set(found))
        raise WorkbookFormatError(
            f"{path.name}/{sheet}: cabeçalho não encontrado nas "
            f"{HEADER_SEARCH_ROWS} primeiras linhas; colunas "
            f"obrigatórias ausentes: {missing}."
        )

    duplicated = sorted(
        name
        for name in set(header.values())
        if list(header.values()).count(name) > 1
    )

    if duplicated:
        raise WorkbookFormatError(
            f"{path.name}/{sheet}: colunas duplicadas no cabeçalho "
            f"(linha {header_row}): {duplicated}."
        )

    unknown = sorted(
        set(header.values())
        - set(REQUIRED_COLUMNS)
        - set(OPTIONAL_COLUMNS)
    )

    if unknown:
        raise WorkbookFormatError(
            f"{path.name}/{sheet}: colunas desconhecidas no cabeçalho "
            f"(linha {header_row}): {unknown}."
        )

    rows: list[WorkbookRow] = []
    empty_rows: list[int] = []

    for row_number in range(header_row + 1, ws.max_row + 1):
        values = {
            name: ws.cell(row=row_number, column=column).value
            for column, name in header.items()
        }

        residual = [
            ws.cell(row=row_number, column=column).coordinate
            for column in range(1, ws.max_column + 1)
            if column not in header
            and not _is_empty(ws.cell(row=row_number, column=column).value)
        ]

        if residual:
            raise WorkbookFormatError(
                f"{path.name}/{sheet}: conteúdo fora das colunas do "
                f"cabeçalho: {residual}."
            )

        if all(_is_empty(value) for value in values.values()):
            empty_rows.append(row_number)
            continue

        if _is_empty(values["name"]):
            filled = sorted(k for k, v in values.items() if not _is_empty(v))
            raise WorkbookFormatError(
                f"{path.name}/{sheet}: linha {row_number} tem conteúdo "
                f"({filled}) mas não tem 'name'."
            )

        rows.append(WorkbookRow(row=row_number, values=values))

    if not rows:
        raise WorkbookFormatError(
            f"{path.name}/{sheet}: nenhuma linha de dados abaixo do "
            f"cabeçalho (linha {header_row})."
        )

    return WorkbookData(
        path=path,
        file_name=path.name,
        sha256=sha256_of(path),
        sheet=sheet,
        header_row=header_row,
        columns=tuple(header.values()),
        rows=tuple(rows),
        empty_rows=tuple(empty_rows),
        last_sheet_row=ws.max_row,
    )
