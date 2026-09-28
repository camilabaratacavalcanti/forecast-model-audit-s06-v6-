"""
Etapa 2.5 — análise contratual (somente leitura).

    python audit/stage2_5_contract_closure/evidence/analysis_stage2_5.py

Entradas: os quatro workbooks aprovados em data/workbooks/ (area_41 v8,
energy v5, MaxHT v9, yield v9) e o candidato production v7 copiado byte a
byte em evidence/inputs/. Os builders são executados EM MEMÓRIA
(reader -> modelo canônico -> seeds), sem gravar em data/seed. A prova de
runtime do desaguamento_oee usa uma cópia temporária dos seeds.

Saídas (ao lado deste script):
    workbooks_sha256.txt, production_v6_v7_diff.csv,
    cross_workbook.csv, variable_value_inventory.csv,
    production_v7_desaguamento.csv, unreachable_reference_inventory.csv,
    identity_repeats_cross_block.csv, summary.json
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import sys
import tempfile
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(REPO))

import openpyxl  # noqa: E402

from app.engine import reference_resolver  # noqa: E402
from app.engine.scope_resolver import ScopeResolver  # noqa: E402
from app.engine.scoped_reference import scope_type_for  # noqa: E402
from app.engine.spatial_candidate_resolver import get_spatial_candidates  # noqa: E402
from tools.workbook_seed.blocks import BLOCKS, seed_json  # noqa: E402
from tools.workbook_seed.canonical import build_canonical_model  # noqa: E402
from tools.workbook_seed.reader import read_workbook, sha256_of  # noqa: E402
from tools.workbook_seed.seeds import build_seeds  # noqa: E402

V7 = OUT / "inputs" / "descritivo_das_variáveis_production_v7.xlsx"
V6 = BLOCKS["production"].workbook_path
FREQ_RANK = {"diário": 0, "mensal": 1, "anual": 2}
SR = ScopeResolver()
summary: dict = {}


def write_csv(name, rows):
    with open(OUT / name, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


# ------------------------------------------------------------------
# 0. workbooks: SHA-256 e diff v6 -> v7
# ------------------------------------------------------------------
sources = {b: (s.workbook_path, s.sheet) for b, s in BLOCKS.items()}
sources["production"] = (V7, "production")
lines = []
for block in sorted(sources):
    path = sources[block][0]
    digest = sha256_of(path)
    approved = BLOCKS[block].sha256
    status = (
        "APROVADO_2.4 (inalterado)" if digest == approved
        else "CANDIDATO (novo; substitui production v6 nesta análise)"
    )
    lines.append(f"{digest}  {path.name}  {block}  {status}")
lines.append(f"{sha256_of(V6)}  {V6.name}  production  v6 aprovado na 2.4 (referência para o diff)")
(OUT / "workbooks_sha256.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")

wa = openpyxl.load_workbook(V6, data_only=True)["production"]
wb = openpyxl.load_workbook(V7, data_only=True)["production"]
header = {c: wb.cell(2, c).value for c in range(1, wb.max_column + 1)}
diff = []
for r in range(1, max(wa.max_row, wb.max_row) + 1):
    for c in range(1, max(wa.max_column, wb.max_column) + 1):
        before, after = wa.cell(r, c).value, wb.cell(r, c).value
        if before != after:
            diff.append({
                "cell": wb.cell(r, c).coordinate, "row": r, "column": header.get(c),
                "name": wb.cell(r, 2).value, "v6": before, "v7": after,
            })
write_csv("production_v6_v7_diff.csv", diff)
summary["production_v6_v7_changed_cells"] = len(diff)

# ------------------------------------------------------------------
# 1. build em memória dos cinco workbooks (production = v7)
# ------------------------------------------------------------------
models, seeds, workbooks = {}, {}, {}
for block, (path, sheet) in sorted(sources.items()):
    wbk = read_workbook(path, sheet)
    model = build_canonical_model(block, wbk, BLOCKS[block].id_base)
    workbooks[block], models[block] = wbk, model
    seeds[block] = build_seeds(model, BLOCKS[block].id_base)

official = {
    b: {n: json.loads((REPO / "data/seed" / b / f"{n}.json").read_text(encoding="utf-8"))
        for n in ("variables", "parameters", "equations", "aggregation_rules", "manifest")}
    for b in BLOCKS
}
seed_delta = []
for name in ("variables", "parameters", "equations", "aggregation_rules"):
    new, old = seeds["production"][name], official["production"][name]
    if len(new) != len(old):
        seed_delta.append(f"{name}: {len(old)} -> {len(new)} registros")
    for a, b in zip(old, new):
        if a != b:
            keys = sorted(k for k in set(a) | set(b) if a.get(k) != b.get(k))
            seed_delta.append(f"{name} {b.get('equation_id') or b.get('variable_id') or b.get('parameter_id') or b.get('aggregation_rule_id')}: {keys}")
summary["production_v7_seed_delta_vs_official_v6"] = seed_delta
summary["production_v7_pending_decisions"] = seeds["production"]["manifest"]["pending_contract_decisions"]
summary["other_blocks_seed_equal_official"] = {
    b: all(seeds[b][n] == official[b][n] for n in official[b]) for b in BLOCKS if b != "production"
}
summary["counts"] = {
    b: {n: len(seeds[b][n]) for n in ("variables", "parameters", "equations", "aggregation_rules")}
    for b in seeds
}

# ------------------------------------------------------------------
# helpers
# ------------------------------------------------------------------
def entity_rows(block):
    return {e.entity_id: e for e in models[block].entities}


def raw_rows(block):
    wbk = workbooks[block]
    return {r.row: r.values for r in wbk.rows}


# ------------------------------------------------------------------
# 2. variáveis com value/version (varredura completa)
# ------------------------------------------------------------------
value_inventory = []
type_counter = Counter()
for block in sorted(sources):
    for r in workbooks[block].rows:
        v = r.values
        type_counter[(block, v["Type"])] += 1
        if v["Type"] != "parameter" and (v.get("value") is not None or v.get("version") is not None):
            value_inventory.append({
                "block": block, "workbook": workbooks[block].file_name, "row": r.row,
                "type": v["Type"], "name": v["name"], "frequency": v["frequency"],
                "scope_type": v["scope_type"], "scope_value": v["scope_value"],
                "value": v.get("value"), "value_physical_type": type(v.get("value")).__name__,
                "version": v.get("version"), "value_type": v["value_type"],
                "variable_type": v.get("variable_type"), "status": v["status"],
                "unit": v["unit"], "source_reference": v.get("source_reference"),
                "has_expression": v.get("expression") is not None,
            })
if not value_inventory:
    value_inventory.append({"block": "(nenhuma)"})
write_csv("variable_value_inventory.csv", value_inventory)
summary["variable_value_occurrences"] = len([x for x in value_inventory if x.get("row")])
# parâmetros relacionados por nome às ocorrências
related = []
for x in value_inventory:
    if not x.get("row"):
        continue
    for block in sorted(sources):
        for r in workbooks[block].rows:
            v = r.values
            if v["Type"] == "parameter" and v["name"] in (x["name"], x["name"] + "_base"):
                related.append((block, r.row, v["name"], v["value"], v["frequency"], v["scope_type"], v["scope_value"]))
summary["parameters_related_to_value_rows"] = related

# ------------------------------------------------------------------
# 3. referências de equações: alcançabilidade espacial/temporal
# ------------------------------------------------------------------
REF = re.compile(r"\b(VAR\d+|PARAM\d+)(?:@(L[1-7](?:_L[1-7])?))?")
unreach = []
ref_class = Counter()
for block in sorted(sources):
    ents = entity_rows(block)
    manifest_rows = {e["equation_id"]: e["row"] for e in seeds[block]["manifest"]["equations"]}
    index = reference_resolver.build_name_index([e.as_index_entry() for e in models[block].entities])
    for eq in seeds[block]["equations"]:
        target = ents[eq["target_variable_id"]]
        consumer_scopes = SR.resolve_scopes(eq["scope_type"], eq["scope_value"])
        for m in REF.finditer(eq["expression"]):
            ref = ents[m.group(1)]
            explicit = m.group(2)
            candidates = index[ref.name]
            same_freq = [c for c in candidates if c.get("frequency") == target.frequency]
            params = [c for c in candidates if c["kind"] == "parameter"]
            inst = set(SR.resolve_scopes(ref.scope_type, ref.scope_value))
            spatial = []
            for cs in consumer_scopes:
                if explicit:
                    spatial.append("EXPLICIT_OK" if (scope_type_for(explicit), explicit) in inst else "EXPLICIT_MISSING")
                else:
                    chain = [tuple(x) for x in get_spatial_candidates(*cs)]
                    hit = next((i for i, x in enumerate(chain) if x in inst), None)
                    spatial.append("UNREACHABLE" if hit is None else ("OWN_SCOPE" if hit == 0 else "PARENT_SCOPE"))
            if ref.kind == "parameter":
                temporal = "PARAMETER (ano)"
                freq_fallback = False
            else:
                rc, rr = FREQ_RANK[target.frequency], FREQ_RANK[ref.frequency]
                temporal = "SAME" if rc == rr else ("COARSER_REF (fallback temporal do runtime)" if rr > rc else "FINER_REF (inalcançável)")
                freq_fallback = not params and not same_freq
            spatial_set = sorted(set(spatial))
            if "UNREACHABLE" in spatial or "EXPLICIT_MISSING" in spatial or temporal.startswith("FINER"):
                verdict = "UNREACHABLE"
            elif freq_fallback:
                verdict = "RESOLVED_BY_FREQUENCY_FALLBACK"
            elif "PARENT_SCOPE" in spatial:
                verdict = "RESOLVED_BY_SPATIAL_PARENT"
            elif temporal.startswith("COARSER"):
                verdict = "RESOLVED_BY_TEMPORAL_FALLBACK"
            else:
                verdict = "DIRECT"
            ref_class[verdict] += 1
            unreach.append({
                "block": block, "equation_id": eq["equation_id"], "row": manifest_rows[eq["equation_id"]],
                "consumer": target.name, "consumer_frequency": target.frequency,
                "consumer_scope": f"{eq['scope_type']}/{eq['scope_value']}",
                "reference": m.group(0), "reference_name": ref.name, "reference_kind": ref.kind,
                "reference_frequency": ref.frequency, "reference_scope": f"{ref.scope_type}/{ref.scope_value}",
                "reference_unit": ref.unit, "candidates_by_name": len(candidates),
                "same_frequency_candidates": len(same_freq), "frequency_fallback": freq_fallback,
                "temporal": temporal, "spatial": "|".join(spatial_set), "verdict": verdict,
            })
write_csv("unreachable_reference_inventory.csv", unreach)
summary["equations_analysed"] = sum(len(seeds[b]["equations"]) for b in seeds)
summary["references_analysed"] = len(unreach)
summary["reference_verdicts"] = dict(ref_class)
summary["unreachable"] = [
    {k: r[k] for k in ("block", "row", "consumer", "reference_name", "reference_frequency", "reference_scope")}
    for r in unreach if r["verdict"] == "UNREACHABLE"
]
summary["frequency_fallback"] = [
    {k: r[k] for k in ("block", "row", "consumer", "consumer_frequency", "reference_name", "reference_frequency")}
    for r in unreach if r["verdict"] == "RESOLVED_BY_FREQUENCY_FALLBACK"
]

# ------------------------------------------------------------------
# 4. production v7: desaguamento_oee -> consumo_mpsa_grupo
# ------------------------------------------------------------------
prod = entity_rows("production")
by_name = defaultdict(list)
for e in models["production"].entities:
    by_name[e.name].append(e)
eqs = {e["target_variable_id"] + "@" + e["scope_value"]: e for e in seeds["production"]["equations"]}
rules = {r["target_variable_id"]: r for r in seeds["production"]["aggregation_rules"]}
dea = next(e for e in by_name["desaguamento_oee"])
dea_eq = next(e for e in seeds["production"]["equations"] if e["target_variable_id"] == dea.entity_id)
chain = []


def describe(step, e, mechanism, detail, check):
    chain.append({
        "step": step, "entity_id": e.entity_id, "kind": e.kind, "name": e.name,
        "frequency": e.frequency, "scope": f"{e.scope_type}/{e.scope_value}", "unit": e.unit,
        "value_type": e.value_type, "workbook_rows": " ".join(str(r.row) for r in e.rows),
        "instances": " ".join(s for _t, s in SR.resolve_scopes(e.scope_type, e.scope_value)),
        "producer": mechanism, "resolution_detail": detail, "check": check,
    })


describe("1 consumidor", dea, f"equação {dea_eq['equation_id']}: {dea_eq['expression']}",
         f"workbook P25: {raw_rows('production')[25]['expression']}", "equação anual linha_grupo/L1_L7")
refs = REF.findall(dea_eq["expression"])
for rid, _ in refs:
    e = prod[rid]
    if e.kind == "variable":
        cands = [c for c in by_name[e.name]]
        same = [c for c in cands if c.frequency == dea.frequency]
        reach = any(tuple(x) in set(SR.resolve_scopes(e.scope_type, e.scope_value))
                    for x in get_spatial_candidates("linha_grupo", "L1_L7"))
        rule = rules.get(e.entity_id)
        describe("2 referência", e,
                 f"AggregationRule {rule['aggregation_rule_id']} ({rule['aggregation_type']}, fator {rule.get('integration_factor', 1.0)})" if rule else "equação",
                 f"candidatos por nome: {[(c.entity_id, c.frequency) for c in cands]}; mesma frequência (anual): {[c.entity_id for c in same]}; fallback de frequência: {not same}",
                 "OK: mesma frequência, mesmo escopo, alcançável" if same == [e] and reach else "FALHA")
        source = prod[rule["source_variable_id"]]
        seq = eqs[source.entity_id + "@" + source.scope_value]
        describe("3 origem da agregação", source, f"equação {seq['equation_id']}: {seq['expression']}",
                 f"regra {rule['source_frequency']} -> {rule['target_frequency']}; unidade {source.unit} -> {e.unit}",
                 "OK: SUM t/d -> t/ano (fator 1)" if source.unit == "t/d" and e.unit == "t/ano" else "VERIFICAR")
        for lid, sc in REF.findall(seq["expression"]):
            le = prod[lid]
            leq = eqs[le.entity_id + "@" + le.scope_value]
            describe(f"4 termo @{sc}", le, f"equação {leq['equation_id']}: {leq['expression']}",
                     f"referência explícita @{sc}", "OK: instância declarada" if ("linha", sc) in SR.resolve_scopes(le.scope_type, le.scope_value) else "FALHA")
    else:
        recs = [p for p in seeds["production"]["parameters"] if p["parameter_id"] == rid]
        describe("2 referência", e, f"parameter value={[p['value'] for p in recs]} version={[p['version'] for p in recs]}",
                 "parâmetros vêm antes na A019; consumidor linha_grupo/L1_L7 alcança planta/PLANTA pela cadeia espacial (Decision E)",
                 "OK: alcançável via planta" if ("planta", "PLANTA") in [tuple(x) for x in get_spatial_candidates("linha_grupo", "L1_L7")] else "FALHA")
old_ref = "consumo_mpsa /" in (raw_rows("production")[25]["expression"] or "")
summary["desaguamento_old_reference_present_in_v7"] = old_ref

# prova de runtime: seeds temporários com production v7
from app.engine.calculation_context import CalculationContext  # noqa: E402
from app.engine.equation_engine import EquationEngine  # noqa: E402
from app.domain.equations.models import EquationInstance  # noqa: E402
from app.engine.temporal_aggregation_service import TemporalAggregationService  # noqa: E402
from app.repositories.seed_loader import SeedLoader  # noqa: E402

tmp = Path(tempfile.mkdtemp())
try:
    shutil.copytree(REPO / "data" / "seed", tmp / "seed")
    for n in ("variables", "parameters", "equations", "aggregation_rules", "manifest"):
        (tmp / "seed" / "production" / f"{n}.json").write_text(seed_json(seeds["production"][n]), encoding="utf-8")
    from app.validation import equation_seed_validator, parameter_seed_validator, variable_seed_validator  # noqa: E402
    vres = variable_seed_validator.validate_seed(tmp / "seed")
    perr, _pw = parameter_seed_validator.validate_seed(tmp / "seed")
    eerr, _ew = equation_seed_validator.validate_seed(tmp / "seed")
    serr = equation_seed_validator.validate_expression_syntax(seeds["production"]["equations"])
    summary["production_v7_validators"] = {
        "variable_errors": vres["errors"], "variable_warnings": len(vres["warnings"]),
        "parameter_errors": perr, "equation_errors": eerr, "production_expression_syntax_errors": serr,
    }
    loader = SeedLoader(tmp / "seed")
    var_defs, _vi, par_defs, _pi, eq_defs, _ei = loader.load_all_definitions_and_instances()
    rule_objs = loader.load_aggregation_rules()
    from app.engine.registry_validator import RegistryIntegrityValidator  # noqa: E402
    RegistryIntegrityValidator().validate_definition_registry(eq_defs, var_defs, par_defs)

    grupo_d = next(e for e in by_name["consumo_mpsa_grupo"] if e.frequency == "diário")
    grupo_a = next(e for e in by_name["consumo_mpsa_grupo"] if e.frequency == "anual")
    mpsa = by_name["consumo_mpsa"][0]
    ctx = CalculationContext()
    days = [date(2026, 1, d) for d in (1, 2, 3)]
    engine = EquationEngine()
    grupo_def = eq_defs.get(eqs[grupo_d.entity_id + "@L1_L7"]["equation_id"])
    for i, day in enumerate(days, start=1):
        for k, line in enumerate(("L1", "L2", "L3", "L4", "L5", "L6", "L7"), start=1):
            ctx.set_variable_value(mpsa.entity_id, 10.0 * k + i, "linha", line, day.isoformat())
        inst = EquationInstance.create(definition=grupo_def, scope_type="linha_grupo", scope_value="L1_L7")
        value = engine.calculate_instance(instance=inst, definition=grupo_def, calculation_context=ctx, period_id=day.isoformat())
        ctx.set_variable_value(grupo_d.entity_id, value, "linha_grupo", "L1_L7", day.isoformat())
    annual = TemporalAggregationService().aggregate(
        rule_objs.get(rules[grupo_a.entity_id]["aggregation_rule_id"]), ctx, "linha_grupo", "L1_L7", days[-1]
    )
    ctx.set_variable_value(grupo_a.entity_id, annual.value, "linha_grupo", "L1_L7", annual.period_id)
    par = next(p for p in par_defs.all() if p.parameter_name == "desaguamento_produtividade")
    ctx.set_parameter_value(par.parameter_definition_id, par.value, "planta", "PLANTA")
    dea_def = eq_defs.get(dea_eq["equation_id"])
    inst = EquationInstance.create(definition=dea_def, scope_type="linha_grupo", scope_value="L1_L7")
    result = engine.calculate_instance(instance=inst, definition=dea_def, calculation_context=ctx, period_id=annual.period_id)
    expected_annual = sum(10.0 * k + i for i in (1, 2, 3) for k in range(1, 8))
    expected = 100 * (expected_annual / 24) / (115 * 13)
    runtime = {
        "consumo_mpsa_grupo_anual (SUM de 3 dias)": annual.value, "esperado": expected_annual,
        "desaguamento_oee": result, "esperado_oee": expected,
        "ok": abs(annual.value - expected_annual) < 1e-9 and abs(result - expected) < 1e-9,
        "registry_integrity": "OK",
    }
finally:
    shutil.rmtree(tmp)
chain.append({
    "step": "5 runtime", "entity_id": dea.entity_id, "kind": "variable", "name": "desaguamento_oee",
    "frequency": "anual", "scope": "linha_grupo/L1_L7", "unit": dea.unit, "value_type": dea.value_type,
    "workbook_rows": "25", "instances": "L1_L7",
    "producer": "SeedLoader(seeds temporários v7) -> EquationEngine + TemporalAggregationService",
    "resolution_detail": json.dumps(runtime, ensure_ascii=False),
    "check": "OK" if runtime["ok"] else "FALHA",
})
write_csv("production_v7_desaguamento.csv", chain)
summary["desaguamento_runtime"] = runtime

# ------------------------------------------------------------------
# 5. cross-workbook: relações declaradas por `fonte` + repetições de identidade
# ------------------------------------------------------------------
FIVE = {"area_41": "area_41", "energy": "energy", "max_ht": "max_ht", "production": "production", "yield": "yield"}
FONTE_BLOCK = re.compile(r"bloco\s+([\w]+)", re.I)


def producers_of(block, name, frequency, scope_type):
    out = []
    for e in models[block].entities:
        if e.name != name:
            continue
        if e.kind == "parameter" or (e.frequency == frequency and e.scope_type == scope_type):
            produced = (
                "parameter" if e.kind == "parameter"
                else ("equation" if any(q["target_variable_id"] == e.entity_id for q in seeds[block]["equations"])
                      else "aggregation" if e.entity_id in {r["target_variable_id"] for r in seeds[block]["aggregation_rules"]}
                      else "input")
            )
            out.append((e, produced))
    return out


cross = []
external = []
for block in sorted(sources):
    for e in models[block].entities:
        if e.kind != "variable":
            continue
        fontes = {raw_rows(block)[r.row].get("fonte") for r in e.rows} - {None}
        for fonte in fontes:
            m = FONTE_BLOCK.search(fonte or "")
            target = m.group(1).lower() if m else None
            if target not in FIVE:
                external.append({"consumer_block": block, "name": e.name, "frequency": e.frequency,
                                 "scope": f"{e.scope_type}/{e.scope_value}", "fonte": fonte})
                continue
            prods = producers_of(target, e.name, e.frequency, e.scope_type)
            issues = []
            representation = []
            if not prods:
                issues.append("PRODUTOR_INEXISTENTE")
            elif len(prods) > 1:
                issues.append("PRODUTOR_AMBIGUO")
            p = prods[0][0] if prods else None
            if p is not None:
                if p.kind != e.kind:
                    issues.append(f"KIND {p.kind}->{e.kind}")
                if p.unit != e.unit:
                    issues.append(f"UNIT {p.unit}!={e.unit}")
                if p.value_type != e.value_type:
                    issues.append("VALUE_TYPE")
                if p.kind == "variable" and p.frequency != e.frequency:
                    issues.append("FREQUENCY")
                if p.kind == "parameter" and p.frequency != e.frequency:
                    issues.append(f"FREQUENCY param {p.frequency}!={e.frequency}")
                if (p.scope_type, p.scope_value) != (e.scope_type, e.scope_value):
                    issues.append(f"SCOPE {p.scope_type}/{p.scope_value}!={e.scope_type}/{e.scope_value}")
                if set(SR.resolve_scopes(p.scope_type, p.scope_value)) != set(SR.resolve_scopes(e.scope_type, e.scope_value)):
                    issues.append("INSTANCES: conjuntos de instâncias diferentes")
                elif p.grouped != e.grouped:
                    representation.append(
                        "mesmas instâncias; produtor declara uma linha por instância, consumidor uma linha L1_L7"
                    )
                if prods[0][1] == "input":
                    issues.append("PRODUTOR_TAMBEM_E_ENTRADA")
            cross.append({
                "consumer_block": block, "consumer_rows": " ".join(str(r.row) for r in e.rows),
                "consumer_identity": f"{e.name}|{e.frequency}|{e.scope_type}/{e.scope_value}",
                "consumer_id": e.entity_id, "consumer_type": f"variable/{e.variable_type}",
                "consumer_fonte": fonte,
                "producer_block": target,
                "producer_identity": f"{p.name}|{p.frequency}|{p.scope_type}/{p.scope_value}" if p else "",
                "producer_id": p.entity_id if p else "",
                "producer_rows": " ".join(str(r.row) for r in p.rows) if p else "",
                "producer_type": (f"{p.kind}/{p.variable_type or ''}".rstrip("/") + f" ({prods[0][1]})") if p else "",
                "frequency": e.frequency, "scope_type": e.scope_type, "scope_value": e.scope_value,
                "unit": e.unit if (p is None or p.unit == e.unit) else f"{e.unit} vs {p.unit}",
                "value_type": e.value_type if (p is None or p.value_type == e.value_type) else f"{e.value_type} vs {p.value_type}",
                "current_link_mechanism": "nenhum: definição local no consumidor (ID próprio); `fonte` é texto livre; sem ligação por ID",
                "compatibility": ("COMPATIVEL" if not issues else "; ".join(issues))
                + (f" ({'; '.join(representation)})" if representation else ""),
                "decision_required": "D24-11",
            })
write_csv("cross_workbook.csv", cross)
write_csv("external_sources_outside_five.csv", external)
summary["cross_workbook_relations"] = len(cross)
summary["cross_workbook_incompatibilities"] = [
    (c["consumer_block"], c["consumer_identity"], c["compatibility"]) for c in cross if not c["compatibility"].startswith("COMPATIVEL")
]
summary["external_sources_outside_five"] = len(external)

# repetições de identidade entre blocos que NÃO vêm de `fonte`
declared = {(c["consumer_block"], c["consumer_identity"]) for c in cross}
identity_blocks = defaultdict(list)
for block in sorted(sources):
    for e in models[block].entities:
        if e.kind != "variable":
            continue
        produced = (
            "equation" if any(q["target_variable_id"] == e.entity_id for q in seeds[block]["equations"])
            else "aggregation" if e.entity_id in {r["target_variable_id"] for r in seeds[block]["aggregation_rules"]}
            else "input"
        )
        fonte = {raw_rows(block)[r.row].get("fonte") for r in e.rows} - {None}
        identity_blocks[(e.name, e.frequency, e.scope_type, e.scope_value)].append(
            (block, e.entity_id, produced, ";".join(sorted(fonte)), e.unit)
        )
repeats = []
for key, occ in sorted(identity_blocks.items()):
    if len({b for b, *_ in occ}) < 2:
        continue
    producers = [o for o in occ if o[2] != "input"]
    inputs = [o for o in occ if o[2] == "input"]
    if len(producers) > 1:
        cls = "IDENTIDADE CALCULADA EM MAIS DE UM BLOCO"
    elif len(producers) == 1 and all(o[3] for o in inputs):
        cls = "LIGAÇÃO DECLARADA POR fonte"
    elif len(producers) == 1:
        cls = "ENTRADA SEM fonte PARA IDENTIDADE PRODUZIDA EM OUTRO BLOCO"
    else:
        cls = "SOMENTE ENTRADAS (produtor fora dos cinco ou inexistente)"
    repeats.append({
        "identity": "|".join(key), "blocks": " ".join(f"{b}:{i}:{p}" for b, i, p, _f, _u in occ),
        "fontes": " | ".join(f"{b}:{f or '-'}" for b, _i, _p, f, _u in occ),
        "units": " ".join(sorted({u for *_x, u in occ})), "class": cls,
        "in_fonte_inventory": any((b, "|".join([key[0], key[1], f"{key[2]}/{key[3]}"])) in declared for b, *_ in occ),
    })
write_csv("identity_repeats_cross_block.csv", repeats)
summary["identity_repeats_cross_block"] = Counter(r["class"] for r in repeats)

(OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=1, default=str))
