"""
Etapa 2.4 — rodada 2: reexecução integral do cross-audit.

Reexecutável a partir da raiz do repositório:

    python audit/stage2_4_cross_audit/final_round2/evidence/audit_round2.py

Somente leitura: lê os workbooks aprovados (data/workbooks), executa os
builders EM MEMÓRIA, compara com os seeds versionados, carrega os seeds
pelo SeedLoader real e roda as verificações de contrato. Grava apenas
os arquivos de evidência ao lado deste script.
"""

from __future__ import annotations

import collections
import csv
import json
import re
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[4]
OUT = Path(__file__).resolve().parent
sys.path.insert(0, str(REPO))

import openpyxl  # noqa: E402

from app.domain.values import RESULT_STATE_TAXONOMY, VALUE_TYPES  # noqa: E402
from app.engine import reference_resolver  # noqa: E402
from app.engine.scope_resolver import ScopeResolver  # noqa: E402
from app.repositories.seed_loader import SeedLoader  # noqa: E402
from app.validation import (  # noqa: E402
    equation_seed_validator,
    parameter_seed_validator,
    variable_seed_validator,
)
from app.validation.aggregation_dimension_validator import (  # noqa: E402
    find_sum_dimension_issues,
)
from tools.workbook_seed.blocks import BLOCKS, SEED_FILES, build_block, read_seed_file  # noqa: E402
from tools.workbook_seed.reader import sha256_of  # noqa: E402

SEED_ROOT = REPO / "data" / "seed"
HIDDEN = re.compile(r"[ ​‌‍﻿\t\r]")
LINES = ["L1", "L2", "L3", "L4", "L5", "L6", "L7"]
summary: dict = {}


def write_csv(name, rows):
    with open(OUT / name, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


# ------------------------------------------------------------------
# 1. hashes
# ------------------------------------------------------------------
hashes = []
for block, spec in sorted(BLOCKS.items()):
    digest = sha256_of(spec.workbook_path)
    hashes.append(f"{digest}  {spec.file_name}  ({block} {spec.version}) approved={'yes' if digest == spec.sha256 else 'NO'}")
(OUT / "workbook_sha256.txt").write_text("\n".join(hashes) + "\n", encoding="utf-8")
summary["workbooks_approved"] = all("approved=yes" in h for h in hashes)

# ------------------------------------------------------------------
# 2. auditoria célula a célula
# ------------------------------------------------------------------
cells = []
cell_issues = []
for block, spec in sorted(BLOCKS.items()):
    ws = openpyxl.load_workbook(spec.workbook_path, data_only=True)[spec.sheet]
    header_row = next(r for r in range(1, 11) if "name" in [c.value for c in ws[r]])
    header = {c.column: c.value for c in ws[header_row] if c.value}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            column = header.get(c.column, "<fora do cabeçalho>")
            cells.append((block, c.coordinate, column, type(c.value).__name__))
            if c.row > header_row and column == "<fora do cabeçalho>":
                cell_issues.append((block, c.coordinate, "conteúdo fora do cabeçalho", repr(c.value)))
            if isinstance(c.value, str) and c.row > header_row and column not in ("description", "OBS", "expression"):
                if HIDDEN.search(c.value) or (column != "allowed_values" and c.value != c.value.strip()):
                    cell_issues.append((block, c.coordinate, "espaço/caractere oculto", repr(c.value)))
            if c.row > header_row and column == "value_type" and c.value not in VALUE_TYPES:
                cell_issues.append((block, c.coordinate, "value_type fora do vocabulário", repr(c.value)))
            if c.row > header_row and column == "value" and ws.cell(c.row, [k for k, v in header.items() if v == "Type"][0]).value == "parameter":
                if isinstance(c.value, bool) or not isinstance(c.value, (int, float)):
                    cell_issues.append((block, c.coordinate, "parâmetro não numérico", repr(c.value)))
            if c.row > header_row and column == "declared_result_states":
                for line in str(c.value).split("\n"):
                    m = re.fullmatch(r'([A-Z_]+) → "([^"]+)"', line)
                    if not m or m.group(1) not in RESULT_STATE_TAXONOMY:
                        cell_issues.append((block, c.coordinate, "declared_result_states fora do formato", repr(c.value)))
summary["cells_audited"] = len(cells)
summary["cell_issues"] = len(cell_issues)
with open(OUT / "cell_audit.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["block", "cell", "column", "physical_type"])
    w.writerows(cells)
with open(OUT / "cell_issues.csv", "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(["block", "cell", "issue", "value"])
    w.writerows(cell_issues)

# ------------------------------------------------------------------
# 3. builders sobre os cinco workbooks + builder == seed
# ------------------------------------------------------------------
builds = {block: build_block(block) for block in sorted(BLOCKS)}
builder_lines = []
builder_equal = True
for block, result in builds.items():
    wb = result.workbook
    counts = {n: len(result.seeds[n]) for n in ("variables", "parameters", "equations", "aggregation_rules")}
    equal = {n: result.seeds[n] == read_seed_file(block, n) for n in SEED_FILES}
    builder_equal &= all(equal.values())
    grouped = [(e.name, e.frequency, e.scope_value, [r.row for r in e.rows]) for e in result.model.entities if e.grouped]
    builder_lines.append(
        f"{block}: {wb.file_name} sheet={wb.sheet} header_row={wb.header_row} data_rows={len(wb.rows)} "
        f"empty_rows_ignored={len(wb.empty_rows)} last_sheet_row={wb.last_sheet_row} "
        f"entities={len(result.model.entities)} seeds={counts} builder==seed={equal}\n"
        f"   definições com instâncias por linha: {grouped}\n"
        f"   entidades sem name: {sum(1 for e in result.model.entities if not e.name)}"
    )
(OUT / "builders.txt").write_text("\n".join(builder_lines) + "\n", encoding="utf-8")
summary["builder_equals_seed"] = builder_equal
summary["rows"] = {b: len(r.workbook.rows) for b, r in builds.items()}

# ------------------------------------------------------------------
# 4. workbook × canonical × seed (todas as linhas)
# ------------------------------------------------------------------
recon = []
for block, result in builds.items():
    manifest = read_seed_file(block, "manifest")
    eq_rows = {e["row"]: e["equation_id"] for e in manifest["equations"]}
    rule_rows = {r["row"]: r["aggregation_rule_id"] for r in manifest["aggregation_rules"]}
    for entity in result.model.entities:
        for r in entity.rows:
            recon.append({
                "workbook": block, "row": r.row, "kind": entity.kind, "name": entity.name,
                "frequency": entity.frequency, "row_scope": f"{entity.scope_type}/{r.scope_value}",
                "definition_id": entity.entity_id,
                "definition_scope": f"{entity.scope_type}/{entity.scope_value}",
                "instances_declared": entity.grouped,
                "value_type": entity.value_type,
                "equation_id": eq_rows.get(r.row, ""),
                "aggregation_rule_id": rule_rows.get(r.row, ""),
            })
write_csv("workbook_vs_seed.csv", recon)
summary["reconciled_rows"] = len(recon)

# ------------------------------------------------------------------
# 5. validadores oficiais + SeedLoader
# ------------------------------------------------------------------
var_result = variable_seed_validator.validate_seed(SEED_ROOT)
param_errors, param_warnings = parameter_seed_validator.validate_seed(SEED_ROOT)
eq_errors, eq_warnings = equation_seed_validator.validate_seed(SEED_ROOT)
syntax_errors = equation_seed_validator.validate_expression_syntax(
    [e for b in BLOCKS for e in read_seed_file(b, "equations")]
)
loader = SeedLoader(SEED_ROOT)
(var_defs, var_inst, par_defs, par_inst, eq_defs, eq_inst) = loader.load_all_definitions_and_instances()
rules = loader.load_aggregation_rules()
rule_instances = loader.load_aggregation_rule_instances(variable_definition_registry=var_defs)
from app.engine.registry_validator import RegistryIntegrityValidator  # noqa: E402
RegistryIntegrityValidator().validate_definition_registry(eq_defs, var_defs, par_defs)
validators = {
    "variable_errors": var_result["errors"], "variable_warnings": var_result["warnings"],
    "parameter_errors": param_errors, "parameter_warnings": param_warnings,
    "equation_errors": eq_errors, "expression_syntax_errors": syntax_errors,
    "sum_dimension_issues": [i.message for i in find_sum_dimension_issues(rules.all(), var_defs)],
    "loaded": {
        "variable_definitions": len(var_defs.all()), "variable_instances": len(var_inst.all()),
        "parameter_definitions": len(par_defs.all()), "parameter_instances": len(par_inst.all()),
        "equation_definitions": len(eq_defs.all()), "equation_instances": len(eq_inst.all()),
        "aggregation_rules": len(rules.all()), "aggregation_rule_instances": len(rule_instances.all()),
    },
    "registry_integrity": "OK",
}
(OUT / "validators.json").write_text(json.dumps(validators, ensure_ascii=False, indent=1), encoding="utf-8")
summary["validator_errors"] = len(var_result["errors"]) + len(param_errors) + len(eq_errors) + len(syntax_errors)

# ------------------------------------------------------------------
# 6. A019 — 20 casos
# ------------------------------------------------------------------
A019 = [
    ("area_41", 22, "valor_retirada", None, "CONFORME"), ("area_41", 39, "hes", None, "CONFORME"),
    ("area_41", 42, "hes", None, "CONFORME"), ("area_41", 52, "retirada_condensado_linha", "diário", "GAP"),
    ("area_41", 53, "retirada_condensado_linha", "diário", "GAP"), ("production", 35, "lth_meta", None, "GAP"),
    ("production", 48, "lth_meta", None, "GAP"), ("production", 50, "lth_meta", None, "CONFORME"),
    ("production", 60, "pick_up", "diário", "GAP"), ("production", 61, "pick_up", None, "CONFORME"),
    ("production", 71, "pick_up", None, "GAP"),
    *[("production", r, "pick_up_yield", None, "CONFORME") for r in range(53, 60)],
    ("yield", 107, "tanque", None, "GAP"), ("yield", 107, "tanque_base", None, "GAP"),
]
a019_rows = []
sr = ScopeResolver()
for block, row, name, lookup_frequency, before in A019:
    model = builds[block].model
    index = reference_resolver.build_name_index([e.as_index_entry() for e in model.entities])
    consumer = next(e for e in model.entities if any(r.row == row for r in e.rows))
    crow = next(r for r in consumer.rows if r.row == row)
    explicit = re.findall(rf"(?<![\w@]){name}@(L\d(?:_L\d)?)", crow.expression)
    resolved = set()
    try:
        for scope in explicit or [None]:
            resolved.add(reference_resolver.resolve_reference(
                name, index, lookup_frequency or consumer.frequency,
                consumer_scope=(consumer.scope_type, crow.scope_value), explicit_scope_value=scope,
            ))
        outcome = "RESOLVED"
    except Exception as ex:  # registrado, nunca escondido
        outcome = f"{type(ex).__name__}: {ex}"
    target = model.by_id()[sorted(resolved)[0]] if resolved else None
    instances = (
        [s for _t, s in sr.resolve_scopes(consumer.scope_type, crow.scope_value)] if not explicit else list(dict.fromkeys(explicit))
    )
    a019_rows.append({
        "n": len(a019_rows) + 1, "workbook": block, "row": row, "consumer": consumer.name,
        "consumer_scope": f"{consumer.scope_type}/{crow.scope_value}", "reference": name,
        "stage_2_3": before, "outcome": outcome, "definition_ids": " ".join(sorted(resolved)),
        "definition_scope": f"{target.scope_type}/{target.scope_value}" if target else "",
        "definition_rows": " ".join(str(r.row) for r in target.rows) if target else "",
        "consumer_instances": " ".join(instances),
    })
write_csv("a019_20.csv", a019_rows)
summary["a019_resolved"] = sum(r["outcome"] == "RESOLVED" for r in a019_rows)

# ------------------------------------------------------------------
# 7. agregações
# ------------------------------------------------------------------
agg_rows = []
for block, result in builds.items():
    names = {e.entity_id: e for e in result.model.entities}
    row_of = {r["aggregation_rule_id"]: r["row"] for r in read_seed_file(block, "manifest")["aggregation_rules"]}
    for rule in result.seeds["aggregation_rules"]:
        t, s = names[rule["target_variable_id"]], names[rule["source_variable_id"]]
        w = names.get(rule.get("weight_variable_id"))
        agg_rows.append({
            "workbook": block, "row": row_of[rule["aggregation_rule_id"]],
            "aggregation_rule_id": rule["aggregation_rule_id"], "type": rule["aggregation_type"],
            "target": f"{t.name}|{t.frequency}|{t.scope_type}/{t.scope_value}|{t.unit}",
            "source": f"{s.name}|{s.frequency}|{s.scope_type}/{s.scope_value}|{s.unit}",
            "weight": f"{w.name}|{w.frequency}" if w else "",
            "integration_factor": rule.get("integration_factor", ""),
            "instances": sum(1 for i in rule_instances.all() if i.rule.aggregation_rule_id == rule["aggregation_rule_id"]),
        })
write_csv("aggregations.csv", agg_rows)
summary["aggregations"] = collections.Counter(r["workbook"] for r in agg_rows)
summary["aggregation_types"] = collections.Counter(r["type"] for r in agg_rows)

# ------------------------------------------------------------------
# 8. cross-workbook (13 vínculos) — estado, sem decisão
# ------------------------------------------------------------------
LINKS = [
    ("area_41", "lth", "diário", "linha", "yield"), ("energy", "producao", "diário", "linha", "production"),
    ("energy", "pick_up", "diário", "linha", "production"),
    ("energy", "pick_up_total", "diário", "linha_grupo", "production"),
    ("energy", "pick_up_total", "mensal", "linha_grupo", "production"),
    ("energy", "lth", "diário", "linha", "production"), ("energy", "lth_total", "diário", "linha_grupo", "production"),
    ("energy", "lth_total", "mensal", "linha_grupo", "production"),
    ("energy", "lth_meta", "anual", "linha", "production"), ("max_ht", "lth", "diário", "linha", "production"),
    ("max_ht", "producao", "diário", "linha", "production"), ("production", "yield", "diário", "linha", "yield"),
    ("yield", "lth", "diário", "linha", "production"),
]
cross = []
for consumer, name, frequency, scope_type, producer in LINKS:
    local = [v for v in read_seed_file(consumer, "variables")
             if v["variable_name"] == name and v["frequency"] == frequency and v["scope_type"] == scope_type]
    prod = [v for v in read_seed_file(producer, "variables")
            if v["variable_name"] == name and v["frequency"] == frequency and v["scope_type"] == scope_type]
    prod_p = [p for p in read_seed_file(producer, "parameters") if p["parameter_name"] == name]
    cross.append({
        "consumer": consumer, "name": name, "frequency": frequency, "scope_type": scope_type,
        "consumer_id": ",".join(v["variable_id"] for v in local),
        "consumer_variable_type": ",".join(v["variable_type"] for v in local),
        "producer": producer,
        "producer_ids": ",".join([v["variable_id"] for v in prod] + sorted({p["parameter_id"] for p in prod_p})),
        "producer_kind": "parameter" if prod_p and not prod else "variable",
        "unit_equal": all(v["unit"] == (prod or prod_p)[0]["unit"] for v in local) if (prod or prod_p) else "",
        "value_type_equal": all(v["value_type"] == (prod or prod_p)[0]["value_type"] for v in local) if (prod or prod_p) else "",
        "linked_by_id": "no",
        "status": "PENDING_CONTRACT_DECISION D24-11",
    })
write_csv("cross_workbook.csv", cross)

# ------------------------------------------------------------------
# 9. decisões contratuais pendentes (manifestos)
# ------------------------------------------------------------------
pending = {b: read_seed_file(b, "manifest")["pending_contract_decisions"] for b in sorted(BLOCKS)}
(OUT / "pending_contract_decisions.json").write_text(json.dumps(pending, ensure_ascii=False, indent=1), encoding="utf-8")

# ------------------------------------------------------------------
# 10. value_type / allowed_values / declared_result_states no domínio
# ------------------------------------------------------------------
domain = []
for d in var_defs.all():
    if d.value_type != "numerico" or d.allowed_values or d.declared_result_states or d.instances:
        domain.append({
            "variable_id": d.variable_definition_id, "name": d.variable_name, "frequency": d.frequency,
            "scope": f"{d.scope_type}/{d.scope_value}", "value_type": d.value_type,
            "allowed_values": "|".join(d.allowed_values or ()),
            "declared_result_states": "|".join(f"{s.state}->{s.literal}" for s in d.declared_result_states or ()),
            "instances": " ".join(i.scope_value for i in d.instances or ()),
        })
write_csv("domain_contract_fields.csv", domain)
summary["value_type_in_domain"] = collections.Counter(d.value_type for d in var_defs.all())

# ------------------------------------------------------------------
# 11. aliases / coerções / constantes históricas (varredura de código)
# ------------------------------------------------------------------
def grep(pattern, *paths):
    out = subprocess.run(["grep", "-rnE", pattern, *paths, "--include=*.py", "--include=*.json"],
                         cwd=REPO, capture_output=True, text=True).stdout.strip().splitlines()
    return [line for line in out if "/__pycache__/" not in line]

scans = {
    "value_type aliases (numeric|categorical) em app/, tools/, data/seed": grep(r"['\"](numeric|categorical)['\"]", "app", "tools", "data/seed"),
    "float()/int() de valor em tools/": grep(r"float\(|int\(entity|int\(row", "tools"),
    "constantes históricas de workbook (energy_v2|MaxHT_v5|yield_v4|production_v1) em tools/ e data/seed": grep(r"energy_v2|MaxHT_v5|yield_v4|production_v1", "tools", "data/seed"),
    "LAST_DATA_ROW/HEADER_ROW fixos em tools/": grep(r"LAST_DATA_ROW|HEADER_ROW|FIRST_DATA_ROW", "tools"),
    "value_type default NUMERIC em app/domain": grep(r"value_type: str = ", "app/domain"),
}
(OUT / "code_scans.json").write_text(json.dumps(scans, ensure_ascii=False, indent=1), encoding="utf-8")
summary["code_scan_hits"] = {k: len(v) for k, v in scans.items()}

(OUT / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=dict), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=1, default=dict))
