"""Stage 2.4 runtime probes (read-only; in-memory; seeds copied to tmp for mutation)."""
import copy, json, shutil, sys, tempfile, traceback
from pathlib import Path

REPO = Path("/home/user/forecast-model-audit-s06-v6-")
sys.path.insert(0, str(REPO))
U = Path("/root/.claude/uploads/afd1cfda-5340-5625-86ea-0588b02fbed0")
WB = {
    "area_41": (U / "27ac153c-descritivo_das_vari_veis_A41_v8.xlsx", "A41"),
    "energy": (U / "b7873d62-descritivo_das_vari_veis_energy_v5.xlsx", "energy"),
    "max_ht": (U / "28284f78-descritivo_das_vari_veis_MaxHT_v9.xlsx", "MaxHT"),
    "production": (U / "f72ff723-descritivo_das_vari_veis_production_v6.xlsx", "production"),
    "yield": (U / "e8cae0eb-descritivo_das_vari_veis_yield_v9.xlsx", "yield"),
}
HEADER = ['Type', 'name', 'description', 'unit', 'value', 'version', 'variable_type', 'frequency',
          'scope_type', 'scope_value', 'source_reference', 'status', 'value_type', 'allowed_values',
          'declared_result_states', 'expression', 'fonte', 'OBS']

from app.engine import reference_resolver as rr
from app.engine.calculation_context import CalculationContext
from app.engine.forecast_engine import ForecastEngine
from app.domain.equations.registry import EquationDefinitionRegistry
from app.domain.equations.models import EquationDefinition
from app.repositories.seed_loader import SeedLoader

OUT = {}


def rows_of(block):
    import openpyxl
    path, sheet = WB[block]
    sh = openpyxl.load_workbook(path, data_only=True)[sheet]
    h = 1 if block == "yield" else 2
    hdr = {sh.cell(h, c).value: c for c in range(1, sh.max_column + 1) if sh.cell(h, c).value}
    out = []
    for r in range(h + 1, sh.max_row + 1):
        d = {k: sh.cell(r, c).value for k, c in hdr.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in d.values()):
            continue
        d["row"] = r
        out.append(d)
    return out


def entities_of(block):
    ents = []
    for i, r in enumerate(rows_of(block)):
        ents.append({"entity_id": f"{block}:{r['row']}", "kind": "parameter" if r["Type"] == "parameter" else "variable", "name": r["name"],
                     "frequency": r["frequency"], "scope_type": r["scope_type"],
                     "scope_value": r["scope_value"], "row": r["row"], "expression": r["expression"]})
    return ents


# ---------------------------------------------------------------- P1 A019 builder path (8 gaps)
def p1():
    cases = [("production", 35, "lth_meta", None), ("production", 48, "lth_meta", None),
             ("production", 71, "pick_up", None), ("yield", 107, "tanque", None),
             ("yield", 107, "tanque_base", None), ("area_41", 52, "retirada_condensado_linha", "diário"),
             ("area_41", 53, "retirada_condensado_linha", "diário"), ("production", 60, "pick_up", "diário")]
    res = []
    idx_cache = {}
    for block, row, name, src_freq in cases:
        if block not in idx_cache:
            ents = entities_of(block)
            idx_cache[block] = (ents, rr.build_name_index(ents))
        ents, idx = idx_cache[block]
        cons = next(e for e in ents if e["row"] == row)
        freq = src_freq or cons["frequency"]
        pool = [(e["row"], e["frequency"], e["scope_type"], e["scope_value"]) for e in idx.get(name, [])]
        try:
            got = rr.resolve_reference(name, idx, freq, consumer_scope=(cons["scope_type"], cons["scope_value"]))
            outcome = f"RESOLVED -> {got}"
        except Exception as ex:
            outcome = f"{type(ex).__name__}: {str(ex)[:260]}"
        res.append(dict(block=block, row=row, consumer=cons["name"],
                        consumer_scope=f"{cons['scope_type']}/{cons['scope_value']}",
                        consumer_frequency=cons["frequency"], reference=name, lookup_frequency=freq,
                        candidates=pool, outcome=outcome))
    OUT["P1_A019_builder_path"] = res


# ---------------------------------------------------------------- P2 A019 runtime per-instance
def p2():
    loader = SeedLoader(REPO / "data" / "seed")
    var_defs, var_inst, par_defs, par_inst, eq_defs, eq_inst = loader.load_all_definitions_and_instances()
    meta = {i.scope_value: i.value for i in par_inst.all() if i.parameter_definition_id == "PARAM12001"}
    lth = {f"L{k}": 900.0 + 13 * k for k in range(1, 8)}
    ctx = CalculationContext()
    for ln, v in lth.items():
        ctx.set_variable_value("VAR12016", v, "linha", ln)
        ctx.set_parameter_value("PARAM12001", meta[ln], "linha", ln)
    sub = EquationDefinitionRegistry(); sub.add(eq_defs.get("EQ12004"))
    ForecastEngine().calculate_from_definition_registry(equation_definition_registry=sub, calculation_context=ctx)
    rows = []
    for ln in sorted(lth):
        got = ctx.get_variable_value("VAR12024", "linha", ln)
        exp = lth[ln] / meta[ln]
        other = [l for l in meta if meta[l] != meta[ln]][0]
        cross = lth[ln] / meta[other]
        rows.append(dict(line=ln, lth=lth[ln], lth_meta=meta[ln], oee=got, expected=exp,
                         ok=abs(got - exp) < 1e-12, cross_assoc_value=cross,
                         cross_assoc_detected=abs(got - cross) > 1e-9))
    # negative control: remove one scoped parameter -> must fail, not fall back to another line
    ctx2 = CalculationContext()
    for ln, v in lth.items():
        ctx2.set_variable_value("VAR12016", v, "linha", ln)
        if ln != "L3":
            ctx2.set_parameter_value("PARAM12001", meta[ln], "linha", ln)
    try:
        ForecastEngine().calculate_from_definition_registry(equation_definition_registry=sub, calculation_context=ctx2)
        neg = "NO ERROR (fallback!)"
    except Exception as ex:
        neg = f"{type(ex).__name__}: {str(ex)[:200]}"
    # parameter definition registry: get without scope
    try:
        par_defs.get("PARAM12001"); pd = "single"
    except Exception as ex:
        pd = f"{type(ex).__name__}: {str(ex)[:160]}"
    OUT["P2_A019_runtime_per_instance"] = dict(equation="EQ12004 oee = VAR12016 / PARAM12001 (linha/L1_L7)",
                                               rows=rows, negative_control_missing_L3=neg,
                                               param_defs_get_without_scope=pd,
                                               param12001_instances=len(meta))


# ---------------------------------------------------------------- P3 value_type / allowed_values / states via loader
def mutate_and_load(mut):
    tmp = Path(tempfile.mkdtemp())
    shutil.copytree(REPO / "data" / "seed", tmp / "seed")
    p = tmp / "seed" / "energy" / "variables.json"
    data = json.loads(p.read_text(encoding="utf-8"))
    mut(data)
    p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    try:
        reg = SeedLoader(tmp / "seed").load_variable_definitions()
        d = reg.get(data[0]["variable_id"])
        return f"LOADED value_type={getattr(d, 'value_type', None)!r}"
    except Exception as ex:
        return f"{type(ex).__name__}: {str(ex)[:220]}"
    finally:
        shutil.rmtree(tmp)


def p3():
    r = {}
    r["baseline_no_value_type"] = mutate_and_load(lambda d: None)
    r["value_type=numerico"] = mutate_and_load(lambda d: d[0].__setitem__("value_type", "numerico"))
    r["value_type=categorico"] = mutate_and_load(lambda d: d[0].__setitem__("value_type", "categorico"))
    r["value_type=numeric"] = mutate_and_load(lambda d: d[0].__setitem__("value_type", "numeric"))
    r["value_type=categorical"] = mutate_and_load(lambda d: d[0].__setitem__("value_type", "categorical"))
    r["allowed_values=[...]"] = mutate_and_load(lambda d: d[0].__setitem__("allowed_values", ["ABERTO", "FECHADO"]))
    r["declared_result_states=[...]"] = mutate_and_load(lambda d: d[0].__setitem__("declared_result_states", ["F"]))
    # seeds: how many carry value_type
    cnt = {}
    for b in ("energy", "max_ht", "production", "yield"):
        v = json.loads((REPO / "data/seed" / b / "variables.json").read_text(encoding="utf-8"))
        cnt[b] = dict(total=len(v), with_value_type=sum("value_type" in x for x in v),
                      with_allowed_values=sum("allowed_values" in x for x in v),
                      with_declared_result_states=sum("declared_result_states" in x for x in v))
    r["seed_field_presence"] = cnt
    # workbooks: categorical rows / allowed_values / states declared
    wb = {}
    for b in WB:
        rs = rows_of(b)
        wb[b] = dict(rows=len(rs), categorico=sum(x["value_type"] == "categorico" for x in rs),
                     allowed_values=sum(x["allowed_values"] not in (None, "") for x in rs),
                     declared_result_states=sum(x["declared_result_states"] not in (None, "") for x in rs))
    r["workbook_field_presence"] = wb
    OUT["P3_value_type_states_loader"] = r


# ---------------------------------------------------------------- P4 "F" and F-001 abort
def p4():
    r = {}
    # F consumed arithmetically -> ConditionalFailureError; F compared -> allowed
    eqs = [EquationDefinition("EQX1", "VAR99011", 1, "linha", "L1", "VAR99001 + 1", "probe", "PUBLISHED"),
           EquationDefinition("EQX2", "VAR99012", 1, "linha", "L1", "VAR99002 * 2", "probe", "PUBLISHED")]
    ctx = CalculationContext()
    ctx.declare_categorical_variables(["VAR99001"])
    ctx.set_variable_value("VAR99001", "F", "linha", "L1")
    ctx.set_variable_value("VAR99002", 21.0, "linha", "L1")
    sub = EquationDefinitionRegistry()
    for e in eqs: sub.add(e)
    try:
        res = ForecastEngine().calculate_from_definition_registry(equation_definition_registry=sub, calculation_context=ctx)
        r["F_arith"] = f"NO ERROR {res}"
    except Exception as ex:
        orig = getattr(ex, "original_error", None)
        r["F_arith"] = f"{type(ex).__name__} (original={type(orig).__name__ if orig else None})"
    try:
        v = ctx.get_variable_value("VAR99012", "linha", "L1")
        r["independent_VX2_after_failure"] = f"computed={v}"
    except Exception as ex:
        r["independent_VX2_after_failure"] = f"NOT COMPUTED ({type(ex).__name__})"
    # F preserved by comparison
    ctx3 = CalculationContext(); ctx3.declare_categorical_variables(["VAR99001", "VAR99013"])
    ctx3.set_variable_value("VAR99001", "F", "linha", "L1")
    sub3 = EquationDefinitionRegistry()
    sub3.add(EquationDefinition("EQY", "VAR99013", 1, "linha", "L1", '"F" if VAR99001 == "F" else "OK"', "probe", "PUBLISHED"))
    try:
        ForecastEngine().calculate_from_definition_registry(equation_definition_registry=sub3, calculation_context=ctx3)
        r["F_propagated_by_comparison"] = repr(ctx3.get_variable_value("VAR99013", "linha", "L1"))
    except Exception as ex:
        r["F_propagated_by_comparison"] = f"{type(ex).__name__}: {str(ex)[:200]}"
    # symbols absent (Stage 3)
    import subprocess
    for sym in ("EvaluationResult", "BLOCKED_BY_UPSTREAM_ERROR", "NO_APPLICABLE_RULE", "allowed_values", "declared_result_states"):
        out = subprocess.run(["grep", "-rln", sym, "app", "tools"], cwd=REPO, capture_output=True, text=True).stdout.split()
        r[f"symbol:{sym}"] = out
    OUT["P4_F_and_abort"] = r


# ---------------------------------------------------------------- P5 identity: same name across frequency/scope in current seeds
def p5():
    from collections import defaultdict
    r = {}
    for b in ("energy", "max_ht", "production", "yield"):
        v = json.loads((REPO / "data/seed" / b / "variables.json").read_text(encoding="utf-8"))
        by = defaultdict(list)
        for x in v:
            by[x["variable_name"]].append((x["variable_id"], x["frequency"], x["scope_type"], x["scope_value"]))
        dup_identity = {n: l for n, l in by.items() if len({t[1:] for t in l}) < len(l)}
        multi = {n: l for n, l in by.items() if len(l) > 1}
        r[b] = dict(names=len(by), names_with_multiple_definitions=len(multi),
                    identity_collisions=dup_identity,
                    per_line_split_names=sorted(n for n, l in multi.items()
                                                if sum(t[2] == "linha" and t[3] in {f"L{k}" for k in range(1, 8)} for t in l) >= 2))
    OUT["P5_identity_seeds"] = r


for f in (p1, p2, p3, p4, p5):
    try:
        f()
    except Exception:
        OUT[f.__name__ + "_CRASH"] = traceback.format_exc()

print(json.dumps(OUT, ensure_ascii=False, indent=1, default=str))
