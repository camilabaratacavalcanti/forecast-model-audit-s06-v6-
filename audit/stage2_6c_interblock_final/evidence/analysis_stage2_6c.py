"""
Etapa 2.6C — fechamento final do contrato interbloco (evidência independente).

    python audit/stage2_6c_interblock_final/evidence/analysis_stage2_6c.py [--no-write]

Independente do builder: NÃO importa nada de tools/ nem de app/. Lê os
cinco workbooks oficiais com openpyxl, declara a taxonomia D26-01 final
(com `max_ht`) por conta própria, reconstrói a classificação de cada
`fonte` e compara com a saída do builder gravada no repositório
(`data/seed/interblock_links.json`; a igualdade desse arquivo com o
build é verificada pela suíte de testes).

Também: diff célula a célula production v10 -> v11 (D26B-01), D26-02,
D26-03, max_ht (D26B-02), auditoria do livro de IDs contra os
manifestos de a126e02 (antes do livro) e eceffd4 (2.6B), e referências
órfãs nos seeds. Qualquer divergência reprova (código de saída != 0).

`--no-write` não grava arquivos (uso pela suíte de testes).
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
REPO = HERE.parents[2]
WB = REPO / "data" / "workbooks"
SEED = REPO / "data" / "seed"
LEDGER = REPO / "data" / "id_ledger"
WRITE = "--no-write" not in sys.argv[1:]
BASELINES = ("a126e02", "eceffd4")  # antes do livro de IDs; fim da 2.6B

OFFICIAL = {
    "area_41": ("descritivo_das_variáveis_A41_v9.xlsx", "A41",
                "36bbae135285ad9b3b88d21ef94e2bc9c92f2ef74736412b66bc62de4dcf3f50"),
    "energy": ("descritivo_das_variáveis_energy_v6.xlsx", "energy",
               "cfc46031fc2f3f375b87ebbba58676d23ad92f3ee43fb81b353fe5a8247e1df8"),
    "max_ht": ("descritivo_das_variáveis_MaxHT_v10.xlsx", "MaxHT",
               "3e40aab12ec0ec07918c5e358145b09f138663f635ca9e7fc0990d1b765b683b"),
    "production": ("descritivo_das_variáveis_production_v11.xlsx", "production",
                   "d946dfd522c8cbb301c0c63caa3c62263b74a849e043d7b2f4bcb0bdb262c922"),
    "yield": ("descritivo_das_variáveis_yield_v11.xlsx", "yield",
              "639e8b980f19bc20429cf9763157ac1828d9067bdbabea3bf3db0badbc2df931"),
}
PRODUCTION_V10 = ("descritivo_das_variáveis_production_v10.xlsx",
                  "302cf18b92ac22c690b3add9bdd1ded0617bfa60636809696d010b8858d72487")
PRODUCTION_V9 = ("descritivo_das_variáveis_production_v9.xlsx",
                 "3e3024e564ee6f18b82e9771957e61f4940becd1e80db5705143123377956a28")

# D26-01 final (D26B-02: max_ht), na ordem do proprietário. Registro histórico da 2.6C.
TAXONOMY_D26_01 = (
    "maintenance", "area_04_13", "forecast_volume", "acido", "yield", "energy",
    "meta_volume_cheio", "custo_budget", "production", "boilers",
    "controle_espaco_vazio_meta", "custo_forecast_bdgt", "max_ht", "volume",
    "lime_dia", "custo_forecast_real", "alumina", "soda",
    "floculante_hidrato_2026", "budget", "temperature_lp", "fator_residuo",
    "floculante_lama_dia", "forecast", "area_41", "vazao_condensado",
    "premissas_ppt_mensal", "shared",
)
assert len(TAXONOMY_D26_01) == len(set(TAXONOMY_D26_01)) == 28

# Historical note: a taxonomia canônica do repositório foi depois normalizada para
# 29 blocos (D-TAX-01, posterior à Stage 3), na ordem das faixas de ID, incluindo
# budget_vs_forecast; a faixa 30000-30999 chama-se thickener_flocculant. Como esta
# análise independente roda hoje como teste (test_17), o oráculo atual é a lista
# D-TAX-01 do proprietário; a D26-01 acima fica só como registro. A evidência
# versionada desta etapa (validation_results.json) não é regenerada.
TAXONOMY = (
    "maintenance", "yield", "production", "max_ht", "alumina", "temperature_lp",
    "area_41", "area_04_13", "energy", "boilers", "volume", "soda",
    "residue_factor", "condensate_flow", "forecast_volume", "full_volume_target",
    "empty_space_target_control", "lime", "hydrated_flocculant",
    "sludge_flocculant", "thickener_flocculant", "acid", "budget_cost",
    "budget_forecast_cost", "actual_forecast_cost", "budget", "forecast",
    "budget_vs_forecast", "shared",
)
assert len(TAXONOMY) == len(set(TAXONOMY)) == 29
LINES = [f"L{i}" for i in range(1, 8)]
CONTRACT_DIMENSIONS = ("kind", "unit", "value_type", "allowed_values", "declared_result_states")
CONTRACT_COLUMNS = (
    "Type", "name", "description", "unit", "value", "version", "variable_type", "frequency",
    "scope_type", "scope_value", "source_reference", "status", "value_type", "allowed_values",
    "declared_result_states", "expression", "fonte",
)
failures: list[str] = []


def check(condition, message):
    if not condition:
        failures.append(message)
    return condition


def write_csv(path, rows):
    if not WRITE:
        return
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def write_text(path, text):
    if WRITE:
        path.write_text(text, encoding="utf-8")


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def git_show(ref, path):
    return subprocess.check_output(["git", "show", f"{ref}:{path}"], cwd=REPO).decode("utf-8")


# ------------------------------------------------------------------
# 0. integridade
# ------------------------------------------------------------------
sha_lines = []
for block, (name, _sheet, expected) in sorted(OFFICIAL.items()):
    digest = sha(WB / name)
    check(digest == expected, f"SHA {block}: {digest}")
    sha_lines.append(f"{digest}  data/workbooks/{name}  (oficial: {block})")
for name, expected, note in (
    (*PRODUCTION_V10, "referência do diff D26B-01; oficial na 2.6B"),
    (*PRODUCTION_V9, "histórico; oficial na 2.6"),
):
    digest = sha(WB / name)
    check(digest == expected, f"SHA {name}")
    sha_lines.append(f"{digest}  data/workbooks/{name}  ({note})")
write_text(HERE / "workbook_sha256.txt", "\n".join(sha_lines) + "\n")


# ------------------------------------------------------------------
# 1. diff production v10 -> v11 (D26B-01)
# ------------------------------------------------------------------
def load_sheet(path, sheet):
    return openpyxl.load_workbook(path)[sheet]


v9 = load_sheet(WB / PRODUCTION_V9[0], "production")
v10 = load_sheet(WB / PRODUCTION_V10[0], "production")
v11 = load_sheet(WB / OFFICIAL["production"][0], "production")
header = [c.value for c in v10[2]]
obs_col = header.index("OBS") + 1
check(
    (v10.max_row, v10.max_column) == (v11.max_row, v11.max_column),
    "dimensões da aba production mudaram v10 -> v11",
)
check(
    openpyxl.load_workbook(WB / PRODUCTION_V10[0]).sheetnames
    == openpyxl.load_workbook(WB / OFFICIAL["production"][0]).sheetnames,
    "abas mudaram v10 -> v11",
)

# Valor esperado em OBS (D26B-01): o antigo `value` de lth_meta (v9) na
# mesma linha; vazio em qualquer outra linha que não tinha OBS textual.
lth_meta_rows = [r for r in range(3, v9.max_row + 1) if v9.cell(r, 2).value == "lth_meta"]
expected_obs = {r: v9.cell(r, 5).value for r in lth_meta_rows}

diff_rows = []
for r in range(1, max(v10.max_row, v11.max_row) + 1):
    for c in range(1, max(v10.max_column, v11.max_column) + 1):
        a, b = v10.cell(r, c).value, v11.cell(r, c).value
        if a == b:
            continue
        column = header[c - 1] if c <= len(header) else f"col{c}"
        if column == "OBS" and b == expected_obs.get(r):
            classification = "D26B-01"
        else:
            classification = "UNEXPECTED"
        diff_rows.append({
            "cell": v11.cell(r, c).coordinate, "row": r, "column": column,
            "name": v11.cell(r, 2).value, "frequency": v11.cell(r, 8).value,
            "scope_value": v11.cell(r, 10).value, "v10": a, "v11": b,
            "expected_v11": expected_obs.get(r), "classification": classification,
        })
write_csv(HERE / "workbook_diff_v10_v11.csv", diff_rows)
diff_summary = Counter(d["classification"] for d in diff_rows)
check(diff_summary.get("UNEXPECTED", 0) == 0, f"diff v10->v11 inesperado: {diff_summary}")
check({d["column"] for d in diff_rows} <= {"OBS"}, "diff v10->v11 fora de OBS")

obs_alignment = []
for r in sorted(set(lth_meta_rows) | {min(lth_meta_rows) - 1, max(lth_meta_rows) + 1}):
    row = {
        "row": r, "name": v11.cell(r, 2).value, "frequency": v11.cell(r, 8).value,
        "scope_value": v11.cell(r, 10).value, "v9_value": v9.cell(r, 5).value,
        "v10_OBS": v10.cell(r, obs_col).value, "v11_OBS": v11.cell(r, obs_col).value,
    }
    row["v11_aligned"] = row["v11_OBS"] == (expected_obs.get(r))
    check(row["v11_aligned"], f"OBS v11 desalinhado na linha {r}")
    obs_alignment.append(row)

# Colunas de contrato idênticas v10 x v11 em todas as linhas.
contract_changes = [
    d for d in diff_rows if d["column"] in CONTRACT_COLUMNS
]
check(not contract_changes, f"colunas de contrato mudaram: {contract_changes}")


# ------------------------------------------------------------------
# 2. leitura crua e definições
# ------------------------------------------------------------------
def raw_rows_from(ws):
    header_row, cols = None, None
    for r in range(1, 11):
        values = [c.value for c in ws[r]]
        if "name" in values and "fonte" in values:
            header_row, cols = r, values
            break
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        values = dict(zip(cols, [c.value for c in ws[r]]))
        if values.get("name") is None:
            continue
        values["_row"] = r
        values["_kind"] = "parameter" if values["Type"] == "parameter" else "variable"
        rows.append(values)
    return rows


def raw_rows(block):
    name, sheet, _ = OFFICIAL[block]
    return raw_rows_from(load_sheet(WB / name, sheet))


def instances(scope_type, scope_value):
    if scope_type == "linha" and "_" in str(scope_value):
        a, b = str(scope_value).split("_")
        return {(scope_type, l) for l in LINES[LINES.index(a): LINES.index(b) + 1]}
    return {(scope_type, scope_value)}


def definitions_of(block, rows):
    groups = defaultdict(list)
    for r in rows:
        per_line = r["scope_type"] == "linha" and r["scope_value"] in LINES
        key = (r["_kind"], r["name"], r["frequency"], r["scope_type"], "*" if per_line else r["scope_value"])
        groups[key].append(r)
    out = []
    for key, grp in groups.items():
        inst = set()
        for r in grp:
            inst |= instances(r["scope_type"], r["scope_value"])
        out.append({"key": key, "rows": grp, "instances": inst, "first": grp[0], "block": block})
    return out


RAW = {block: raw_rows(block) for block in OFFICIAL}
DEFS = {block: definitions_of(block, RAW[block]) for block in OFFICIAL}


def counts(rows):
    defs = definitions_of("x", rows)
    return {
        "rows": len(rows),
        "variables": sum(1 for d in defs if d["key"][0] == "variable"),
        "parameters": sum(1 for d in defs if d["key"][0] == "parameter"),
        "parameter_rows": sum(1 for r in rows if r["_kind"] == "parameter"),
        "expression_rows": sum(1 for r in rows if r["expression"] is not None),
        "fonte_rows": sum(1 for r in rows if r["fonte"] is not None),
    }


production_counts = {
    "v10": counts(raw_rows_from(v10)),
    "v11": counts(raw_rows_from(v11)),
}
check(production_counts["v10"] == production_counts["v11"], "contagens v10 != v11")


# ------------------------------------------------------------------
# 3. classificação independente
# ------------------------------------------------------------------
def classify(d):
    row = d["first"]
    source = row["fonte"]
    findings = []
    if row["_kind"] != "variable" or any(r["expression"] is not None for r in d["rows"]):
        findings.append(("CONTRACT_MISMATCH", "consumer_context"))
    if source not in TAXONOMY:
        return None, findings + [("SOURCE_BLOCK_UNKNOWN", "source_block")]
    if source not in OFFICIAL:
        return None, findings + [("SOURCE_BLOCK_NOT_LOADED", "source_block")]
    wanted = d["instances"]
    same_axes = [
        p for p in DEFS[source]
        if (p["key"][1], p["key"][2], p["key"][3]) == (row["name"], row["frequency"], row["scope_type"])
    ]
    covering = [p for p in same_axes if wanted <= p["instances"]]
    if len(covering) > 1:
        return None, findings + [("AMBIGUOUS", "identity")]
    if not covering:
        if not [p for p in DEFS[source] if p["key"][1] == row["name"]]:
            return None, findings + [("SOURCE_NOT_FOUND", "name")]
        if same_axes:
            union = set().union(*(p["instances"] for p in same_axes))
            if wanted <= union:
                return None, findings + [("AMBIGUOUS", "identity")]
            return None, findings + [("INSTANCE_MISMATCH", "instances")]
        return None, findings + [("SOURCE_NOT_FOUND", "frequency/scope")]
    producer = covering[0]
    p = producer["first"]
    for dimension in CONTRACT_DIMENSIONS:
        mine = row["_kind"] if dimension == "kind" else row[dimension]
        theirs = p["_kind"] if dimension == "kind" else p[dimension]
        if mine != theirs:
            findings.append(("CONTRACT_MISMATCH", dimension))
    return producer, findings


links = []
for block in sorted(OFFICIAL):
    for d in DEFS[block]:
        if d["first"]["fonte"] is None:
            continue
        check(len({r["fonte"] for r in d["rows"]}) == 1, f"fonte divergente em {block} {d['key']}")
        producer, findings = classify(d)
        links.append({"block": block, "def": d, "producer": producer, "findings": findings})

# Ciclos: vínculos + dependências intrabloco sobre-aproximadas.
node = lambda d: (d["block"], d["key"])
graph = defaultdict(set)
for block, defs in DEFS.items():
    by_name = defaultdict(list)
    for d in defs:
        by_name[d["key"][1]].append(d)
    for d in defs:
        for r in d["rows"]:
            for token in set(re.findall(r"[A-Za-z_À-ÿ][\wÀ-ÿ]*", r["expression"] or "")):
                for target in by_name.get(token, []):
                    if target is not d:
                        graph[node(d)].add(node(target))
for l in links:
    if l["producer"]:
        graph[node(l["def"])].add(node(l["producer"]))

color, cycles = {}, []


def dfs(start):
    stack = [(start, iter(sorted(graph.get(start, ()), key=str)))]
    path = [start]
    color[start] = 1
    while stack:
        n, it = stack[-1]
        nxt = next(it, None)
        if nxt is None:
            stack.pop(); path.pop(); color[n] = 2
            continue
        if color.get(nxt) == 1:
            cycles.append(path[path.index(nxt):] + [nxt])
        elif not color.get(nxt):
            color[nxt] = 1; path.append(nxt)
            stack.append((nxt, iter(sorted(graph.get(nxt, ()), key=str))))


for n in sorted(graph, key=str):
    if not color.get(n):
        dfs(n)
interblock_cycles = [c for c in cycles if any(a[0] != b[0] for a, b in zip(c, c[1:]))]
for l in links:
    if l["producer"] and any(node(l["def"]) in c for c in interblock_cycles):
        l["findings"].append(("CYCLE", "cycle"))
for l in links:
    l["status"] = l["findings"][0][0] if l["findings"] else "VALID"

by_consumer = {node(l["def"]): l for l in links}
for l in links:
    chain, status, cur, seen = [f"{l['block']}.{l['def']['key'][1]}"], "VALID", l, set()
    while True:
        if node(cur["def"]) in seen:
            status = "CYCLE"; break
        seen.add(node(cur["def"]))
        src = cur["def"]["first"]["fonte"]
        if cur["status"] != "VALID":
            status = ("PENDING_AT:" if cur["status"] == "SOURCE_BLOCK_NOT_LOADED" else "BROKEN_AT:") + \
                f"{cur['block']}.{cur['def']['key'][1]}"
            chain.append(f"{src}.{cur['producer']['key'][1] if cur['producer'] else '?'}")
            break
        chain.append(f"{src}.{cur['producer']['key'][1]}")
        nxt = by_consumer.get(node(cur["producer"]))
        if nxt is None:
            break
        cur = nxt
    l["chain"], l["chain_status"] = " → ".join(chain), status


# ------------------------------------------------------------------
# 4. comparação com o builder (seed gravado) e IDs de definição
# ------------------------------------------------------------------
seed = json.loads((SEED / "interblock_links.json").read_text(encoding="utf-8"))
builder = {}
for rec in seed["links"]:
    builder[(rec["consumer_block"], tuple(rec["consumer_rows"]))] = ("VALID", rec["source_block"])
for group in ("pending", "rejected"):
    for rec in seed[group]:
        builder[(rec["consumer_block"], tuple(rec["consumer_rows"]))] = (rec["validation_status"], rec["source_block"])
independent = {
    (l["block"], tuple(r["_row"] for r in l["def"]["rows"])): (l["status"], l["def"]["first"]["fonte"])
    for l in links
}
divergences = [
    {"key": list(k), "independent": independent.get(k), "builder": builder.get(k)}
    for k in sorted(set(independent) | set(builder))
    if independent.get(k) != builder.get(k)
]
check(not divergences, f"divergências independente x builder: {divergences}")
check(tuple(seed["taxonomy"]["official_blocks"]) == TAXONOMY, "taxonomia do seed != D-TAX-01 (29 blocos)")

manifests = {
    block: json.loads((SEED / block / "manifest.json").read_text(encoding="utf-8")) for block in OFFICIAL
}


def manifest_entity(block, rows):
    return next(
        (e for e in manifests[block]["entities"] if e["rows"] == list(rows)), None
    )


# Os IDs do seed (consumidor/produtor) correspondem às linhas do workbook.
for rec in seed["links"]:
    consumer = manifest_entity(rec["consumer_block"], rec["consumer_rows"])
    check(consumer and consumer["entity_id"] == rec["consumer_definition"],
          f"consumer_definition órfão {rec['consumer_definition']}")
    producer_ids = {e["entity_id"] for e in manifests[rec["source_block"]]["entities"]}
    check(rec["source_definition"] in producer_ids, f"source_definition órfão {rec['source_definition']}")
for rec in seed["pending"]:
    check(rec["source_definition"] is None and rec["source_block"] not in OFFICIAL,
          f"produtor fictício em pending: {rec}")


# ------------------------------------------------------------------
# 5. D26-02, D26-03, max_ht
# ------------------------------------------------------------------
def link_of(block, name):
    return next(l for l in links if l["block"] == block and l["def"]["key"][1] == name)


lth_meta = link_of("energy", "lth_meta")
p = lth_meta["producer"]
d26_02 = {
    "status": lth_meta["status"],
    "consumer": {k: lth_meta["def"]["first"][k] for k in ("Type", "variable_type", "frequency", "scope_type", "scope_value", "unit", "value_type")},
    "producer": {k: p["first"][k] for k in ("Type", "variable_type", "frequency", "scope_type", "unit", "value_type")} if p else None,
    "producer_instances": sorted(v for _t, v in p["instances"]) if p else None,
    "production_lth_meta_parameter_rows": [r["_row"] for r in RAW["production"] if r["name"] == "lth_meta" and r["_kind"] == "parameter"],
}
check(d26_02["status"] == "VALID", "D26-02 não é VALID")
check(p and p["key"][0] == "variable" and p["first"]["variable_type"] == "entrada_externa", "D26-02 produtor")
check(lth_meta["def"]["first"]["variable_type"] == "entrada", "D26-02 consumidor")
check(not d26_02["production_lth_meta_parameter_rows"], "lth_meta parameter ainda existe")

d26_03 = {}
for name in ("fator_mpsa", "fator_mrn", "fator_mpsa_kg_t", "fator_mrn_kg_t"):
    [row] = [r for r in RAW["production"] if r["name"] == name]
    d26_03[name] = {k: row[k] for k in ("Type", "variable_type", "unit", "frequency", "scope_type", "scope_value", "expression", "fonte")}
for name in ("fator_mpsa", "fator_mrn"):
    check(d26_03[name]["fonte"] is None and d26_03[name]["expression"] == f"{name}_kg_t / 1000"
          and d26_03[name]["variable_type"] == "calculado", f"D26-03 {name}")
for name in ("fator_mpsa_kg_t", "fator_mrn_kg_t"):
    check(d26_03[name]["fonte"] == "forecast" and d26_03[name]["variable_type"] == "entrada"
          and d26_03[name]["expression"] is None, f"D26-03 {name}")
    check(link_of("production", name)["status"] == "SOURCE_BLOCK_NOT_LOADED", f"D26-03 {name} status")
check(not [l for l in links if any(k == "CONTRACT_MISMATCH" and dim == "consumer_context" for k, dim in l["findings"])],
      "consumidor com expressão + fonte")

fontes = sorted({l["def"]["first"]["fonte"] for l in links})
max_ht = {
    "max_ht_in_taxonomy": "max_ht" in TAXONOMY,
    "mx_ht_in_taxonomy": "mx_ht" in TAXONOMY,
    "loaded_block_ids": sorted(OFFICIAL),
    "loaded_blocks_outside_taxonomy": sorted(b for b in OFFICIAL if b not in TAXONOMY),
    "seed_loaded_blocks": seed["taxonomy"]["loaded_blocks"],
    "manifest_block": manifests["max_ht"]["block"],
    "fonte_using_max_ht": [l["block"] for l in links if l["def"]["first"]["fonte"] == "max_ht"],
    "fonte_using_mx_ht": [l["block"] for l in links if l["def"]["first"]["fonte"] == "mx_ht"],
}
check(max_ht["max_ht_in_taxonomy"] and not max_ht["mx_ht_in_taxonomy"], "taxonomia max_ht/mx_ht")
check(not max_ht["loaded_blocks_outside_taxonomy"], "bloco carregado fora da taxonomia")
check(all(b["in_taxonomy"] for b in seed["taxonomy"]["loaded_blocks"]), "seed: bloco fora da taxonomia")
check(max_ht["manifest_block"] == "max_ht", "manifesto max_ht")
active_mx_ht = []
for path in list((REPO / "data").rglob("*.json")) + list((REPO / "tools").rglob("*.py")):
    text = path.read_text(encoding="utf-8")
    if "mx_ht" in text:
        active_mx_ht.append({"file": str(path.relative_to(REPO)), "count": text.count("mx_ht")})
max_ht["active_mx_ht_occurrences"] = active_mx_ht
# Única ocorrência aceita: o docstring de taxonomy.py que declara mx_ht não oficial.
check(
    [o["file"] for o in active_mx_ht] in ([], ["tools/workbook_seed/taxonomy.py"])
    and all(o["count"] == 1 for o in active_mx_ht),
    f"mx_ht em artefato ativo: {active_mx_ht}",
)


# ------------------------------------------------------------------
# 6. livro de IDs
# ------------------------------------------------------------------
def ids_of(text):
    return {
        (e["kind"], e["name"], e["frequency"], e["scope_type"], e["scope_value"]): e["entity_id"]
        for e in json.loads(text)["entities"]
    }


ledger_rows = []
id_totals = Counter()
ledgers = {}
for block in OFFICIAL:
    current = ids_of((SEED / block / "manifest.json").read_text(encoding="utf-8"))
    baseline = {ref: ids_of(git_show(ref, f"data/seed/{block}/manifest.json")) for ref in BASELINES}
    ledger = json.loads((LEDGER / f"{block}.json").read_text(encoding="utf-8"))
    ledgers[block] = ledger
    entries = {
        (e["kind"], e["name"], e["frequency"], e["scope_type"], e["scope_value"]): e["entity_id"]
        for e in ledger["entries"]
    }
    retired = {r["entity_id"]: r for r in ledger["retired"]}
    check(entries == current, f"livro {block} != manifesto")
    check(len(set(entries.values())) == len(entries), f"livro {block}: ID repetido")
    check(not set(retired) & set(entries.values()), f"livro {block}: aposentado reutilizado")
    all_ids = set(current.values())
    for key in sorted(set(current) | set(baseline["a126e02"]) | set(baseline["eceffd4"]), key=str):
        before = baseline["a126e02"].get(key)
        mid = baseline["eceffd4"].get(key)
        now = current.get(key)
        if before and now:
            status = "PRESERVED" if before == now else "RENUMBERED"
        elif now and not before:
            status = "NEW"
        else:
            status = "RETIRED"
        if status == "RETIRED":
            check(before in retired, f"{block}: {before} removido sem registro de aposentadoria")
            check(before not in all_ids, f"{block}: {before} aposentado reutilizado")
        check(status != "RENUMBERED", f"{block}: {key} renumerado {before} -> {now}")
        if mid is not None and now is not None:
            check(mid == now, f"{block}: {key} mudou desde eceffd4")
        id_totals[status] += 1
        ledger_rows.append({
            "block": block, "kind": key[0], "name": key[1], "frequency": key[2],
            "scope_type": key[3], "scope_value": key[4],
            "id_a126e02": before or "", "id_eceffd4": mid or "", "id_current": now or "",
            "ledger": entries.get(key) or (before if before in retired else ""),
            "status": status,
            "retired_in": retired[before]["retired_in"] if before in retired else "",
        })
write_csv(ROOT / "id_ledger_audit.csv", ledger_rows)
new_ids = [(r["block"], r["name"], r["id_current"]) for r in ledger_rows if r["status"] == "NEW"]
retired_ids = [(r["block"], r["name"], r["id_a126e02"]) for r in ledger_rows if r["status"] == "RETIRED"]
check(new_ids == [("production", "lth_meta", "VAR12066")], f"IDs novos inesperados: {new_ids}")
check(retired_ids == [("production", "lth_meta", "PARAM12003")], f"IDs aposentados inesperados: {retired_ids}")

# Referências órfãs nos seeds: toda referência de equação/agregação existe.
orphans = []
for block in OFFICIAL:
    ids = {e["entity_id"] for e in manifests[block]["entities"]}
    for eq in json.loads((SEED / block / "equations.json").read_text(encoding="utf-8")):
        for ref in re.findall(r"\b(VAR\d+|PARAM\d+)", eq["expression"]) + [eq["target_variable_id"]]:
            if ref not in ids:
                orphans.append((block, eq["equation_id"], ref))
    for rule in json.loads((SEED / block / "aggregation_rules.json").read_text(encoding="utf-8")):
        for ref in (rule["source_variable_id"], rule["target_variable_id"]):
            if ref not in ids:
                orphans.append((block, rule["aggregation_rule_id"], ref))
check(not orphans, f"referências órfãs: {orphans}")
param12003 = [
    str(p.relative_to(REPO)) for p in SEED.rglob("*.json")
    if "PARAM12003" in p.read_text(encoding="utf-8")
]
check(not param12003, f"PARAM12003 ainda referenciado: {param12003}")
lth_meta_refs = []
for eq in json.loads((SEED / "production" / "equations.json").read_text(encoding="utf-8")):
    if "VAR12066" in eq["expression"]:
        lth_meta_refs.append(eq["equation_id"])
old_refs = [
    eq["equation_id"] for eq in json.loads(git_show("a126e02", "data/seed/production/equations.json"))
    if "PARAM12003" in eq["expression"]
]
check(lth_meta_refs == old_refs, f"equações de lth_meta: antes {old_refs}, agora {lth_meta_refs}")


# ------------------------------------------------------------------
# 7. CSVs de vínculos, contrato e grafo
# ------------------------------------------------------------------
link_rows, contract_rows = [], []
for l in links:
    d, p = l["def"], l["producer"]
    c = d["first"]
    pf = p["first"] if p else None
    key = (l["block"], tuple(x["_row"] for x in d["rows"]))
    producer_id = next(
        (rec["source_definition"] for rec in seed["links"]
         if (rec["consumer_block"], tuple(rec["consumer_rows"])) == key), "",
    )
    for r in d["rows"]:
        link_rows.append({
            "consumer_block": l["block"],
            "consumer_row": r["_row"],
            "consumer_name": c["name"],
            "consumer_unit": c["unit"],
            "consumer_frequency": c["frequency"],
            "consumer_scope": f"{r['scope_type']}/{r['scope_value']}",
            "fonte": c["fonte"],
            "source_in_taxonomy": c["fonte"] in TAXONOMY,
            "source_block_loaded": c["fonte"] in OFFICIAL,
            "classification": l["status"],
            "state": "RESOLVED" if l["status"] == "VALID" else (
                "PENDING_LOAD" if l["status"] == "SOURCE_BLOCK_NOT_LOADED" else "REJECTED"),
            "producer": f"{c['fonte']}.{pf['name']} [{pf['frequency']} {p['key'][3]}]" if pf else "",
            "producer_definition": producer_id,
            "producer_kind": pf["_kind"] if pf else "",
            "producer_unit": pf["unit"] if pf else "",
            "consumer_variable_type": c["variable_type"],
            "producer_variable_type": pf["variable_type"] if pf else "",
            "errors_or_pending": ";".join(f"{k}:{dim}" for k, dim in l["findings"]),
            "chain": l["chain"],
            "chain_status": l["chain_status"],
            "builder_classification": builder.get(key, ("?",))[0],
            "agreement": "AGREE" if independent[key] == builder.get(key) else "DIVERGE",
        })
    if not p:
        contract_rows.append({
            "consumer_block": l["block"], "consumer_rows": ",".join(str(x["_row"]) for x in d["rows"]),
            "consumer_name": c["name"], "fonte": c["fonte"], "dimension": "source_block",
            "consumer_value": c["fonte"],
            "producer_value": "fora da taxonomia" if c["fonte"] not in TAXONOMY else "workbook não carregado",
            "result": l["status"], "is_contract_dimension": True,
        })
        continue
    values = {
        "kind": (c["_kind"], pf["_kind"]),
        "name": (c["name"], pf["name"]),
        "unit": (c["unit"], pf["unit"]),
        "value_type": (c["value_type"], pf["value_type"]),
        "allowed_values": (c["allowed_values"], pf["allowed_values"]),
        "declared_result_states": (c["declared_result_states"], pf["declared_result_states"]),
        "frequency": (c["frequency"], pf["frequency"]),
        "scope_type": (c["scope_type"], pf["scope_type"]),
        "instances": (sorted(v for _t, v in d["instances"]), sorted(v for _t, v in p["instances"])),
        "variable_type": (c["variable_type"], pf["variable_type"]),
    }
    for dimension, (mine, theirs) in values.items():
        ok = set(mine) <= set(theirs) if dimension == "instances" else mine == theirs
        contract_rows.append({
            "consumer_block": l["block"], "consumer_rows": ",".join(str(x["_row"]) for x in d["rows"]),
            "consumer_name": c["name"], "fonte": c["fonte"], "dimension": dimension,
            "consumer_value": mine, "producer_value": theirs,
            "result": ("OK" if ok else "DIFFERENT (não é dimensão do contrato)")
            if dimension == "variable_type" else ("OK" if ok else "MISMATCH"),
            "is_contract_dimension": dimension != "variable_type",
        })
write_csv(ROOT / "interblock_links.csv", link_rows)
write_csv(ROOT / "contract_validation.csv", contract_rows)

graph_rows = []
block_edges = Counter()
for l in links:
    block_edges[(l["block"], l["def"]["first"]["fonte"], l["status"])] += 1
    graph_rows.append({
        "level": "definition",
        "consumer": f"{l['block']}.{l['def']['key'][1]}[{l['def']['key'][2]} {l['def']['key'][3]}]",
        "producer": (f"{l['def']['first']['fonte']}.{l['producer']['key'][1]}"
                     f"[{l['producer']['key'][2]} {l['producer']['key'][3]}]")
        if l["producer"] else f"{l['def']['first']['fonte']}.?",
        "edges": 1, "status": l["status"], "note": f"{l['chain']} ({l['chain_status']})",
    })
pairs = {(a, b) for (a, b, _s) in block_edges}
for (a, b, status), n in sorted(block_edges.items()):
    graph_rows.append({
        "level": "block", "consumer": a, "producer": b, "edges": n, "status": status,
        "note": "dependência mútua entre blocos; sem ciclo no nível de definição" if (b, a) in pairs else "",
    })
write_csv(ROOT / "interblock_graph.csv", graph_rows)


# ------------------------------------------------------------------
# 8. resumo
# ------------------------------------------------------------------
summary = {
    "baselines": list(BASELINES),
    "workbooks": {b: {"file": v[0], "sha256": v[2]} for b, v in sorted(OFFICIAL.items())},
    "production_v10_v11_diff": {"cells": len(diff_rows), "by_classification": dict(diff_summary),
                                "cells_list": [d["cell"] for d in diff_rows]},
    "production_counts": production_counts,
    "lth_meta_obs_alignment": obs_alignment,
    "taxonomy": list(TAXONOMY),
    "fonte_names": {f: {"in_taxonomy": f in TAXONOMY, "loaded": f in OFFICIAL} for f in fontes},
    "max_ht": max_ht,
    "d26_02": d26_02,
    "d26_03": d26_03,
    "fonte_rows": len(link_rows),
    "declared_links_definitions": len(links),
    "links_by_class_definitions": dict(sorted(Counter(l["status"] for l in links).items())),
    "links_by_class_rows": dict(sorted(Counter(r["classification"] for r in link_rows).items())),
    "cycles_overapproximated": [[str(n) for n in c] for c in interblock_cycles],
    "id_ledger": {
        "totals": dict(id_totals),
        "new": new_ids,
        "retired": retired_ids,
        "lth_meta_equations": lth_meta_refs,
        "orphans": orphans,
    },
    "independent_vs_builder_divergences": divergences,
    "failures": failures,
}
write_text(HERE / "validation_results.json", json.dumps(summary, ensure_ascii=False, indent=2, default=list) + "\n")
print(json.dumps({k: summary[k] for k in (
    "production_v10_v11_diff", "fonte_rows", "declared_links_definitions",
    "links_by_class_definitions", "links_by_class_rows", "cycles_overapproximated",
    "id_ledger", "independent_vs_builder_divergences", "failures",
)}, ensure_ascii=False, indent=1, default=list))
sys.exit(1 if failures else 0)
