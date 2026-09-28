import sys,re,json,collections,csv,ast
sys.path.insert(0,"/home/user/forecast-model-audit-s06-v6-")
import openpyxl
from wbp import WB
from app.domain.values import VALUE_TYPES
from app.validation import variable_seed_validator as V, parameter_seed_validator as P
from app.engine.scope_resolver import ScopeResolver
SR=ScopeResolver()
TAXONOMY={"NO_APPLICABLE_RULE","INVALID_INPUT","VALIDATION_FAILED"}
HIDDEN=re.compile(r"[ ​‌‍﻿\t\r]")
NAME_RE=re.compile(r"^[A-Za-zÀ-ÿ_][\wÀ-ÿ]*$")
O=json.load(open("probe2_out.json"))
findings=[]; counts=collections.Counter(); param_rows=[]; status_rows=[]; meta=collections.defaultdict(collections.Counter)
def add(b,sheet,cell,row,col,val,cls,rule,exp=""):
    findings.append(dict(workbook=b,sheet=sheet,cell=cell,row=row,column=col,value=repr(val),physical_type=type(val).__name__,expected=exp,classification=cls,rule=rule))
for b,p in WB.items():
    ws=openpyxl.load_workbook(p).worksheets[0]; h=1 if b=="yield" else 2
    hdr={ws.cell(h,c).value:c for c in range(1,ws.max_column+1) if ws.cell(h,c).value}
    # every non-empty cell of the sheet, any row/column
    for row in ws.iter_rows():
        for c in row:
            v=c.value
            if isinstance(v,str):
                counts[(b,"str_cells")]+=1
                if HIDDEN.search(v): add(b,ws.title,c.coordinate,c.row,c.column_letter,v,"BLOCKER" if c.row>h and ws.cell(h,c.column).value not in("description","OBS","expression") else "INFORMATIONAL","caractere oculto/não imprimível")
    for r in range(h+1,ws.max_row+1):
        cell=lambda k: ws.cell(r,hdr[k]) if k in hdr else None
        vals={k:ws.cell(r,c).value for k,c in hdr.items()}
        if all(v is None or (isinstance(v,str) and not v.strip()) for v in vals.values()):
            if any(ws.cell(r,c).value not in (None,"") for c in range(1,ws.max_column+1)): add(b,ws.title,f"row {r}",r,"*","conteúdo fora das colunas de cabeçalho","BLOCKER","célula fora do contrato")
            continue
        counts[(b,"rows")]+=1
        t=vals.get("Type"); name=vals.get("name")
        kind="parameter" if t=="parameter" else "variable"
        # string hygiene on contract fields
        for k in ("Type","name","unit","variable_type","frequency","scope_type","scope_value","status","value_type","declared_result_states"):
            v=vals.get(k)
            if isinstance(v,str) and v!=v.strip(): add(b,ws.title,cell(k).coordinate,r,k,v,"BLOCKER","espaço antes/depois em campo de contrato",v.strip())
            if v is not None and k!="version" and not isinstance(v,str): add(b,ws.title,cell(k).coordinate,r,k,v,"BLOCKER","campo de contrato deve ser texto")
        if t not in ("variable","variable / equation","parameter"): add(b,ws.title,cell("Type").coordinate,r,"Type",t,"BLOCKER","Type fora do vocabulário")
        if not (isinstance(name,str) and NAME_RE.match(name)): add(b,ws.title,cell("name").coordinate,r,"name",name,"BLOCKER","name não é identificador válido")
        # value / version physical types
        val=vals.get("value"); ver=vals.get("version")
        if kind=="parameter":
            ok=isinstance(val,(int,float)) and not isinstance(val,bool)
            param_rows.append(dict(workbook=b,row=r,cell=cell("value").coordinate,name=name,value=repr(val),physical_type=type(val).__name__,value_type=vals.get("value_type"),version=repr(ver),version_type=type(ver).__name__,unit=vals.get("unit"),scope=f"{vals.get('scope_type')}/{vals.get('scope_value')}",status=vals.get("status"),ok="sim" if ok else "NÃO"))
            if not ok: add(b,ws.title,cell("value").coordinate,r,"value",val,"BLOCKER","parâmetro exige valor numérico físico (int/float)","número")
            if vals.get("value_type")!="numerico": add(b,ws.title,cell("value_type").coordinate,r,"value_type",vals.get("value_type"),"BLOCKER","parâmetros são só numéricos","numerico")
            if not (isinstance(ver,int) and not isinstance(ver,bool)): add(b,ws.title,cell("version").coordinate,r,"version",ver,"BLOCKER","version de parâmetro deve ser inteiro","inteiro")
            if vals.get("expression"): add(b,ws.title,cell("expression").coordinate,r,"expression",vals["expression"],"BLOCKER","parâmetro sem expressão")
        else:
            if val is not None: add(b,ws.title,cell("value").coordinate,r,"value",val,"NON_BLOCKING","variável com value preenchido (value é campo de parâmetro)")
            if ver is not None: add(b,ws.title,cell("version").coordinate,r,"version",ver,"INFORMATIONAL","variável com version preenchida")
        # enums
        if kind=="variable":
            if vals.get("variable_type") not in V.ALLOWED_VARIABLE_TYPES: add(b,ws.title,cell("variable_type").coordinate,r,"variable_type",vals.get("variable_type"),"BLOCKER","variable_type fora do enum")
            if vals.get("frequency") not in V.ALLOWED_FREQUENCIES: add(b,ws.title,cell("frequency").coordinate,r,"frequency",vals.get("frequency"),"BLOCKER","frequency fora do enum")
            if t=="variable / equation" and not vals.get("expression"): add(b,ws.title,cell("expression").coordinate,r,"expression",None,"BLOCKER","variable / equation sem expressão")
            if t=="variable" and vals.get("expression"): add(b,ws.title,cell("expression").coordinate,r,"expression",vals["expression"],"NON_BLOCKING","Type=variable com expressão")
        else:
            if vals.get("frequency") is not None and vals.get("frequency") not in V.ALLOWED_FREQUENCIES: add(b,ws.title,cell("frequency").coordinate,r,"frequency",vals.get("frequency"),"BLOCKER","frequency fora do enum")
        if vals.get("scope_type") not in V.ALLOWED_SCOPE_TYPES: add(b,ws.title,cell("scope_type").coordinate,r,"scope_type",vals.get("scope_type"),"BLOCKER","scope_type fora do enum")
        if vals.get("scope_value") not in V.ALLOWED_SCOPE_VALUES: add(b,ws.title,cell("scope_value").coordinate,r,"scope_value",vals.get("scope_value"),"BLOCKER","scope_value fora do enum")
        try: SR.resolve_scopes(vals.get("scope_type"),vals.get("scope_value"))
        except Exception as e: add(b,ws.title,cell("scope_value").coordinate,r,"scope_type/scope_value",f"{vals.get('scope_type')}/{vals.get('scope_value')}","BLOCKER",f"ScopeResolver rejeita: {e}")
        u=vals.get("unit")
        if u is None or (isinstance(u,str) and not u.strip()): add(b,ws.title,cell("unit").coordinate,r,"unit",u,"BLOCKER","unit obrigatória")
        elif u not in V.ALLOWED_UNITS: add(b,ws.title,cell("unit").coordinate,r,"unit",u,"BLOCKER","unit fora de ALLOWED_UNITS")
        if u=="tpd": add(b,ws.title,cell("unit").coordinate,r,"unit",u,"BLOCKER","padrão é t/d","t/d")
        st=vals.get("status")
        if st not in V.ALLOWED_STATUSES: status_rows.append((b,r,name,st)); add(b,ws.title,cell("status").coordinate,r,"status",st,"BLOCKER","status ausente/inválido","ativo")
        vt=vals.get("value_type")
        if vt not in VALUE_TYPES: add(b,ws.title,cell("value_type").coordinate,r,"value_type",vt,"BLOCKER","value_type ∉ {numerico, categorico}")
        av=vals.get("allowed_values"); drs=vals.get("declared_result_states")
        if vt=="categorico" and not av: add(b,ws.title,cell("allowed_values").coordinate,r,"allowed_values",av,"BLOCKER","categórica sem allowed_values")
        if vt!="categorico" and av: add(b,ws.title,cell("allowed_values").coordinate,r,"allowed_values",av,"BLOCKER","allowed_values só em categórica")
        if isinstance(av,str):
            items=av.split("\n")
            if any(i!=i.strip() or not i for i in items) or len(set(items))!=len(items): add(b,ws.title,cell("allowed_values").coordinate,r,"allowed_values",av,"BLOCKER","allowed_values com item vazio/duplicado/espaço")
        if drs:
            m=re.fullmatch(r'([A-Z_]+) → "([^"]*)"',drs)
            if not m: add(b,ws.title,cell("declared_result_states").coordinate,r,"declared_result_states",drs,"BLOCKER",'formato canônico RÓTULO → "literal"')
            elif m.group(1) not in TAXONOMY: add(b,ws.title,cell("declared_result_states").coordinate,r,"declared_result_states",drs,"BLOCKER","rótulo fora da taxonomia D2")
        # metadata
        for k in ("description","source_reference"):
            if not vals.get(k): meta[b][k]+=1
json.dump(dict(findings=findings,params=param_rows,status=status_rows,meta={b:dict(c) for b,c in meta.items()},counts={f"{k[0]}:{k[1]}":v for k,v in counts.items()}),open("integral_out.json","w"),ensure_ascii=False,indent=1,default=str)
w=csv.DictWriter(open("parameters.csv","w"),fieldnames=list(param_rows[0])); w.writeheader(); w.writerows(param_rows)
print("rows:",{k:v for k,v in counts.items() if k[1]=="rows"})
print("params:",len(param_rows),"não numéricos:",[p for p in param_rows if p["ok"]!="sim"])
print("status inválido:",status_rows)
print("metadados ausentes:",{b:dict(c) for b,c in meta.items()})
print("findings by class:",collections.Counter((f["workbook"],f["classification"],f["rule"]) for f in findings))
for f in findings:
    if f["classification"]=="BLOCKER": print("  BLOCKER",f)
