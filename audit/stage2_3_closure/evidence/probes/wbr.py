import openpyxl
from wbp import WB
HDR={"yield":1}
def rows(b):
    wb=openpyxl.load_workbook(WB[b]); ws=wb.worksheets[0]
    h=HDR.get(b,2); hdr=[ws.cell(h,c).value for c in range(1,ws.max_column+1)]
    out=[]
    for r in range(h+1, ws.max_row+1):
        vals={(hdr[c-1] or f"col{c}"): ws.cell(r,c).value for c in range(1,ws.max_column+1)}
        if any(v is not None and str(v).strip()!="" for v in vals.values()):
            vals["row"]=r; out.append(vals)
    return out, ws
