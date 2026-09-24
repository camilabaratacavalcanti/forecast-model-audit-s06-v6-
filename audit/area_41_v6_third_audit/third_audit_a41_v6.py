"""
Terceira auditoria independente do workbook A41 v6.

Independência: este script NÃO importa o harness da Etapa 2
(audit/area_41_v5/check_a41_v5.py, audit/area_41_v6/check_a41_v6.py).
Ele reconstrói por conta própria:

  - leitura do xlsx e inventário de entidades;
  - resolução de referências (nome + frequência + escopo) a partir do
    contrato de grupos (L1,L2,L3->L1_L3; L4,L5->L4_L5; L6,L7->L6_L7);
  - DAG producer -> consumer por instância concreta;
  - avaliação do TEXTO das expressões do workbook por um interpretador
    próprio (ast + math), sem o ExpressionEvaluator da plataforma;
  - fórmulas de especificação escritas à mão a partir das decisões
    vinculantes (A020, rateio §7, total §10, ramos §12);
  - álgebra de unidades.

A plataforma é usada apenas como RUNTIME AUDITADO (ForecastEngine,
TemporalAggregationService, ScopeResolver, resolvedor central), e seus
resultados são comparados com as duas fontes independentes acima.

Read-only: não altera workbook, plataforma, seeds, builders ou testes.
Os IDs VAR98xxx/PARAM98xxx/EQ98xxx são efêmeros (só em memória).

Uso:
    python3 audit/area_41_v6_third_audit/third_audit_a41_v6.py
Grava os CSVs e evidence/checks.json; código de saída 0 se nenhuma
verificação FAIL.
"""

import ast
import csv
import hashlib
import json
import math
import os
import random
import re
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT))

EVIDENCE = HERE / "evidence"
PROVENANCE = EVIDENCE / "provenance"

# A41_AUDIT_TARGET só existe para os testes de mutação (evidence/mutation_tests.txt):
# um alvo diferente do v6 validado falha, no mínimo, na verificação de SHA.
AUDITED = Path(os.environ.get("A41_AUDIT_TARGET", ROOT / "audit" / "area_41_v6" / "descritivo_das_variáveis_A41_v6.xlsx"))
AUDITED_SHA = "4b0c41aef6c56a5b655faa543cf96f33f270cdfad8ecd2b4c6a5e1a09144a352"

CHAIN = [
    ("v3", PROVENANCE / "A41_v3.xlsx", "59f82304f13895c0fb45306f260388d650c5a7daec897bfa633d706d07fcfd85"),
    ("v4", PROVENANCE / "A41_v4.xlsx", "383c462b60c9d9806ac93531477533486fb8779350a4808ae9abb676d7282a4c"),
    ("v5_input", PROVENANCE / "A41_v5_input.xlsx", "ea3d2e51863376c34c91c1e3fc55f1ad320b2d715f668634ed586625f9152188"),
    ("v5_validado", ROOT / "audit" / "area_41_v5" / "descritivo_das_variáveis_A41_v5.xlsx", "c8be0a3c70bb04682163ebc489b83a1bc6af58c8105a7bb58c63381c2f915125"),
    ("v6_recebido", ROOT / "audit" / "area_41_v6" / "descritivo_das_variáveis_A41_v6_recebido.xlsx", "aad9b3069069bff47e51ddcd40722a7c3d6ba40a6ccb53fdb732e0bc9613231f"),
    ("v6_validado", AUDITED, AUDITED_SHA),
]

HEADERS = [
    "Type", "name", "description", "unit", "value", "version",
    "variable_type", "frequency", "scope_type", "scope_value",
    "source_reference", "status", "expression", "fonte", "OBS",
    "value_type",
]

LINES = [f"L{i}" for i in range(1, 8)]
GROUPS = {
    "L1_L3": ["L1", "L2", "L3"],
    "L4_L5": ["L4", "L5"],
    "L6_L7": ["L6", "L7"],
    "L1_L7": LINES,
}
SMALL_GROUP = {line: g for g in ("L1_L3", "L4_L5", "L6_L7") for line in GROUPS[g]}
STATES = ["Normal", "LC", "Overhaul/Parada", "1 By pass", "1 By pass e LC"]
AGG_TEXT = re.compile(r"^(Média|Somatório) (mensal|anual) de todos os resultados diários de '([^']+)'$")
KEYWORDS = {"if", "else", "and", "or", "not", "ln", "in", "is"}
TOL = 1e-9

CHECKS: list[dict] = []
FAILFAST: set[str] = set()


def check(gate: str, name: str, ok: bool, detail="") -> bool:
    CHECKS.append({"gate": gate, "check": name, "result": "PASS" if ok else "FAIL", "detail": detail})
    return ok


def close(a, b) -> bool:
    if isinstance(a, str) or isinstance(b, str):
        return a == b
    return math.isclose(a, b, rel_tol=TOL, abs_tol=TOL)


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


MUTATION = "A41_AUDIT_TARGET" in os.environ


def write_csv(name: str, rows: list[dict]) -> None:
    if MUTATION:
        return
    path = HERE / name
    fields = list(rows[0]) if rows else ["empty"]
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


# ============================================================
# 1. Leitura independente do workbook
# ============================================================


def read_workbook(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)
    wb_values = openpyxl.load_workbook(path, data_only=True)
    ws = wb["A41"]
    cells = {}
    formulas = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                cells[cell.coordinate] = cell.value
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append(cell.coordinate)
    cached_diff = [
        c.coordinate for row in wb_values["A41"].iter_rows() for c in row
        if c.value is not None and cells.get(c.coordinate) != c.value
    ]
    header = [ws.cell(2, c).value for c in range(1, ws.max_column + 1)]
    entities = []
    empty_gap = []
    for r in range(3, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(1, len(header) + 1)]
        if not values[1]:
            if any(v not in (None, "") for v in values):
                empty_gap.append(r)
            continue
        e = dict(zip(header, values))
        e["row"] = r
        e.setdefault("value_type", None)
        e["kind"] = "parameter" if e["Type"] == "parameter" else "variable"
        expression = e.get("expression")
        match = AGG_TEXT.match(expression) if isinstance(expression, str) else None
        e["agg"] = match.groups() if match else None
        e["role"] = (
            "parameter" if e["kind"] == "parameter"
            else "aggregation" if match
            else "equation" if expression
            else "input"
        )
        entities.append(e)
    last_entity_row = max(e["row"] for e in entities)
    return {
        "path": path,
        "sheets": wb.sheetnames,
        "dims": (ws.max_row, ws.max_column),
        "merged": sorted(str(m) for m in ws.merged_cells.ranges),
        "header": header,
        "header_groups": [ws.cell(1, c).value for c in range(1, ws.max_column + 1)],
        "cells": cells,
        "formulas": formulas,
        "cached_values_differ": cached_diff,
        "entities": entities,
        "rows_with_content_but_no_name": empty_gap,
        "last_entity_row": last_entity_row,
    }


def identity(e) -> tuple:
    return (e["name"], e["frequency"], e["scope_type"], e["scope_value"])


def concrete(scope_type, scope_value) -> list[tuple[str, str]]:
    if scope_type == "linha":
        if scope_value in LINES:
            return [("linha", scope_value)]
        a, b = scope_value.split("_")
        return [("linha", f"L{i}") for i in range(int(a[1:]), int(b[1:]) + 1)]
    if scope_type == "linha_grupo":
        return [("linha_grupo", scope_value)]
    if scope_type == "planta":
        return [("planta", "PLANTA")]
    raise ValueError(scope_type)


def spatial_chain(instance) -> list[tuple[str, str]]:
    """Contrato de escopo: do consumidor para fora, nunca para dentro."""

    st, sv = instance
    if st == "linha":
        return [instance, ("linha_grupo", SMALL_GROUP[sv]), ("linha_grupo", "L1_L7"), ("planta", "PLANTA")]
    if st == "linha_grupo":
        chain = [instance]
        if sv != "L1_L7":
            chain.append(("linha_grupo", "L1_L7"))
        return chain + [("planta", "PLANTA")]
    return [instance]


# ============================================================
# 2. Resolução independente de referências
# ============================================================

_STRING = re.compile(r"(\"[^\"]*\"|'[^']*')")
_REF = re.compile(r"(?<![\w@])([A-Za-z_][A-Za-z0-9_]*)(?:@(L[1-7](?:_L[1-7])?))?")


def references(expression: str) -> list[tuple[str, str | None]]:
    out = []
    for i, seg in enumerate(_STRING.split(expression)):
        if i % 2:
            continue
        for m in _REF.finditer(seg):
            if m.group(1) not in KEYWORDS:
                out.append((m.group(1), m.group(2)))
    return out


FREQ_RANK = {"diário": 0, "mensal": 1, "anual": 2, None: 3}


def resolve(name, suffix, consumer, instance, by_name) -> dict:
    candidates = by_name.get(name, [])
    if not candidates:
        return {"hits": [], "mode": "orphan", "temporal": None}
    if all(c["kind"] == "parameter" for c in candidates):
        pool, temporal = candidates, "parâmetro"
    else:
        same = [c for c in candidates if c["frequency"] == consumer["frequency"]]
        pool = same or candidates
        temporal = "mesma frequência" if same else "frequência mais grossa (período que contém o dia)"
    if suffix:
        target = ("linha", suffix) if suffix in LINES else ("linha_grupo", suffix)
        hits = [c for c in pool if target in concrete(c["scope_type"], c["scope_value"])]
        return {"hits": hits, "mode": "explícita", "at": target, "temporal": temporal}
    for scope in spatial_chain(instance):
        hits = [c for c in pool if scope in concrete(c["scope_type"], c["scope_value"])]
        if hits:
            return {"hits": hits, "mode": "implícita", "at": scope, "temporal": temporal}
    return {"hits": [], "mode": "implícita-inalcançável", "temporal": temporal}


# ============================================================
# 3. Interpretador independente do texto das expressões
# ============================================================


class Failure(Exception):
    """Consumo do marcador "F" (falha da rotina IF)."""


def encode(expression: str) -> str:
    parts = _STRING.split(expression)
    return "".join(
        seg if i % 2 else re.sub(r"(?<![\w@])([A-Za-z_]\w*)@(L[1-7](?:_L[1-7])?)", r"\1__AT__\2", seg)
        for i, seg in enumerate(parts)
    ).strip()


def interpret(expression: str, lookup) -> object:
    tree = ast.parse(encode(expression), mode="eval")

    def num(v):
        if v == "F":
            raise Failure()
        if isinstance(v, bool) or not isinstance(v, (int, float)):
            raise TypeError(repr(v))
        return v

    def ev(n):
        if isinstance(n, ast.Constant):
            return n.value
        if isinstance(n, ast.Name):
            name, _, suffix = n.id.partition("__AT__")
            return lookup(name, suffix or None)
        if isinstance(n, ast.BinOp):
            a, b = num(ev(n.left)), num(ev(n.right))
            return {ast.Add: a + b, ast.Sub: a - b, ast.Mult: a * b}.get(type(n.op)) if not isinstance(n.op, ast.Div) else a / b
        if isinstance(n, ast.UnaryOp):
            v = num(ev(n.operand))
            return -v if isinstance(n.op, ast.USub) else v
        if isinstance(n, ast.Call):
            assert n.func.id == "ln" and len(n.args) == 1
            v = num(ev(n.args[0]))
            if v <= 0:
                raise ValueError("ln domain")
            return math.log(v)
        if isinstance(n, ast.Compare):
            (op,), (right,) = n.ops, n.comparators
            a, b = ev(n.left), ev(right)
            if "F" in (a, b) and not (isinstance(n.left, ast.Constant) or isinstance(right, ast.Constant)):
                raise Failure()
            return {ast.Eq: a == b, ast.NotEq: a != b, ast.Gt: a > b, ast.Lt: a < b, ast.GtE: a >= b, ast.LtE: a <= b}[type(op)]
        if isinstance(n, ast.BoolOp):
            if isinstance(n.op, ast.And):
                return all(ev(v) for v in n.values)
            return any(ev(v) for v in n.values)
        if isinstance(n, ast.IfExp):
            cond = ev(n.test)
            if cond == "F":
                raise Failure()
            return ev(n.body) if cond else ev(n.orelse)
        raise TypeError(type(n).__name__)

    return ev(tree.body)


ALLOWED_AST = (
    ast.Expression, ast.BinOp, ast.UnaryOp, ast.Constant, ast.Name, ast.Load,
    ast.Add, ast.Sub, ast.Mult, ast.Div, ast.USub, ast.UAdd, ast.Compare,
    ast.Eq, ast.NotEq, ast.Gt, ast.Lt, ast.GtE, ast.LtE, ast.BoolOp, ast.And,
    ast.Or, ast.IfExp, ast.Call,
)


# ============================================================
# 4. Especificação (fórmulas escritas à mão a partir das decisões)
# ============================================================


def corr(x):
    return 19.475 * math.log(x) - 85.999


def spec(inputs: dict) -> dict:
    """inputs: {(name, scope_value): valor}. Retorna {(name, scope_value): valor}."""

    out = {}
    lth = {l: inputs[("lth", l)] for l in LINES}
    out[("lth_grupo", "L1_L3")] = (lth["L1"] + lth["L2"] + lth["L3"]) / 3
    out[("lth_grupo", "L4_L5")] = (lth["L4"] + lth["L5"]) / 2
    out[("lth_grupo", "L6_L7")] = (lth["L6"] + lth["L7"]) / 2
    for g in ("L1_L3", "L4_L5", "L6_L7"):
        out[("retirada_cond_corr_ltp", g)] = corr(inputs[("vazao_ltp", g)])
        out[("retirada_cond_corr_lth", g)] = corr(out[("lth_grupo", g)] * inputs[("fator_retirada_cond_corr_lth", g)])
    d = {g: out[("retirada_cond_corr_ltp", g)] - out[("retirada_cond_corr_lth", g)] for g in ("L1_L3", "L4_L5", "L6_L7")}
    out[("retirada_condensado_grupo", "L1_L3")] = (
        inputs[("valor_retirada", "L1")] + inputs[("valor_retirada", "L2")] + inputs[("valor_retirada", "L3")] - 3 * d["L1_L3"]
    )

    def conditional(a, b, delta, bypass, bypass_lc):
        base = 100 - delta * 2
        if a == "Normal" and b == "Normal":
            return base
        if a == "LC" or b == "LC":
            return (base * 14 + (base / 2) * 10) / 24
        if a == "Overhaul/Parada" or b == "Overhaul/Parada":
            return 50 - delta * 2
        if a == "1 By pass" or b == "1 By pass":
            return base - bypass
        if a == "1 By pass e LC" and b == "1 By pass e LC":
            return (base * 14 + ((base - bypass_lc) / 2) * 10) / 24
        return "F"

    # Descontos conforme a decisão 41c/41d e as expressões do workbook:
    # L4_L5: "1 By pass" -> 41d@L4_L5 ; "1 By pass e LC" -> 41c@L4_L5
    # L6_L7: ambos os ramos -> 41c@L6_L7
    out[("retirada_condensado_grupo", "L4_L5")] = conditional(
        inputs[("hes_l4", "L4")], inputs[("hes_l5", "L5")], d["L4_L5"],
        inputs[("desconto_retirada_41d", "L4_L5")], inputs[("desconto_retirada_41c", "L4_L5")],
    )
    out[("retirada_condensado_grupo", "L6_L7")] = conditional(
        inputs[("hes_l6", "L6")], inputs[("hes_l7", "L7")], d["L6_L7"],
        inputs[("desconto_retirada_41c", "L6_L7")], inputs[("desconto_retirada_41c", "L6_L7")],
    )
    for g in ("L1_L3", "L4_L5", "L6_L7"):
        grp = out[("retirada_condensado_grupo", g)]
        denominator = sum(lth[l] for l in GROUPS[g])
        for l in GROUPS[g]:
            out[("retirada_condensado_linha", l)] = "F" if grp == "F" else grp * lth[l] / denominator
    groups = [out[("retirada_condensado_grupo", g)] for g in ("L1_L3", "L4_L5", "L6_L7")]
    out[("retirada_condensado_total", "L1_L7")] = "F" if "F" in groups else sum(groups)
    return out


# ============================================================
# 5. Modelo independente (DAG + avaliação do texto)
# ============================================================


class Model:
    def __init__(self, wb: dict, order_seed: int | None = None):
        entities = list(wb["entities"])
        if order_seed is not None:
            random.Random(order_seed).shuffle(entities)
        self.entities = entities
        self.by_name = defaultdict(list)
        for e in entities:
            self.by_name[e["name"]].append(e)
        self.equations = [e for e in entities if e["role"] == "equation"]
        self.edges = []  # (producer_key, consumer_key, ref, detail)
        self.resolution = {}  # (consumer_row, instance, name, suffix) -> hit entity
        for c in self.equations:
            for inst in concrete(c["scope_type"], c["scope_value"]):
                for name, suffix in dict.fromkeys(references(c["expression"])):
                    r = resolve(name, suffix, c, inst, self.by_name)
                    self.resolution[(c["row"], inst, name, suffix)] = r
                    if len(r["hits"]) == 1:
                        p = r["hits"][0]
                        self.edges.append(((p["row"], r["at"]), (c["row"], inst), name + (f"@{suffix}" if suffix else ""), r))

    def topo(self) -> list:
        nodes = {(c["row"], i) for c in self.equations for i in concrete(c["scope_type"], c["scope_value"])}
        deps = defaultdict(set)
        for p, c, _, _ in self.edges:
            if p in nodes:
                deps[c].add(p)
        order, seen, visiting = [], set(), set()

        def visit(n):
            if n in seen:
                return
            if n in visiting:
                raise RuntimeError(f"ciclo em {n}")
            visiting.add(n)
            for p in sorted(deps[n]):
                visit(p)
            visiting.discard(n)
            seen.add(n)
            order.append(n)

        for n in sorted(nodes):
            visit(n)
        return order

    def evaluate(self, inputs: dict) -> dict:
        """Avalia o TEXTO do workbook. Retorna {(name, scope_value): valor | exceção}."""

        row_of = {e["row"]: e for e in self.entities}
        values = {}
        for e in self.entities:
            if e["role"] == "input":
                for inst in concrete(e["scope_type"], e["scope_value"]):
                    values[(e["row"], inst)] = inputs[(e["name"], inst[1])]
        results = {}
        for node in self.topo():
            entity = row_of[node[0]]

            def lookup(name, suffix, _node=node):
                r = self.resolution[(_node[0], _node[1], name, suffix)]
                (hit,) = r["hits"]
                return values[(hit["row"], r["at"])] if (hit["row"], r["at"]) in values else values[(hit["row"], concrete(hit["scope_type"], hit["scope_value"])[0])]

            try:
                v = interpret(entity["expression"], lookup)
            except Failure:
                v = FailureMarker("consumo de F")
            except ZeroDivisionError:
                v = FailureMarker("divisão por zero")
            except ValueError as exc:
                v = FailureMarker(str(exc))
            except KeyError:
                v = FailureMarker("dependência não calculada")
            if not isinstance(v, FailureMarker):
                values[node] = v
            results[(entity["name"], node[1][1])] = v
        return results

    def descendants(self, input_key) -> set:
        """input_key: (row, instance). Retorna {(name, scope_value)} afetados."""

        row_of = {e["row"]: e for e in self.entities}
        children = defaultdict(set)
        for p, c, _, _ in self.edges:
            children[p].add(c)
        out, stack = set(), [input_key]
        while stack:
            n = stack.pop()
            for c in children[n]:
                if c not in out:
                    out.add(c)
                    stack.append(c)
        return {(row_of[r]["name"], i[1]) for r, i in out}


class FailureMarker:
    def __init__(self, why):
        self.why = why

    def __repr__(self):
        return f"<falha: {self.why}>"


# ============================================================
# 6. Runtime auditado (plataforma genérica)
# ============================================================

from app.domain.equations.models import EquationDefinition  # noqa: E402
from app.domain.equations.registry import EquationDefinitionRegistry  # noqa: E402
from app.domain.forecast.aggregation import AggregationRule  # noqa: E402
from app.domain.variables.models import VariableDefinition  # noqa: E402
from app.domain.variables.registry import VariableDefinitionRegistry  # noqa: E402
from app.engine import reference_resolver  # noqa: E402
from app.engine.calculation_context import CalculationContext, CalculationKey  # noqa: E402
from app.engine.exceptions import (  # noqa: E402
    AggregationFailureError,
    ConditionalFailureError,
    DivisionByZeroError,
    EquationEvaluationError,
    ExpressionTypeError,
    MathDomainError,
    UnsafeExpressionError,
)
from app.engine.expression_evaluator import ExpressionEvaluator  # noqa: E402
from app.engine.expression_parser import ExpressionParser  # noqa: E402
from app.engine.forecast_engine import ForecastEngine  # noqa: E402
from app.engine.scope_resolver import ScopeResolver  # noqa: E402
from app.engine.temporal_aggregation_service import TemporalAggregationService  # noqa: E402


class Runtime:
    def __init__(self, wb: dict):
        self.entities = wb["entities"]
        counters = defaultdict(int)
        self.ids = {}
        docs = []
        for e in self.entities:
            counters[e["kind"]] += 1
            prefix = "PARAM98" if e["kind"] == "parameter" else "VAR98"
            eid = f"{prefix}{counters[e['kind']]:03d}"
            self.ids[e["row"]] = eid
            docs.append({"entity_id": eid, "name": e["name"], "kind": e["kind"], "frequency": e["frequency"],
                         "scope_type": e["scope_type"], "scope_value": e["scope_value"]})
        self.index = reference_resolver.build_name_index(docs)
        self.by_id = {self.ids[e["row"]]: e for e in self.entities}
        self.variables = VariableDefinitionRegistry()
        for e in self.entities:
            if e["kind"] == "variable":
                self.variables.add(VariableDefinition(
                    self.ids[e["row"]], e["name"], e["description"], e["unit"], e["variable_type"], e["frequency"],
                    e["scope_type"], e["scope_value"], e["source_reference"], e["status"], e["value_type"] or "numeric"))
        self.definitions = []
        for n, e in enumerate(self.entities, start=1):
            if e["role"] == "equation":
                text = reference_resolver.translate_expression(
                    e["expression"], self.index, e["frequency"], consumer_scope=(e["scope_type"], e["scope_value"]))
                self.definitions.append(EquationDefinition(
                    f"EQ98{n:03d}", self.ids[e["row"]], 1, e["scope_type"], e["scope_value"], text,
                    e["source_reference"], "PUBLISHED"))

    def platform_resolution(self, consumer, name, suffix) -> dict:
        eid = reference_resolver.resolve_reference(
            name, self.index, consumer["frequency"],
            consumer_scope=(consumer["scope_type"], consumer["scope_value"]),
            explicit_scope_value=suffix)
        return self.by_id[eid]

    def context(self):
        ctx = CalculationContext()
        ctx.declare_categorical_variables(
            d.variable_definition_id for d in self.variables.all() if d.is_categorical)
        return ctx

    def set_inputs(self, ctx, day: date, inputs: dict):
        for e in self.entities:
            if e["role"] != "input":
                continue
            for st, sv in concrete(e["scope_type"], e["scope_value"]):
                ctx.set_variable_value(
                    self.ids[e["row"]], inputs[(e["name"], sv)], st, sv,
                    period_id=day.isoformat()[:7] if e["frequency"] == "mensal" else day.isoformat())

    def run(self, inputs: dict, day=date(2026, 7, 1), ctx=None, order=None):
        ctx = ctx or self.context()
        self.set_inputs(ctx, day, inputs)
        registry = EquationDefinitionRegistry()
        for d in (order if order is not None else self.definitions):
            registry.add(d)
        error = None
        try:
            ForecastEngine().calculate_from_definition_registry(registry, ctx, self.variables, run_date=day)
        except EquationEvaluationError as exc:
            error = exc
        return ctx, error

    def outputs(self, ctx, day=date(2026, 7, 1)) -> dict:
        out = {}
        for e in self.entities:
            if e["role"] != "equation":
                continue
            for st, sv in concrete(e["scope_type"], e["scope_value"]):
                key = CalculationKey(self.ids[e["row"]], st, sv, day.isoformat())
                if key in ctx._scoped_variables:
                    out[(e["name"], sv)] = ctx._scoped_variables[key]
        return out


# ============================================================
# 7. Cenários
# ============================================================


def scenario(lth, valor=(47.0, 47.0, 48.0), vazao=1320.0, fator=1.1, states=None, d41c45=5.0, d41c67=6.0, d41d45=4.0):
    states = states or {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal"}
    s = {("lth", l): v for l, v in zip(LINES, lth)}
    for g in ("L1_L3", "L4_L5", "L6_L7"):
        s[("vazao_ltp", g)] = vazao if not isinstance(vazao, dict) else vazao[g]
        s[("fator_retirada_cond_corr_lth", g)] = fator if not isinstance(fator, dict) else fator[g]
    for l, v in zip(("L1", "L2", "L3"), valor):
        s[("valor_retirada", l)] = v
    for l in ("L4", "L5", "L6", "L7"):
        s[(f"hes_{l.lower()}", l)] = states[l]
    s[("desconto_retirada_41c", "L4_L5")] = d41c45
    s[("desconto_retirada_41c", "L6_L7")] = d41c67
    s[("desconto_retirada_41d", "L4_L5")] = d41d45
    return s


VALUE_SETS = {
    "diferentes": scenario([1000.0, 1100.0, 1200.0, 900.0, 1040.0, 1000.0, 1200.0],
                           vazao={"L1_L3": 1320.0, "L4_L5": 1300.0, "L6_L7": 1250.0},
                           fator={"L1_L3": 1.1, "L4_L5": 1.05, "L6_L7": 1.15}),
    "iguais": scenario([1100.0] * 7),
    "assimetricos": scenario([10.0, 2000.0, 5.0, 300.0, 1.0, 999.0, 0.5]),
    "proximos_de_zero": scenario([1e-6, 2e-6, 3e-6, 1e-6, 5e-6, 2e-6, 1e-6]),
    "enunciado_21": scenario([10.0, 20.0, 30.0, 40.0, 60.0, 70.0, 30.0]),
}


# ============================================================
# 8. Auditoria
# ============================================================


def main() -> int:
    EVIDENCE.mkdir(exist_ok=True)
    evidence = {}
    mutation_run = "A41_AUDIT_TARGET" in os.environ

    # ---------- identidade e proveniência ----------
    books = {}
    for tag, path, expected_sha in CHAIN:
        actual = sha256(path)
        check("identidade", f"sha256 {tag}", actual == expected_sha, f"{rel(path)} {actual}")
        books[tag] = read_workbook(path)
    wb = books["v6_validado"]
    evidence["identity"] = {
        "audited": rel(AUDITED), "sha256": sha256(AUDITED), "size": AUDITED.stat().st_size,
        "sheets": wb["sheets"], "dims": wb["dims"], "merged": wb["merged"], "formulas": wb["formulas"],
        "cached_values_differ": wb["cached_values_differ"], "entities": len(wb["entities"]),
        "last_entity_row": wb["last_entity_row"],
    }
    check("identidade", "1 aba A41", wb["sheets"] == ["A41"])
    check("integridade", "sem fórmulas Excel", wb["formulas"] == [])
    check("integridade", "valores em cache == valores armazenados", wb["cached_values_differ"] == [])
    check("integridade", "cabeçalho = 15 colunas históricas + value_type", wb["header"] == HEADERS, str(wb["header"]))
    check("integridade", "merged ranges históricos preservados", wb["merged"] == books["v4"]["merged"], str(wb["merged"]))
    check("integridade", "nenhuma linha com conteúdo sem name", wb["rows_with_content_but_no_name"] == [])
    check("integridade", "nenhuma célula preenchida após a última entidade",
          all(int(re.sub(r"[A-Z]", "", c)) <= wb["last_entity_row"] for c in wb["cells"]))

    # change matrix
    changes = []

    def cell_diff(a_tag, b_tag, expected: dict):
        a, b = books[a_tag]["cells"], books[b_tag]["cells"]
        for coord in sorted(set(a) | set(b), key=lambda c: (int(re.sub(r"[A-Z]", "", c)), re.sub(r"\d", "", c))):
            if a.get(coord) != b.get(coord):
                row = int(re.sub(r"[A-Z]", "", coord))
                ent = books[b_tag]["cells"].get(f"B{row}") or books[a_tag]["cells"].get(f"B{row}")
                scope = books[b_tag]["cells"].get(f"J{row}") or books[a_tag]["cells"].get(f"J{row}")
                rule = next((v for k, v in expected.items() if re.fullmatch(k, coord)), None)
                changes.append({
                    "transicao": f"{a_tag} -> {b_tag}", "celula": coord,
                    "entidade": f"{ent}@{scope}" if row > 2 else "(cabeçalho)",
                    "antes": "" if a.get(coord) is None else str(a.get(coord)),
                    "depois": "" if b.get(coord) is None else str(b.get(coord)),
                    "natureza": rule[0] if rule else "NÃO PREVISTA",
                    "esperada": "sim" if rule else "NÃO",
                    "justificativa": rule[1] if rule else "",
                })

    def identity_diff(a_tag, b_tag):
        def bag(tag):
            out = defaultdict(list)
            for e in books[tag]["entities"]:
                out[identity(e)].append(e)
            return out
        a, b = bag(a_tag), bag(b_tag)
        for key in sorted(set(a) | set(b), key=str):
            if len(a.get(key, [])) != len(b.get(key, [])):
                changes.append({"transicao": f"{a_tag} -> {b_tag}", "celula": "(identidade)", "entidade": "@".join(map(str, key)),
                                "antes": f"{len(a.get(key, []))} linha(s)", "depois": f"{len(b.get(key, []))} linha(s)",
                                "natureza": "estrutura de entidades", "esperada": "sim",
                                "justificativa": "Consolidação v3->v4 (duplicatas Somatório/Média removidas; lth unificado em L1_L7; linha diária por linha; PLANTA canônico)."})
            elif a.get(key) and b.get(key):
                for field in ("expression", "source_reference", "description", "unit", "frequency", "OBS", "variable_type", "status"):
                    if a[key][0].get(field) != b[key][0].get(field):
                        changes.append({"transicao": f"{a_tag} -> {b_tag}", "celula": f"({field})", "entidade": "@".join(map(str, key)),
                                        "antes": str(a[key][0].get(field)), "depois": str(b[key][0].get(field)),
                                        "natureza": field, "esperada": "sim",
                                        "justificativa": "Correções v3->v4 (A020 média de lth_grupo; ln(lth_grupo·fator); source_reference; OBS 41c/41d)."})

    identity_diff("v3", "v4")
    cell_diff("v4", "v5_input", {r"C\d+": ("description", "Descrições aprovadas (vazias no v4)."),
                                 r"M4[5-9]|M5[01]": ("expression", "§7: denominador Σ lth do grupo em vez de lth_grupo.")})
    cell_diff("v5_input", "v5_validado", {r"P\d+": ("value_type (D1)", "Contrato value_type da Etapa 1."),
                                          r"M39|M42": ("expression (D2)", "hes_lN -> hes_lN@LN (referência implícita não desce do grupo à linha).")})
    cell_diff("v5_validado", "v6_recebido", {r"K30|K31": ("source_reference", "Decisão 41c/41d do usuário."),
                                             r"P\d+": ("value_type (D1) AUSENTE", "v6 recebido derivado do v5_input: perdeu D1."),
                                             r"M39|M42": ("expression (D2) AUSENTE", "v6 recebido derivado do v5_input: perdeu D2.")})
    cell_diff("v6_recebido", "v6_validado", {r"P\d+": ("value_type (D1) reaplicado", "Correção mínima do fechamento da Etapa 2."),
                                             r"M39|M42": ("expression (D2) reaplicada", "Correção mínima do fechamento da Etapa 2.")})
    cell_diff("v5_validado", "v6_validado", {r"K30|K31": ("source_reference", "Decisão 41c/41d do usuário.")})
    write_csv("change_matrix.csv", changes)
    unexpected = [c for c in changes if c["esperada"] != "sim"]
    check("proveniência", "todas as alterações da cadeia são previstas", unexpected == [], f"{len(unexpected)} não previstas")
    net = [c for c in changes if c["transicao"] == "v5_validado -> v6_validado"]
    check("proveniência", "v6_validado = v5_validado + K30,K31", sorted(c["celula"] for c in net) == ["K30", "K31"], str([c["celula"] for c in net]))
    check("proveniência", "v6_recebido = v5_input + K30,K31",
          sorted(c for c in set(books["v5_input"]["cells"]) | set(books["v6_recebido"]["cells"])
                 if books["v5_input"]["cells"].get(c) != books["v6_recebido"]["cells"].get(c)) == ["K30", "K31"])
    check("identidade", "dimensões/merged/cabeçalho iguais entre v5_validado e v6_validado",
          (books["v5_validado"]["dims"], books["v5_validado"]["merged"], books["v5_validado"]["header"]) ==
          (wb["dims"], wb["merged"], wb["header"]))

    # ---------- inventário ----------
    entities = wb["entities"]
    inventory = [{k: ("" if e.get(k) is None else e.get(k)) for k in ["row"] + HEADERS + ["role"]} for e in entities]
    for row in inventory:
        row["expression"] = str(row["expression"]).replace("\n", " ")
    write_csv("entity_inventory.csv", inventory)
    ids = [identity(e) for e in entities]
    check("inventário", "54 entidades (17 entradas, 20 equações, 10 agregações, 7 parâmetros)",
          (len(entities), sum(e["role"] == "input" for e in entities), sum(e["role"] == "equation" for e in entities),
           sum(e["role"] == "aggregation" for e in entities), sum(e["role"] == "parameter" for e in entities)) == (54, 17, 20, 10, 7))
    check("inventário", "identidades únicas", len(ids) == len(set(ids)))
    check("inventário", "mesmo conjunto de identidades de v5_validado",
          sorted(map(str, ids)) == sorted(str(identity(e)) for e in books["v5_validado"]["entities"]))
    check("inventário", "descrições não vazias", all(isinstance(e["description"], str) and e["description"].strip() for e in entities))
    bad_scopes = []
    for e in entities:
        try:
            concrete(e["scope_type"], e["scope_value"])
            ScopeResolver().resolve_scopes(e["scope_type"], e["scope_value"])
        except Exception as exc:  # noqa: BLE001
            bad_scopes.append((e["row"], str(exc)))
    check("inventário", "escopos válidos", bad_scopes == [], str(bad_scopes))

    # ---------- value_type / variable_type ----------
    vt_rows = []
    categorical_expected = {"hes_l4", "hes_l5", "hes_l6", "hes_l7"}
    text_compared = defaultdict(set)
    for e in entities:
        if e["role"] == "equation":
            for m in re.finditer(r"([A-Za-z_]\w*)(?:@\w+)?\s*==\s*\"([^\"]*)\"", e["expression"]):
                text_compared[m.group(1)].add(m.group(2))
    for e in entities:
        expected = None if e["kind"] == "parameter" else ("categorical" if e["name"] in categorical_expected else "numeric")
        used_as = (
            f"comparado a textos {sorted(text_compared[e['name']])}" if text_compared.get(e["name"])
            else "parâmetro numérico" if e["kind"] == "parameter" else "operando numérico / resultado"
        )
        vt_rows.append({
            "row": e["row"], "entidade": f"{e['name']}@{e['scope_value']}", "kind": e["kind"],
            "variable_type": e["variable_type"] or "", "value_type": e["value_type"] or "",
            "conteudo_esperado": expected or "(não se aplica: parâmetro numérico)", "conteudo_utilizado": used_as,
            "valido": "sim" if (e["value_type"] or None) == expected else "NÃO",
        })
    write_csv("value_type_matrix.csv", vt_rows)
    check("value_type", "value_type correto em todas as entidades", all(r["valido"] == "sim" for r in vt_rows))
    check("value_type", "variable_type e value_type não são sinônimos (conjuntos de valores disjuntos)",
          not ({e["variable_type"] for e in entities} & {e["value_type"] for e in entities} - {None}))
    check("value_type", "variáveis comparadas a texto são exatamente as categóricas",
          set(text_compared) == categorical_expected, str(sorted(text_compared)))
    check("value_type", "estados usados nas equações = 5 estados previstos",
          set().union(*text_compared.values()) == set(STATES), str(sorted(set().union(*text_compared.values()))))

    # ---------- A019: resolução ----------
    model = Model(wb)
    runtime = Runtime(wb)
    scope_rows = []
    for (row, inst, name, suffix), r in model.resolution.items():
        consumer = next(e for e in entities if e["row"] == row)
        hit = r["hits"][0] if len(r["hits"]) == 1 else None
        platform = runtime.platform_resolution(consumer, name, suffix)
        scope_rows.append({
            "consumidor": f"{consumer['name']}@{consumer['scope_value']}", "linha_consumidor": row,
            "referencia": name + (f"@{suffix}" if suffix else ""), "frequencia_consumidor": consumer["frequency"],
            "escopo_consumidor": f"{inst[0]}/{inst[1]}", "modo": r["mode"],
            "produtor_resolvido": f"{hit['name']} (linha {hit['row']})" if hit else "",
            "frequencia_produtor": (hit["frequency"] or "parâmetro") if hit else "",
            "escopo_produtor": f"{hit['scope_type']}/{hit['scope_value']}" if hit else "",
            "instancia_lida": f"{r['at'][0]}/{r['at'][1]}" if hit else "",
            "temporal": r["temporal"],
            "produtor_plataforma": f"{platform['name']} (linha {platform['row']})",
            "deterministico": "sim" if hit else "NÃO",
            "independente_igual_plataforma": "sim" if hit and hit["row"] == platform["row"] else "NÃO",
        })
    write_csv("scope_matrix.csv", scope_rows)
    check("A019", "todas as referências resolvem para exatamente 1 produtor", all(r["deterministico"] == "sim" for r in scope_rows), f"{len(scope_rows)} referências")
    check("A019", "resolução independente == resolvedor central da plataforma", all(r["independente_igual_plataforma"] == "sim" for r in scope_rows))
    check("A019", "nenhum produtor mais fino que o consumidor (só mesma frequência ou mais grossa)",
          all(FREQ_RANK[next(e for e in entities if f"{e['name']} (linha {e['row']})" == r['produtor_resolvido'])["frequency"]] >= FREQ_RANK[r["frequencia_consumidor"]] for r in scope_rows))
    reordered = [Model(wb, order_seed=s) for s in (1, 2, 3)]
    check("A019", "resolução independente invariante à ordem das entidades",
          all({k: [h["row"] for h in v["hits"]] for k, v in m.resolution.items()} ==
              {k: [h["row"] for h in v["hits"]] for k, v in model.resolution.items()} for m in reordered))

    # ---------- @grupo ----------
    grp_rows = [r for r in scope_rows if r["referencia"].startswith("retirada_condensado_grupo@")]
    expected_group = {"L1": "L1_L3", "L2": "L1_L3", "L3": "L1_L3", "L4": "L4_L5", "L5": "L4_L5", "L6": "L6_L7", "L7": "L6_L7"}
    ok = all(
        r["referencia"] == f"retirada_condensado_grupo@{expected_group[r['escopo_consumidor'].split('/')[1]]}"
        and r["escopo_produtor"] == f"linha_grupo/{expected_group[r['escopo_consumidor'].split('/')[1]]}"
        for r in grp_rows if r["consumidor"].startswith("retirada_condensado_linha")
    )
    check("@grupo", "L1,L2,L3->L1_L3; L4,L5->L4_L5; L6,L7->L6_L7", ok and len([r for r in grp_rows if r["consumidor"].startswith("retirada_condensado_linha")]) == 7)
    check("@grupo", "total referencia os três grupos explicitamente",
          sorted(r["referencia"] for r in grp_rows if r["consumidor"].startswith("retirada_condensado_total")) ==
          ["retirada_condensado_grupo@L1_L3", "retirada_condensado_grupo@L4_L5", "retirada_condensado_grupo@L6_L7"])
    explicit = [r for r in scope_rows if r["modo"] == "explícita"]
    check("@grupo", "toda referência explícita lê exatamente o escopo declarado",
          all(r["instancia_lida"].split("/")[1] == r["referencia"].split("@")[1] for r in explicit), f"{len(explicit)} explícitas")

    # ---------- dependências ----------
    dep_rows = []
    row_of = {e["row"]: e for e in entities}
    for (prow, pinst), (crow, cinst), ref, r in model.edges:
        dep_rows.append({"produtor": f"{row_of[prow]['name']}@{pinst[1]}", "papel_produtor": row_of[prow]["role"],
                         "consumidor": f"{row_of[crow]['name']}@{cinst[1]}", "referencia": ref, "modo": r["mode"]})
    agg_edges = []
    for e in entities:
        if e["role"] == "aggregation":
            for inst in concrete(e["scope_type"], e["scope_value"]):
                r = resolve(e["agg"][2], None, {"frequency": "diário"}, inst, model.by_name)
                (hit,) = r["hits"]
                agg_edges.append((hit, e, inst))
                dep_rows.append({"produtor": f"{hit['name']}@{inst[1]} (diário)", "papel_produtor": hit["role"],
                                 "consumidor": f"{e['name']}@{inst[1]} ({e['frequency']})", "referencia": e["agg"][2], "modo": f"agregação {e['agg'][0]}"})
    write_csv("dependency_matrix.csv", dep_rows)
    try:
        order = model.topo()
        acyclic = True
    except RuntimeError as exc:
        acyclic, order = False, str(exc)
    check("dependências", "DAG acíclico", acyclic)
    consumed = {row_of[p[0]]["name"] for p, _, _, _ in model.edges} | {h["name"] for h, _, _ in agg_edges}
    unused = sorted({e["name"] for e in entities if e["role"] in ("input", "parameter")} - consumed)
    evidence["unused_inputs_and_parameters"] = unused
    check("dependências", "entradas/parâmetros não consumidos = só os 7 parâmetros documentados como não utilizados",
          unused == sorted(e["name"] for e in entities if e["role"] == "parameter"), str(unused))
    check("dependências", "toda equação diária produz valor consumido ou agregado",
          all(e["name"] in consumed or e["name"] in {"retirada_condensado_total", "retirada_condensado_linha"}
              for e in entities if e["role"] == "equation"))

    # ---------- segurança DSL ----------
    forbidden = re.compile(r"\b(eval|exec|__import__|open|getattr|setattr|lambda|import|globals|locals|compile)\b|\[|\]|\{|\}|[A-Za-z_]\s*\.\s*[A-Za-z_]|(?<![=!<>])=(?!=)")
    sec_issues = []
    for e in entities:
        if e["role"] != "equation":
            continue
        code = "".join(seg for i, seg in enumerate(_STRING.split(e["expression"])) if i % 2 == 0)
        if forbidden.search(code):
            sec_issues.append((e["row"], forbidden.search(code).group(0)))
        for node in ast.walk(ast.parse(encode(e["expression"]), mode="eval")):
            if not isinstance(node, ALLOWED_AST):
                sec_issues.append((e["row"], type(node).__name__))
            if isinstance(node, ast.Call) and (not isinstance(node.func, ast.Name) or node.func.id != "ln" or len(node.args) != 1 or node.keywords):
                sec_issues.append((e["row"], "chamada não permitida"))
        try:
            ExpressionParser().parse(runtime.definitions[[d.target_variable_id for d in runtime.definitions].index(runtime.ids[e["row"]])].expression)
        except UnsafeExpressionError as exc:
            sec_issues.append((e["row"], f"parser: {exc}"))
    check("segurança DSL", "nenhuma construção proibida; só ln/1 argumento; parser da plataforma aceita as 20 equações", sec_issues == [], str(sec_issues))

    # ---------- ln ----------
    ln_calls = [(e["row"], ast.unparse(n)) for e in entities if e["role"] == "equation"
                for n in ast.walk(ast.parse(encode(e["expression"]), mode="eval")) if isinstance(n, ast.Call)]
    evidence["ln_calls"] = ln_calls
    check("ln", "6 ocorrências de ln, todas com 1 argumento", len(ln_calls) == 6, str(ln_calls))
    ev = ExpressionEvaluator
    pp = ExpressionParser()

    def platform_eval(expr, variables=None, categorical=()):
        ctx = CalculationContext(variables or {}, categorical_variable_ids=categorical)
        return ev(ctx).evaluate(pp.parse(expr))

    ln_ok = (math.isclose(platform_eval("ln(1)"), 0.0) and math.isclose(platform_eval("ln(VAR98999)", {"VAR98999": math.e}), 1.0))
    errors = []
    for value in (0, -1.0):
        try:
            platform_eval("ln(VAR98999)", {"VAR98999": value})
            errors.append(f"ln({value}) não falhou")
        except MathDomainError:
            pass
    try:
        ctx = CalculationContext(categorical_variable_ids={"VAR98998"})
        ctx.set_variable("VAR98998", "Normal")
        ev(ctx).evaluate(pp.parse("ln(VAR98998)"))
        errors.append("ln(categorical) não falhou")
    except ExpressionTypeError:
        pass
    for bad in ("log(VAR98999)", "exp(1)", "abs(1)", "ln(1, 2)", "ln(x=1)", "__import__('os')", "VAR98999.real", "getattr(VAR98999, 'x')"):
        try:
            pp.parse(bad)
            errors.append(f"aceitou {bad}")
        except UnsafeExpressionError:
            pass
    check("ln", "ln(1)=0, ln(e)=1, ln(0)/ln(<0)->MathDomainError, ln(categorical)->ExpressionTypeError, allowlist fechada", ln_ok and not errors, str(errors))

    # ---------- equações: spec vs texto independente vs runtime ----------
    eq_rows = []

    def compare(label, inputs):
        s = spec(inputs)
        ind = model.evaluate(inputs)
        ctx, err = runtime.run(inputs)
        rt = runtime.outputs(ctx)
        all_ok = True
        for key in sorted(s):
            sv, iv, rv = s[key], ind.get(key), rt.get(key)
            if sv == "F" or (isinstance(iv, FailureMarker) and sv == "F"):
                pass
            if isinstance(iv, FailureMarker):
                iv_show = repr(iv)
            else:
                iv_show = iv
            expected_failure = sv == "F"
            scenario_has_failure = "F" in s.values()
            status = None
            if expected_failure:
                # "F" direto (grupo) deve ser gravado; consumidores não são calculados.
                is_group = key[0] == "retirada_condensado_grupo"
                ok = (iv == "F" and rv == "F") if is_group else (isinstance(iv, FailureMarker) and rv is None)
            elif rv is None and scenario_has_failure:
                # fail-fast da plataforma: a rodada do dia para no primeiro
                # consumo de "F"; saídas independentes podem não ser
                # calculadas. Registrado como observação (FAILFAST), não
                # como divergência de valor.
                ok = not isinstance(iv, FailureMarker) and close(sv, iv)
                status = "NAO_CALCULADO (fail-fast)" if ok else "FAIL"
                FAILFAST.add(f"{key[0]}@{key[1]}")
            else:
                ok = not isinstance(iv, FailureMarker) and rv is not None and close(sv, iv) and close(sv, rv)
            all_ok &= ok
            eq_rows.append({"cenario": label, "entidade": f"{key[0]}@{key[1]}", "spec": sv, "texto_independente": iv_show,
                            "runtime": "" if rv is None else rv, "status": status or ("PASS" if ok else "FAIL")})
        return all_ok, s, rt, err

    for label, inputs in VALUE_SETS.items():
        ok, s, rt, err = compare(f"valores:{label}", inputs)
        check("equações", f"spec == texto == runtime ({label})", ok and err is None, str(err))
        rateio_ok = all(close(sum(rt[("retirada_condensado_linha", l)] for l in GROUPS[g]), rt[("retirada_condensado_grupo", g)])
                        for g in ("L1_L3", "L4_L5", "L6_L7"))
        check("rateio", f"Σ linhas == grupo ({label})", rateio_ok)
        total_ok = close(rt[("retirada_condensado_total", "L1_L7")], sum(rt[("retirada_condensado_grupo", g)] for g in ("L1_L3", "L4_L5", "L6_L7")))
        check("total", f"total == Σ grupos ({label})", total_ok)
        if label == "enunciado_21":
            check("A020", "lth_grupo = 20/50/50 (média, não soma)",
                  [rt[("lth_grupo", g)] for g in ("L1_L3", "L4_L5", "L6_L7")] == [20.0, 50.0, 50.0])
    for g, (a, b) in {"L4_L5": ("L4", "L5"), "L6_L7": ("L6", "L7")}.items():
        state_ok = True
        for sa in STATES + ["Overhaul"]:
            for sb in STATES:
                states = {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal", a: sa, b: sb}
                inputs = dict(VALUE_SETS["diferentes"])
                inputs.update({(f"hes_{l.lower()}", l): states[l] for l in ("L4", "L5", "L6", "L7")})
                ok, s, rt, err = compare(f"estados {g}: {sa} / {sb}", inputs)
                state_ok &= ok
                if s[("retirada_condensado_grupo", g)] == "F":
                    state_ok &= err is not None and isinstance(err.original_error, ConditionalFailureError)
        check("texto/booleano", f"30 combinações de estados {g}: spec == texto == runtime (inclui F)", state_ok)
    # ordem dos ramos demonstrada por combinações que satisfazem 2 condições
    precedence = {("LC", "Overhaul/Parada"): "LC", ("Overhaul/Parada", "1 By pass"): "Overhaul", ("1 By pass", "1 By pass e LC"): "By pass",
                  ("LC", "1 By pass e LC"): "LC", ("Normal", "1 By pass e LC"): "F"}
    prec_detail = {}
    base = VALUE_SETS["diferentes"]
    s0 = spec(base)
    d45 = s0[("retirada_cond_corr_ltp", "L4_L5")] - s0[("retirada_cond_corr_lth", "L4_L5")]
    b45 = 100 - 2 * d45
    branch_value = {"LC": (b45 * 14 + b45 / 2 * 10) / 24, "Overhaul": 50 - 2 * d45, "By pass": b45 - 4.0, "F": "F"}
    for (sa, sb), branch in precedence.items():
        inputs = dict(base)
        inputs.update({("hes_l4", "L4"): sa, ("hes_l5", "L5"): sb})
        ctx, err = runtime.run(inputs)
        got = runtime.outputs(ctx).get(("retirada_condensado_grupo", "L4_L5"))
        prec_detail[f"{sa} / {sb}"] = {"ramo_esperado": branch, "runtime": got}
        check("texto/booleano", f"precedência {sa} / {sb} -> {branch}", close(got, branch_value[branch]) if got is not None else False)
    evidence["precedence"] = prec_detail
    write_csv("equation_validation.csv", eq_rows)

    # ---------- "F" ----------
    f_inputs = dict(VALUE_SETS["diferentes"])
    f_inputs.update({("hes_l4", "L4"): "Normal", ("hes_l5", "L5"): "1 By pass e LC"})
    ctx, err = runtime.run(f_inputs)
    stored = runtime.outputs(ctx).get(("retirada_condensado_grupo", "L4_L5"))
    check('"F"', 'grupo L4_L5 grava exatamente o texto "F" (não 0/False/None/"0")',
          isinstance(stored, str) and stored == "F" and stored not in (0, False, None, "0"))
    check('"F"', "consumidores falham com ConditionalFailureError (erro contratual)",
          err is not None and isinstance(err.original_error, ConditionalFailureError))
    computed_under_f = sorted(f"{k[0]}@{k[1]}" for k in runtime.outputs(ctx))
    not_computed = sorted(f"{k[0]}@{k[1]}" for k in spec(VALUE_SETS["diferentes"]) if f"{k[0]}@{k[1]}" not in computed_under_f)
    grp45 = next(e for e in entities if e["name"] == "retirada_condensado_grupo" and e["frequency"] == "diário" and e["scope_value"] == "L4_L5")
    dependents = {f"{a}@{b}" for a, b in model.descendants((grp45["row"], ("linha_grupo", "L4_L5")))}
    collateral = sorted(set(not_computed) - dependents)
    evidence["F_fail_fast"] = {"computed": computed_under_f, "not_computed": not_computed,
                               "dependents_of_failed_group": sorted(dependents),
                               "collateral_not_computed_independent_of_failure": collateral}
    check('"F"', "saídas não calculadas sob F ⊇ dependentes do grupo em falha", dependents <= set(not_computed))
    for text, exc in (('"F" if VAR98999 > 0 else 1', None),):
        pass
    try:
        platform_eval("VAR98999 + 1", {"VAR98999": "F"})
        f_consume = False
    except ConditionalFailureError:
        f_consume = True
    try:
        platform_eval('1 if VAR98999 == 0 else 2', {"VAR98999": "F"})
        f_equal_zero = False
    except ConditionalFailureError:
        f_equal_zero = True
    check('"F"', '"F"+1 e "F"==0 falham; "F"=="F" é a única leitura', f_consume and f_equal_zero and platform_eval('1 if VAR98999 == "F" else 2', {"VAR98999": "F"}) == 1)

    # ---------- agregações ----------
    days = [date(2026, 1, 1) + timedelta(days=i) for i in range(40)]
    run_date = days[-1]
    agg_ctx = runtime.context()
    daily_spec = {}
    for i, day in enumerate(days):
        lth = [v * (1 + 0.013 * ((i * 7 + k) % 11)) for k, v in enumerate([1000.0, 1100.0, 1200.0, 900.0, 1040.0, 1000.0, 1200.0])]
        states = {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal"}
        if i % 9 == 4:
            states.update({"L4": "LC"})
        if i % 13 == 6:
            states.update({"L6": "1 By pass"})
        inputs = scenario(lth, vazao={"L1_L3": 1320.0 + i % 3, "L4_L5": 1300.0, "L6_L7": 1250.0},
                          fator={"L1_L3": 1.1, "L4_L5": 1.05, "L6_L7": 1.15}, states=states)
        _, err = runtime.run(inputs, day=day, ctx=agg_ctx)
        assert err is None, err
        daily_spec[day] = spec(inputs)
    service = TemporalAggregationService()
    agg_rows = []
    for e in entities:
        if e["role"] != "aggregation":
            continue
        op, freq, source = e["agg"]
        by_source = defaultdict(list)
        for inst in concrete(e["scope_type"], e["scope_value"]):
            r = resolve(source, None, {"frequency": "diário"}, inst, model.by_name)
            (hit,) = r["hits"]
            by_source[hit["row"]].append(inst)
        materialized = []
        for src_row, insts in by_source.items():
            rule = AggregationRule(f"AGG98-{e['row']}-{src_row}", runtime.ids[src_row], "diário", runtime.ids[e["row"]], freq,
                                   {"Média": "AVERAGE", "Somatório": "SUM"}[op])
            for instance in ScopeResolver().resolve_aggregation_rule(rule, runtime.variables.get(runtime.ids[e["row"]]),
                                                                     runtime.variables.get(runtime.ids[src_row])):
                window = [d for d in days if (d.month == run_date.month if freq == "mensal" else True)]
                expected = sum(daily_spec[d][(source, instance.scope_value)] for d in window) / len(window)
                got = service.aggregate(rule, agg_ctx, instance.scope_type, instance.scope_value, run_date)
                materialized.append(instance.scope_value)
                agg_rows.append({
                    "entidade": e["name"], "frequencia": freq, "escopo_declarado": f"{e['scope_type']}/{e['scope_value']}",
                    "regra_workbook": op, "regra_aplicada": rule.aggregation_type, "instancia": instance.scope_value,
                    "dias_na_janela": len(window), "esperado_independente": expected, "runtime": got.value,
                    "unidade_origem": row_of[src_row]["unit"], "unidade_destino": e["unit"],
                    "status": "PASS" if close(expected, got.value) and rule.aggregation_type == "AVERAGE" and got.scope_value == instance.scope_value else "FAIL",
                })
        expected_count = len(concrete(e["scope_type"], e["scope_value"]))
        check("A021", f"{e['name']} {freq} {e['scope_type']}/{e['scope_value']}: {expected_count} instâncias esperadas",
              sorted(materialized) == sorted(i[1] for i in concrete(e["scope_type"], e["scope_value"])), str(materialized))
    write_csv("aggregation_matrix.csv", agg_rows)
    check("agregações", "todas as 22 instâncias = média independente dos diários (AVERAGE, sem fator)",
          len(agg_rows) == 22 and all(r["status"] == "PASS" for r in agg_rows))
    check("agregações", "total mensal/anual == Σ médias dos grupos (mesmos dias)",
          all(close(next(r["runtime"] for r in agg_rows if r["entidade"] == "retirada_condensado_total" and r["frequencia"] == f),
                    sum(r["runtime"] for r in agg_rows if r["entidade"] == "retirada_condensado_grupo" and r["frequencia"] == f))
              for f in ("mensal", "anual")))
    check("/h gate", "A41: nenhuma SUM; unidade origem == destino em todas as agregações",
          all(r["regra_workbook"] == "Média" and r["unidade_origem"] == r["unidade_destino"] for r in agg_rows))
    # "F" através da agregação
    f_ctx = runtime.context()
    for i, day in enumerate([date(2026, 3, 1), date(2026, 3, 2), date(2026, 3, 3)]):
        inputs = dict(VALUE_SETS["diferentes"])
        if i == 1:
            inputs.update({("hes_l4", "L4"): "Normal", ("hes_l5", "L5"): "1 By pass e LC"})
        runtime.run(inputs, day=day, ctx=f_ctx)
    grp = next(e for e in entities if e["name"] == "retirada_condensado_grupo" and e["frequency"] == "mensal" and e["scope_value"] == "L4_L5")
    src = next(e for e in entities if e["name"] == "retirada_condensado_grupo" and e["frequency"] == "diário" and e["scope_value"] == "L4_L5")
    try:
        service.aggregate(AggregationRule("AGG98-F", runtime.ids[src["row"]], "diário", runtime.ids[grp["row"]], "mensal", "AVERAGE"),
                          f_ctx, "linha_grupo", "L4_L5", date(2026, 3, 3))
        f_agg = None
    except AggregationFailureError as exc:
        f_agg = exc.failed_period_ids
    check('"F"', "média mensal com um dia F falha com AggregationFailureError(['2026-03-02'])", f_agg == ["2026-03-02"], str(f_agg))

    # ---------- ordem ----------
    order_rows = []
    base_inputs = VALUE_SETS["diferentes"]
    ref_ctx, _ = runtime.run(base_inputs)
    reference = runtime.outputs(ref_ctx)
    orders = {"original": runtime.definitions, "reversa": list(reversed(runtime.definitions))}
    for seed in (7, 41, 2026):
        shuffled = list(runtime.definitions)
        random.Random(seed).shuffle(shuffled)
        orders[f"aleatória seed={seed}"] = shuffled
    for label, order_defs in orders.items():
        for scen, inputs in (("valores diferentes", base_inputs), ("falha F L4_L5", f_inputs)):
            ctx, err = runtime.run(inputs, order=order_defs)
            out = runtime.outputs(ctx)
            ref = reference if scen == "valores diferentes" else runtime.outputs(runtime.run(f_inputs)[0])
            same = out.keys() == ref.keys() and all(close(out[k], ref[k]) for k in out)
            order_rows.append({"ordem": label, "cenario": scen, "valores_calculados": len(out),
                               "erro": type(err.original_error).__name__ if err else "", "igual_a_original": "sim" if same else "NÃO"})
    write_csv("order_invariance_results.csv", order_rows)
    check("ordem", "resultados idênticos em 5 ordens × 2 cenários", all(r["igual_a_original"] == "sim" for r in order_rows))

    # ---------- perturbação / isolamento ----------
    pert_rows = []

    def perturb(label, base_in, key, new_value):
        ctx0, e0 = runtime.run(base_in)
        out0 = runtime.outputs(ctx0)
        changed_in = dict(base_in)
        changed_in[key] = new_value
        ctx1, e1 = runtime.run(changed_in)
        out1 = runtime.outputs(ctx1)
        observed = {k for k in out0 if not close(out0[k], out1.get(k, float("nan")))}
        input_entity = next(e for e in entities if e["name"] == key[0] and key[1] in [i[1] for i in concrete(e["scope_type"], e["scope_value"])])
        inst = next(i for i in concrete(input_entity["scope_type"], input_entity["scope_value"]) if i[1] == key[1])
        expected = model.descendants((input_entity["row"], inst))
        ok = observed == expected
        pert_rows.append({"perturbacao": label, "entrada": f"{key[0]}@{key[1]}", "de": base_in[key], "para": new_value,
                          "afetados_esperados_DAG": " ".join(sorted(f"{a}@{b}" for a, b in expected)),
                          "afetados_observados": " ".join(sorted(f"{a}@{b}" for a, b in observed)),
                          "status": "PASS" if ok else "FAIL"})
        PERTURBED[key] = observed
        return ok

    PERTURBED = {}
    b = VALUE_SETS["diferentes"]
    bypass45 = dict(b); bypass45.update({("hes_l4", "L4"): "1 By pass"})
    bypasslc45 = dict(b); bypasslc45.update({("hes_l4", "L4"): "1 By pass e LC", ("hes_l5", "L5"): "1 By pass e LC"})
    bypass67 = dict(b); bypass67.update({("hes_l6", "L6"): "1 By pass"})
    tests = [("lth@L1 +10%", b, ("lth", "L1"), 1100.0), ("lth@L4 +10%", b, ("lth", "L4"), 990.0),
             ("lth@L7 +10%", b, ("lth", "L7"), 1320.0), ("valor_retirada@L2 +1", b, ("valor_retirada", "L2"), 48.0),
             ("vazao_ltp@L6_L7", b, ("vazao_ltp", "L6_L7"), 1400.0), ("fator@L4_L5", b, ("fator_retirada_cond_corr_lth", "L4_L5"), 1.2),
             ("hes_l4 Normal->LC", b, ("hes_l4", "L4"), "LC"), ("hes_l7 Normal->Overhaul/Parada", b, ("hes_l7", "L7"), "Overhaul/Parada"),
             ("41d@L4_L5 (ramo 1 By pass ativo)", bypass45, ("desconto_retirada_41d", "L4_L5"), 9.0),
             ("41c@L4_L5 (ramo 1 By pass e LC ativo)", bypasslc45, ("desconto_retirada_41c", "L4_L5"), 9.0),
             ("41c@L6_L7 (ramo 1 By pass ativo)", bypass67, ("desconto_retirada_41c", "L6_L7"), 9.0)]
    pert_ok = all(perturb(*t) for t in tests)
    # desconto sem ramo ativo não altera nada (efeito nulo esperado)
    ctx0, _ = runtime.run(b)
    changed = dict(b); changed[("desconto_retirada_41c", "L4_L5")] = 99.0
    ctx1, _ = runtime.run(changed)
    inactive_ok = runtime.outputs(ctx0) == runtime.outputs(ctx1)
    pert_rows.append({"perturbacao": "41c@L4_L5 com estados Normal (ramo inativo)", "entrada": "desconto_retirada_41c@L4_L5", "de": 5.0, "para": 99.0,
                      "afetados_esperados_DAG": "(nenhum: ramo inativo)", "afetados_observados": "" if inactive_ok else "ALTERADOS", "status": "PASS" if inactive_ok else "FAIL"})
    write_csv("perturbation_results.csv", pert_rows)
    check("perturbação", "11 perturbações: afetados observados == descendentes do DAG independente", pert_ok)
    check("perturbação", "desconto com ramo inativo não altera resultados", inactive_ok)
    iso_ok = True
    for line in ("L1", "L4", "L7"):
        group = SMALL_GROUP[line]
        allowed = set(GROUPS[group]) | {group, "L1_L7"}
        iso_ok &= all(sv in allowed for _, sv in PERTURBED[("lth", line)])
    check("isolamento", "perturbar lth de uma linha só altera o próprio grupo, suas linhas e o total", iso_ok)
    evidence["failfast_not_computed"] = sorted(FAILFAST)

    # denominador zero (rateio) e domínio de ln
    zero = scenario([1000.0, 1100.0, 1200.0, 0.0, 0.0, 1000.0, 1200.0])
    _, err = runtime.run(zero)
    evidence["denominador_zero_L4_L5"] = type(err.original_error).__name__ if err else None
    check("rateio", "Σ lth = 0 no grupo: falha explícita (nunca valor silencioso)", err is not None and isinstance(err.original_error, (MathDomainError, DivisionByZeroError)),
          evidence["denominador_zero_L4_L5"])

    # ---------- unidades ----------
    def dims(unit):
        if unit in (None, "-"):
            return {}
        num, _, den = unit.partition("/")
        d = {num: 1}
        if den:
            d[den] = d.get(den, 0) - 1
        return {k: v for k, v in d.items() if v}

    dim_rows = []
    for e in entities:
        if e["role"] != "equation":
            continue
        inst = concrete(e["scope_type"], e["scope_value"])[0]
        notes = []

        def unit_of(node):
            if isinstance(node, ast.Constant):
                return None
            if isinstance(node, ast.Name):
                name, _, suffix = node.id.partition("__AT__")
                r = model.resolution[(e["row"], inst, name, suffix or None)]
                return dims(r["hits"][0]["unit"])
            if isinstance(node, ast.BinOp):
                a, b = unit_of(node.left), unit_of(node.right)
                if isinstance(node.op, (ast.Add, ast.Sub)):
                    if a is not None and b is not None and a != b:
                        notes.append(f"INCOMPATÍVEL {a} ± {b}")
                    if (a is None) != (b is None):
                        notes.append("constante somada a grandeza (unidade implícita da constante)")
                    return a if a is not None else b
                if a is None or b is None:
                    return a if b is None else b
                out = dict(a)
                for k, v in b.items():
                    out[k] = out.get(k, 0) + (v if isinstance(node.op, ast.Mult) else -v)
                return {k: v for k, v in out.items() if v}
            if isinstance(node, ast.UnaryOp):
                return unit_of(node.operand)
            if isinstance(node, ast.Call):
                arg = unit_of(node.args[0])
                if arg:
                    notes.append(f"ln de grandeza dimensional {arg} (coeficientes de regressão em m³/h)")
                return None
            if isinstance(node, ast.IfExp):
                a, b = unit_of(node.body), unit_of(node.orelse)
                return a if a is not None else b
            return None

        result = unit_of(ast.parse(encode(e["expression"]), mode="eval").body)
        target = dims(e["unit"])
        compatible = result is None or result == target
        dim_rows.append({"entidade": f"{e['name']}@{e['scope_value']}", "unidade_declarada": e["unit"],
                         "unidade_calculada": "(adimensional/constante)" if result is None else json.dumps(result, ensure_ascii=False),
                         "observacoes": "; ".join(dict.fromkeys(notes)), "compativel": "sim" if compatible and not any("INCOMPAT" in n for n in notes) else "NÃO"})
    for r in agg_rows:
        dim_rows.append({"entidade": f"{r['entidade']}@{r['instancia']} ({r['frequencia']})", "unidade_declarada": r["unidade_destino"],
                         "unidade_calculada": r["unidade_origem"], "observacoes": "AVERAGE de taxa preserva a unidade",
                         "compativel": "sim" if r["unidade_origem"] == r["unidade_destino"] else "NÃO"})
    write_csv("dimensional_validation.csv", dim_rows)
    check("unidades", "nenhuma soma incompatível; unidade calculada == declarada", all(r["compativel"] == "sim" for r in dim_rows))

    # ---------- Forecast A41 ----------
    fc_rows = []
    param_row = {}
    for e in entities:
        m = re.fullmatch(r"retirada_meta_(41[a-z])", e["name"])
        if m:
            param_row[int(re.search(r"\d+$", e["source_reference"]).group(0))] = m.group(1)
    decision = {("desconto_retirada_41c", "L4_L5"): "Forecast A41!C26", ("desconto_retirada_41c", "L6_L7"): "Forecast A41!C28",
                ("desconto_retirada_41d", "L4_L5"): "Forecast A41!C27"}
    for e in entities:
        refs = e["source_reference"]
        corroboration = ""
        if (e["name"], e["scope_value"]) in decision:
            row = int(re.search(r"(\d+)$", refs).group(1))
            corroboration = (f"decisão do usuário: {decision[(e['name'], e['scope_value'])]} ({'igual' if refs == decision[(e['name'], e['scope_value'])] else 'DIFERENTE'}); "
                             f"linha {row} da aba corresponde, pelos parâmetros do próprio workbook, a {param_row.get(row, 'NENHUMA condição declarada')}")
        shared = [f"{x['name']}@{x['scope_value']}" for x in entities if x is not e and x["source_reference"] == refs]
        fc_rows.append({"workbook_linha": e["row"], "celula_forecast": refs, "entidade": f"{e['name']}@{e['scope_value']} ({e['frequency'] or 'parâmetro'})",
                        "interpretacao": e["description"], "valor_informado_no_workbook": e["value"] if e["value"] is not None else (e["OBS"] or "").split(".")[0][:80],
                        "celula_compartilhada_com": " ".join(shared), "corroboracao_interna": corroboration,
                        "resultado": "fonte Forecast A41 indisponível", "status": "NOT VERIFIED"})
    write_csv("forecast_reconciliation.csv", fc_rows)
    evidence["forecast_param_rows"] = param_row
    for key, cell in decision.items():
        e = next(x for x in entities if (x["name"], x["scope_value"]) == key)
        check("41c/41d", f"{key[0]}@{key[1]} = {cell} (decisão do usuário)", e["source_reference"] == cell)
    discounts = [e for e in entities if e["name"].startswith("desconto_retirada_41")]
    check("41c/41d", "exatamente 3 descontos; sem 41d@L6_L7; nenhuma célula compartilhada",
          len(discounts) == 3 and not any(e["name"] == "desconto_retirada_41d" and e["scope_value"] == "L6_L7" for e in discounts)
          and len({e["source_reference"] for e in discounts}) == 3)

    # ---------- max_ht (somente identificação) ----------
    from app.repositories.seed_loader import SeedLoader
    from app.validation.aggregation_dimension_validator import find_sum_dimension_issues
    loader = SeedLoader(ROOT / "data" / "seed")
    rules = loader.load_aggregation_rules()
    issues = find_sum_dimension_issues(rules.all(), loader.load_variable_definitions())
    evidence["max_ht"] = {"affected_rules": sorted(i.aggregation_rule_id for i in issues if i.aggregation_rule_id.startswith("AGR-MAX_HT-")),
                          "count": sum(i.aggregation_rule_id.startswith("AGR-MAX_HT-") for i in issues),
                          "integration_factors_in_seeds": sorted({r.integration_factor for r in rules})}
    check("max_ht", "identificação: 34 SUMs /h -> /mês|/ano; nenhum fator aplicado nos seeds",
          evidence["max_ht"]["count"] == 34 and evidence["max_ht"]["integration_factors_in_seeds"] == [1.0])

    # ---------- saída ----------
    evidence["checks"] = CHECKS
    if mutation_run:
        failed = [c for c in CHECKS if c["result"] != "PASS"]
        print(f"MUTATION RUN: {len(failed)} FAIL")
        for c in failed:
            print(f"  FAIL [{c['gate']}] {c['check']}")
        return 1 if failed else 0
    (EVIDENCE / "checks.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8")
    failed = [c for c in CHECKS if c["result"] != "PASS"]
    for c in CHECKS:
        print(f"{c['result']:4} [{c['gate']}] {c['check']}" + (f"  -- {c['detail']}" if c["result"] != "PASS" and c["detail"] else ""))
    print(f"\n{len(CHECKS) - len(failed)} PASS / {len(failed)} FAIL")
    return 1 if failed else 0


def run() -> int:
    """Uma exceção inesperada nunca vira PASS: a auditoria aborta como FAIL."""

    try:
        return main()
    except Exception as exc:  # noqa: BLE001
        failed = [c for c in CHECKS if c["result"] != "PASS"]
        print(f"{'MUTATION RUN' if MUTATION else 'AUDIT'} ABORTED (FAIL): {type(exc).__name__}: {exc}")
        for c in failed:
            print(f"  FAIL [{c['gate']}] {c['check']}")
        return 2


if __name__ == "__main__":
    raise SystemExit(run())
