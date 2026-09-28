"""
Etapa 2.6 — contrato interbloco declarado em `fonte` (evidência).

    python audit/stage2_6_interblock_contract/evidence/analysis_stage2_6.py

Duas leituras independentes dos cinco workbooks oficiais
(data/workbooks, SHA-256 conferido contra tools/workbook_seed/blocks.py):

  A. leitura crua com openpyxl, sem o reader/modelo canônico do
     projeto: cada linha com `fonte` é classificada por regras
     reescritas aqui (bloco carregado?, contexto, produtor por
     name + frequency + scope_type + instância, contrato);
  B. o builder oficial (`tools.workbook_seed.blocks.build_all`).

As duas classificações têm de coincidir linha a linha; qualquer
divergência é registrada e reprova a evidência. Nada em data/, app/,
tools/ ou tests/ é escrito.

Saídas: interblock_links.csv, contract_validation.csv,
interblock_graph.csv (raiz da etapa) e evidence/workbook_sha256.txt,
evidence/validation_results.json.
"""

from __future__ import annotations

import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
REPO = HERE.parents[2]
sys.path.insert(0, str(REPO))

import openpyxl  # noqa: E402

from app.validation.variable_seed_validator import VARIABLE_ID_RANGES  # noqa: E402
from tools.workbook_seed.blocks import BLOCKS, build_all  # noqa: E402
from tools.workbook_seed.reader import sha256_of  # noqa: E402

LINES = [f"L{i}" for i in range(1, 8)]
DIMENSIONS = ("kind", "unit", "value_type", "allowed_values", "declared_result_states", "instances")


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def instances(scope_type, scope_value):
    if scope_type == "linha" and "_" in str(scope_value):
        a, b = scope_value.split("_")
        return {(scope_type, l) for l in LINES[LINES.index(a): LINES.index(b) + 1]}
    return {(scope_type, scope_value)}


# ------------------------------------------------------------------
# 0. integridade
# ------------------------------------------------------------------
sha_lines = []
for block, spec in sorted(BLOCKS.items()):
    digest = sha256_of(spec.workbook_path)
    assert digest == spec.sha256, (block, digest)
    sha_lines.append(f"{digest}  data/workbooks/{spec.file_name}  ({block} {spec.version})")
(HERE / "workbook_sha256.txt").write_text("\n".join(sha_lines) + "\n", encoding="utf-8")


# ------------------------------------------------------------------
# A. leitura crua
# ------------------------------------------------------------------
def raw_rows(spec):
    ws = openpyxl.load_workbook(spec.workbook_path, data_only=False)[spec.sheet]
    header_row, header = None, None
    for r in range(1, 11):
        values = [c.value for c in ws[r]]
        if "name" in values and "fonte" in values:
            header_row, header = r, values
            break
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        values = dict(zip(header, [c.value for c in ws[r]]))
        if values.get("name") is None:
            continue
        values["_row"] = r
        values["_kind"] = "parameter" if values["Type"] == "parameter" else "variable"
        rows.append(values)
    return rows


RAW = {block: raw_rows(spec) for block, spec in BLOCKS.items()}


def raw_definitions(block):
    """Agrupa linhas atômicas por linha como o contrato de 2.3 define."""
    groups = defaultdict(list)
    for r in RAW[block]:
        sv = r["scope_value"]
        per_line = r["scope_type"] == "linha" and sv in LINES
        key = (r["_kind"], r["name"], r["frequency"], r["scope_type"], "*" if per_line else sv)
        groups[key].append(r)
    defs = []
    for key, rows in groups.items():
        inst = set()
        for r in rows:
            inst |= instances(r["scope_type"], r["scope_value"])
        defs.append({"key": key, "rows": rows, "instances": inst, "first": rows[0]})
    return defs


DEFS = {block: raw_definitions(block) for block in BLOCKS}


def classify_raw(block, row):
    """Classificação independente de uma linha com fonte."""
    source = row["fonte"]
    findings = []
    if row["_kind"] != "variable" or row["expression"] is not None:
        findings.append(("INTERBLOCK_CONTRACT_MISMATCH", "CONTRACT_MISMATCH", "consumer_context"))
    if source not in BLOCKS:
        findings.append(("INTERBLOCK_SOURCE_NOT_FOUND", "SOURCE_NOT_FOUND", "source_block"))
        return None, findings
    wanted = instances(row["scope_type"], row["scope_value"])
    consumer_def = next(
        d for d in DEFS[block] if row in d["rows"]
    )
    wanted = consumer_def["instances"]
    candidates = [
        d for d in DEFS[source]
        if (d["key"][1], d["key"][2], d["key"][3]) == (row["name"], row["frequency"], row["scope_type"])
    ]
    covering = [d for d in candidates if wanted <= d["instances"]]
    if len(covering) != 1:
        if not [d for d in DEFS[source] if d["key"][1] == row["name"]]:
            findings.append(("INTERBLOCK_SOURCE_NOT_FOUND", "SOURCE_NOT_FOUND", "name"))
        elif len(covering) > 1:
            findings.append(("INTERBLOCK_SOURCE_AMBIGUOUS", "AMBIGUOUS", "identity"))
        elif candidates:
            findings.append(("INTERBLOCK_CONTRACT_MISMATCH", "INSTANCE_MISMATCH", "instances"))
        else:
            findings.append(("INTERBLOCK_SOURCE_NOT_FOUND", "SOURCE_NOT_FOUND", "frequency/scope"))
        return None, findings
    producer = covering[0]
    p = producer["first"]
    checks = {
        "kind": (row["_kind"], p["_kind"]),
        "unit": (row["unit"], p["unit"]),
        "value_type": (row["value_type"], p["value_type"]),
        "allowed_values": (row["allowed_values"], p["allowed_values"]),
        "declared_result_states": (row["declared_result_states"], p["declared_result_states"]),
    }
    for dimension, (mine, theirs) in checks.items():
        if mine != theirs:
            findings.append(("INTERBLOCK_CONTRACT_MISMATCH", "CONTRACT_MISMATCH", dimension))
    return producer, findings


# ------------------------------------------------------------------
# B. builder oficial
# ------------------------------------------------------------------
built = build_all()
result = built.interblock
by_row = {}
for link in result.links:
    for r in link.consumer_rows:
        by_row[(link.consumer_block, r)] = link

link_rows, contract_rows, divergences = [], [], []
raw_count = 0

for block in sorted(BLOCKS):
    for row in RAW[block]:
        if row["fonte"] is None:
            continue
        raw_count += 1
        producer, raw_findings = classify_raw(block, row)
        raw_status = raw_findings[0][1] if raw_findings else "VALID"
        link = by_row.get((block, row["_row"]))

        if link is None:
            divergences.append({"block": block, "row": row["_row"], "issue": "linha com fonte sem vínculo no builder"})
            continue

        impl = [(f.code, f.link_class, f.dimension) for f in link.findings]
        impl_dims = sorted(d for _c, _k, d in impl)
        raw_dims = sorted(d if d != "frequency/scope" else d for _c, _k, d in raw_findings)
        if raw_status != link.validation_status or (
            [c for c, _k, _d in raw_findings] != [c for c, _k, _d in impl]
        ):
            divergences.append({
                "block": block, "row": row["_row"],
                "issue": f"raw={raw_findings} builder={impl}",
            })

        p = link.producer
        consumer_scope = f"{row['scope_type']}/{row['scope_value']}"
        link_rows.append({
            "consumer_block": block,
            "consumer_row": row["_row"],
            "consumer_name": row["name"],
            "consumer_unit": row["unit"],
            "consumer_frequency": row["frequency"],
            "consumer_scope": consumer_scope,
            "source_block": row["fonte"],
            "producer_name": p.name if p else "",
            "producer_unit": p.unit if p else "",
            "producer_frequency": p.frequency if p else "",
            "producer_scope": f"{p.scope_type}/{p.scope_value}" if p else "",
            "resolution_status": link.resolution_status,
            "validation_status": link.validation_status,
            "consumer_definition": link.consumer.entity_id,
            "producer_definition": p.entity_id if p else "",
            "producer_kind": p.kind if p else "",
            "consumer_variable_type": row["variable_type"],
            "source_in_official_taxonomy": row["fonte"] in VARIABLE_ID_RANGES,
            "source_block_loaded": row["fonte"] in BLOCKS,
            "chain": " → ".join(link.chain),
            "chain_status": link.chain_status,
            "error_codes": ";".join(f.code for f in link.findings),
            "divergent_dimensions": ";".join(f.dimension for f in link.findings),
            "independent_check": "AGREE" if not divergences or divergences[-1]["row"] != row["_row"] or divergences[-1]["block"] != block else "DIVERGE",
        })

        # contrato, dimensão a dimensão
        consumer = link.consumer
        by_dim = {f.dimension: f for f in link.findings}
        for f in link.findings:
            if f.dimension not in DIMENSIONS:
                contract_rows.append({
                    "consumer_block": block, "consumer_row": row["_row"], "consumer_name": row["name"],
                    "source_block": row["fonte"], "dimension": f.dimension,
                    "consumer_value": f.consumer_value, "producer_value": f.producer_value,
                    "result": f.link_class, "error_code": f.code, "message": f.message,
                })
        if p is None:
            continue
        values = {
            "kind": (consumer.kind, p.kind),
            "unit": (consumer.unit, p.unit),
            "value_type": (consumer.value_type, p.value_type),
            "allowed_values": (consumer.allowed_values, p.allowed_values),
            "declared_result_states": (consumer.declared_result_states, p.declared_result_states),
            "instances": (
                sorted(v for _t, v in link.instances()),
                sorted(v for _t, v in instances(p.scope_type, p.scope_value)),
            ),
        }
        for dimension in DIMENSIONS:
            mine, theirs = values[dimension]
            f = by_dim.get(dimension)
            contract_rows.append({
                "consumer_block": block, "consumer_row": row["_row"], "consumer_name": row["name"],
                "source_block": row["fonte"], "dimension": dimension,
                "consumer_value": mine, "producer_value": theirs,
                "result": f.link_class if f else "OK",
                "error_code": f.code if f else "",
                "message": f.message if f else "",
            })

write_csv(ROOT / "interblock_links.csv", link_rows)
write_csv(ROOT / "contract_validation.csv", contract_rows)


# ------------------------------------------------------------------
# grafo
# ------------------------------------------------------------------
graph_rows = []
block_edges = Counter()
block_valid = Counter()
for link in result.links:
    block_edges[(link.consumer_block, link.source_block)] += 1
    block_valid[(link.consumer_block, link.source_block)] += link.is_valid
    graph_rows.append({
        "level": "definition",
        "consumer": f"{link.consumer_block}.{link.consumer.name}[{link.consumer.frequency} {link.consumer.scope_type}/{link.consumer.scope_value}]",
        "consumer_definition": link.consumer.entity_id,
        "producer": (
            f"{link.source_block}.{link.producer.name}[{link.producer.frequency} {link.producer.scope_type}/{link.producer.scope_value}]"
            if link.producer else f"{link.source_block}.?"
        ),
        "producer_definition": link.producer.entity_id if link.producer else "",
        "edges": 1,
        "valid_edges": int(link.is_valid),
        "status": link.validation_status,
        "note": " → ".join(link.chain) + f" ({link.chain_status})",
    })
mutual = {(a, b) for (a, b) in block_edges if (b, a) in block_edges}
for (a, b), n in sorted(block_edges.items()):
    graph_rows.append({
        "level": "block",
        "consumer": a,
        "consumer_definition": "",
        "producer": b,
        "producer_definition": "",
        "edges": n,
        "valid_edges": block_valid[(a, b)],
        "status": "LOADED" if b in BLOCKS else "SOURCE_BLOCK_NOT_LOADED",
        "note": (
            "dependência mútua entre blocos; no nível de definição (vínculos + "
            "equações + agregações) não há ciclo"
            if (a, b) in mutual else ""
        ),
    })
write_csv(ROOT / "interblock_graph.csv", graph_rows)


# ------------------------------------------------------------------
# resumo
# ------------------------------------------------------------------
status_by_link = Counter(l.validation_status for l in result.links)
status_by_row = Counter(r["validation_status"] for r in link_rows)
findings_by_code = Counter(f.code for _l, f in result.findings)
findings_by_class = Counter(f.link_class for _l, f in result.findings)
summary = {
    "workbooks": {b: {"file": s.file_name, "version": s.version, "sha256": s.sha256} for b, s in sorted(BLOCKS.items())},
    "fonte_rows": raw_count,
    "declared_links_definitions": len(result.links),
    "valid_links_definitions": len(result.valid),
    "valid_links_rows": status_by_row["VALID"],
    "rejected_links_definitions": len(result.rejected),
    "rejected_links_rows": raw_count - status_by_row["VALID"],
    "links_by_class_definitions": dict(sorted(status_by_link.items())),
    "links_by_class_rows": dict(sorted(status_by_row.items())),
    "findings_total": len(result.findings),
    "findings_by_code": dict(sorted(findings_by_code.items())),
    "findings_by_class": dict(sorted(findings_by_class.items())),
    "links_to_unloaded_taxonomy_blocks": sorted(
        {l.source_block for l in result.links if l.source_block not in BLOCKS}
    ),
    "links_to_unloaded_blocks_definitions": sum(1 for l in result.links if l.source_block not in BLOCKS),
    "links_to_loaded_blocks_definitions": sum(1 for l in result.links if l.source_block in BLOCKS),
    "fonte_outside_taxonomy": sorted(
        {str(l.source_block) for l in result.links if l.source_block not in VARIABLE_ID_RANGES}
    ),
    "cycles": [list(map(list, c)) for c in result.cycles],
    "block_level_mutual_dependencies": sorted(sorted(p) for p in {frozenset(x) for x in mutual}),
    "independent_check_divergences": divergences,
    "interblock_seed_matches_builder": (
        (REPO / "data/seed/interblock_links.json").read_text(encoding="utf-8")
        == json.dumps(built.interblock_seed, ensure_ascii=False, indent=2) + "\n"
    ),
}
(HERE / "validation_results.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)
print(json.dumps({k: summary[k] for k in (
    "fonte_rows", "declared_links_definitions", "valid_links_definitions", "valid_links_rows",
    "links_by_class_definitions", "links_by_class_rows", "findings_by_class", "cycles",
    "independent_check_divergences", "interblock_seed_matches_builder",
)}, ensure_ascii=False, indent=1))
assert not divergences, divergences
