"""
Harness de evidência do workbook A41 v6 (fechamento da Etapa 2).

Derivado de audit/area_41_v5/check_a41_v5.py. Diferenças:
  - 41c/41d: verifica as três atribuições decididas pelo usuário
    (C26 -> 41c@L4_L5, C28 -> 41c@L6_L7, C27 -> 41d@L4_L5) em vez de
    inferir a condição pelas linhas dos parâmetros;
  - scope_matrix com consumer_frequency / resolved_entity /
    resolved_frequency / resolved_scope / status;
  - "F" através das agregações;
  - change_matrix.csv (v5 -> v6 recebido -> v6 validado).

NÃO é implementação do A41: não cria seed, builder, IDs produtivos nem
lógica específica no runtime. Lê o workbook e o submete apenas aos
componentes GENÉRICOS da plataforma (resolvedor central de referências,
ExpressionParser, DependencyExtractor, ForecastEngine, ScopeResolver,
TemporalAggregationService, validador dimensional).

Os IDs usados aqui (VAR99xxx / PARAM99xxx / EQ99xxx) são efêmeros, só
existem em memória durante a verificação e ficam fora de qualquer faixa
produtiva.

Os valores esperados das reconciliações são calculados por fórmulas de
referência independentes, escritas a partir das decisões vinculantes do
prompt da Etapa 2 (§5, §7, §10, §12) — nunca a partir da saída do
próprio engine.

Uso (reprodutível pelo auditor):

    python3 -m pytest audit/area_41_v6/check_a41_v6.py -v
    A41_WORKBOOK=<outro.xlsx> python3 -m pytest audit/area_41_v6/check_a41_v6.py
    python3 audit/area_41_v6/check_a41_v6.py   # grava evidence/*
"""

import hashlib
import json
import math
import os
import re
import sys
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from app.domain.equations.models import EquationDefinition  # noqa: E402
from app.domain.equations.registry import EquationDefinitionRegistry  # noqa: E402
from app.domain.forecast.aggregation import AggregationRule  # noqa: E402
from app.domain.values import VALUE_TYPES  # noqa: E402
from app.domain.variables.models import VariableDefinition  # noqa: E402
from app.domain.variables.registry import VariableDefinitionRegistry  # noqa: E402
from app.engine import reference_resolver  # noqa: E402
from app.engine.calculation_context import CalculationContext  # noqa: E402
from app.engine.dependency_extractor import DependencyExtractor  # noqa: E402
from app.engine.exceptions import (  # noqa: E402
    CalculationValueError,
    ConditionalFailureError,
    EquationEvaluationError,
    VariableNotFoundError,
)
from app.engine.expression_parser import ExpressionParser  # noqa: E402
from app.engine.forecast_engine import ForecastEngine  # noqa: E402
from app.engine.scope_resolver import ScopeResolver  # noqa: E402
from app.engine.spatial_candidate_resolver import (  # noqa: E402
    get_spatial_candidates,
)
from app.engine.temporal_aggregation_service import (  # noqa: E402
    TemporalAggregationService,
)
from app.validation import variable_seed_validator as vsv  # noqa: E402
from app.validation.aggregation_dimension_validator import (  # noqa: E402
    find_sum_dimension_issues,
)
from tools.max_ht_seed_builder import (  # noqa: E402
    DSL_OPERATION_TO_AGGREGATION_TYPE,
    DSL_PATTERN,
)

HERE = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = HERE / "descritivo_das_variáveis_A41_v6.xlsx"
WORKBOOK = Path(os.environ.get("A41_WORKBOOK", DEFAULT_WORKBOOK))
SHEET = "A41"

BASE_HEADERS = [
    "Type", "name", "description", "unit", "value", "version",
    "variable_type", "frequency", "scope_type", "scope_value",
    "source_reference", "status", "expression", "fonte", "OBS",
]
OPTIONAL_HEADERS = ["value_type"]

STATES = ("Normal", "LC", "Overhaul/Parada", "1 By pass", "1 By pass e LC")
FAILURE = "F"

SCOPE_RESOLVER = ScopeResolver()


# ============================================================
# Leitura
# ============================================================


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_workbook(path: Path = WORKBOOK) -> dict:
    workbook = openpyxl.load_workbook(path)
    sheet = workbook[SHEET]
    headers = [c.value for c in sheet[2]]
    while headers and headers[-1] is None:
        headers.pop()

    entities = []
    trailing = []
    formulas = []
    counters = {"variable": 0, "parameter": 0}

    for row in sheet.iter_rows(min_row=1):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append(cell.coordinate)

    for row in sheet.iter_rows(min_row=3):
        values = [c.value for c in row][: len(headers)]
        record = dict(zip(headers, values))
        if not record.get("name"):
            if any(v not in (None, "") for v in values):
                trailing.append(row[0].row)
            continue
        if trailing:
            raise AssertionError(f"linhas com conteúdo sem name: {trailing}")

        kind = "parameter" if record["Type"] == "parameter" else "variable"
        counters[kind] += 1
        prefix = "PARAM" if kind == "parameter" else "VAR"
        expression = record.get("expression")
        dsl = DSL_PATTERN.match(expression) if expression else None

        if kind == "parameter":
            role = "parameter"
        elif dsl:
            role = "aggregation"
        elif expression:
            role = "equation"
        else:
            role = "input"

        entities.append(
            {
                **record,
                "row": row[0].row,
                "kind": kind,
                "role": role,
                "entity_id": f"{prefix}99{counters[kind]:03d}",
                "dsl": (
                    (
                        DSL_OPERATION_TO_AGGREGATION_TYPE[dsl.group(1)],
                        dsl.group(2),
                        dsl.group(3),
                    )
                    if dsl
                    else None
                ),
            }
        )

    return {
        "path": path,
        "headers": headers,
        "entities": entities,
        "formulas": formulas,
        "max_row": sheet.max_row,
    }


def identity(entity) -> tuple:
    return (
        entity["name"],
        entity["frequency"],
        entity["scope_type"],
        entity["scope_value"],
    )


def find(entities, name, frequency, scope_value):
    matches = [
        e for e in entities
        if e["name"] == name
        and e["frequency"] == frequency
        and e["scope_value"] == scope_value
    ]
    assert len(matches) == 1, (name, frequency, scope_value, matches)
    return matches[0]


# ============================================================
# Referências
# ============================================================

_STRING = re.compile(r"(\"[^\"]*\"|'[^']*')")
_TOKEN = re.compile(r"(?<![\w@])([A-Za-z_][A-Za-z0-9_]*)(@L[1-7](?:_L[1-7])?)?")


def references(expression: str, index) -> list[tuple[str, str | None]]:
    """(nome, sufixo) de cada referência a entidade na expressão."""

    found = []
    for position, segment in enumerate(_STRING.split(expression)):
        if position % 2:
            continue
        for match in _TOKEN.finditer(segment):
            name, suffix = match.group(1), match.group(2)
            if name in index:
                found.append((name, suffix[1:] if suffix else None))
    return found


def translate(entity, index) -> str:
    return reference_resolver.translate_expression(
        entity["expression"],
        index,
        entity["frequency"],
        consumer_scope=(entity["scope_type"], entity["scope_value"]),
    )


def scope_matrix(model) -> list[dict]:
    """
    Uma linha por referência de cada equação: como o resolvedor central
    vincula o nome e se o runtime alcança o valor a partir de CADA
    instância concreta do consumidor.

    Referência explícita (@): o runtime lê exatamente aquele escopo.
    Referência implícita: o runtime procura a partir do escopo da
    instância consumidora para fora (linha -> grupo -> planta), nunca
    para dentro (get_spatial_candidates).
    """

    entities = model["entities"]
    index = reference_resolver.build_name_index(entities)
    by_id = {e["entity_id"]: e for e in entities}
    rows = []

    for consumer in entities:
        if consumer["role"] != "equation":
            continue
        consumer_scopes = SCOPE_RESOLVER.resolve_scopes(
            consumer["scope_type"], consumer["scope_value"]
        )
        seen = set()
        for name, suffix in references(consumer["expression"], index):
            if (name, suffix) in seen:
                continue
            seen.add((name, suffix))

            resolved_id = reference_resolver.resolve_reference(
                name,
                index,
                consumer["frequency"],
                consumer_scope=(consumer["scope_type"], consumer["scope_value"]),
                explicit_scope_value=suffix,
            )
            resolved = by_id[resolved_id]
            resolved_scopes = set(
                SCOPE_RESOLVER.resolve_scopes(
                    resolved["scope_type"], resolved["scope_value"]
                )
            ) if resolved["kind"] == "variable" else {
                (resolved["scope_type"], resolved["scope_value"])
            }

            if suffix:
                relationship = (
                    "explícita-grupo" if "_" in suffix else "explícita-linha"
                )
                wanted = {(("linha_grupo" if "_" in suffix else "linha"), suffix)}
                reachable = wanted <= resolved_scopes
                expected = f"{name}@{suffix}"
            else:
                reachable = all(
                    any(
                        candidate in resolved_scopes
                        for candidate in get_spatial_candidates(*instance)
                    )
                    for instance in consumer_scopes
                )
                inward = not reachable and any(
                    scope[0] == "linha" for scope in resolved_scopes
                ) and consumer["scope_type"] != "linha"
                relationship = (
                    "implícita-mesmo-escopo"
                    if (resolved["scope_type"], resolved["scope_value"])
                    == (consumer["scope_type"], consumer["scope_value"])
                    else (
                        "implícita-para-dentro (grupo lendo linha)"
                        if inward
                        else "implícita-para-fora"
                    )
                )
                expected = "valor alcançável a partir de cada instância"

            rows.append(
                {
                    "consumer": consumer["name"],
                    "consumer_row": consumer["row"],
                    "consumer_frequency": consumer["frequency"],
                    "consumer_scope": f"{consumer['scope_type']}/{consumer['scope_value']}",
                    "reference": name + (f"@{suffix}" if suffix else ""),
                    "resolved_entity": f"{resolved['name']} (linha {resolved['row']})",
                    "resolved_frequency": resolved["frequency"] or "parâmetro",
                    "resolved_scope": f"{resolved['scope_type']}/{resolved['scope_value']}",
                    "reference_scope": (
                        f"{resolved['scope_type']}/{resolved['scope_value']}"
                        f" ({resolved['frequency'] or 'parâmetro'}, linha "
                        f"{resolved['row']})"
                    ),
                    "relationship": relationship,
                    "expected_resolution": expected,
                    "status": "PASS" if reachable else "FAIL",
                    "result": "PASS" if reachable else "FAIL",
                }
            )
    return rows


# ============================================================
# Registries genéricos a partir do workbook
# ============================================================


def build_registries(model):
    entities = model["entities"]
    index = reference_resolver.build_name_index(entities)
    variables = VariableDefinitionRegistry()
    equations = EquationDefinitionRegistry()
    translated = {}

    for entity in entities:
        if entity["kind"] != "variable":
            continue
        variables.add(
            VariableDefinition(
                entity["entity_id"], entity["name"], entity["description"],
                entity["unit"], entity["variable_type"], entity["frequency"],
                entity["scope_type"], entity["scope_value"],
                entity["source_reference"], entity["status"],
                entity.get("value_type") or "numeric",
            )
        )

    for number, entity in enumerate(entities, start=1):
        if entity["role"] != "equation":
            continue
        expression = translate(entity, index)
        translated[entity["entity_id"]] = expression
        equations.add(
            EquationDefinition(
                f"EQ99{number:03d}", entity["entity_id"], 1,
                entity["scope_type"], entity["scope_value"], expression,
                entity["source_reference"], "PUBLISHED",
            )
        )

    return variables, equations, translated


# ============================================================
# Fórmulas de referência (independentes do engine)
# ============================================================


def ref_corr(x: float) -> float:
    return 19.475 * math.log(x) - 85.999


def ref_group_l1_l3(v1, v2, v3, ltp, lth_corr):
    return v1 + v2 + v3 - 3 * (ltp - lth_corr)


def ref_group_conditional(s_a, s_b, ltp, lth_corr, bypass_discount, bypass_lc_discount):
    base = 100 - ((ltp - lth_corr) * 2)
    if s_a == "Normal" and s_b == "Normal":
        return base
    if s_a == "LC" or s_b == "LC":
        return (base * 14 + (base / 2) * 10) / 24
    if s_a == "Overhaul/Parada" or s_b == "Overhaul/Parada":
        return 50 - ((ltp - lth_corr) * 2)
    if s_a == "1 By pass" or s_b == "1 By pass":
        return base - bypass_discount
    if s_a == "1 By pass e LC" and s_b == "1 By pass e LC":
        return (base * 14 + ((base - bypass_lc_discount) / 2) * 10) / 24
    return FAILURE


# ============================================================
# Cenário determinístico
# ============================================================

INPUTS = {
    "vazao_ltp": {"L1_L3": 1320.0, "L4_L5": 1300.0, "L6_L7": 1250.0},
    "fator_retirada_cond_corr_lth": {"L1_L3": 1.1, "L4_L5": 1.05, "L6_L7": 1.15},
    "valor_retirada": {"L1": 47.0, "L2": 47.0, "L3": 48.0},
    "desconto_retirada_41c": {"L4_L5": 3.0, "L6_L7": 4.0},
    "desconto_retirada_41d": {"L4_L5": 7.0},
}
# §21: L1=10, L2=20, L3=30 -> média 20 (não 60).
LTH = {"L1": 10.0, "L2": 20.0, "L3": 30.0, "L4": 40.0, "L5": 60.0, "L6": 70.0, "L7": 30.0}


def new_context(variables) -> CalculationContext:
    context = CalculationContext()
    context.declare_categorical_variables(
        d.variable_definition_id for d in variables.all() if d.is_categorical
    )
    return context


def set_inputs(model, context, day: date, lth=None, states=None):
    entities = model["entities"]
    lth = lth or LTH
    states = states or {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal"}
    daily = day.isoformat()
    monthly = daily[:7]

    for name, per_scope in INPUTS.items():
        for scope_value, value in per_scope.items():
            entity = next(
                e for e in entities
                if e["name"] == name and e["scope_value"] == scope_value
            )
            context.set_variable_value(
                entity["entity_id"], value, entity["scope_type"], scope_value,
                period_id=monthly if entity["frequency"] == "mensal" else daily,
            )

    lth_entity = find(entities, "lth", "diário", "L1_L7")
    for line, value in lth.items():
        context.set_variable_value(
            lth_entity["entity_id"], value, "linha", line, period_id=daily
        )

    for line, state in states.items():
        entity = find(entities, f"hes_{line.lower()}", "diário", line)
        context.set_variable_value(
            entity["entity_id"], state, "linha", line, period_id=daily
        )


def run_day(model, registries, context, day, **inputs):
    variables, equations, _ = registries
    set_inputs(model, context, day, **inputs)
    return ForecastEngine().calculate_from_definition_registry(
        equation_definition_registry=equations,
        calculation_context=context,
        variable_definition_registry=variables,
        run_date=day,
    )


def value_of(model, context, name, scope_value, day, frequency="diário"):
    entity = find(model["entities"], name, frequency, scope_value)
    concrete = scope_value
    return context.get_variable_value(
        entity["entity_id"],
        "linha" if entity["scope_type"] == "linha" else entity["scope_type"],
        concrete,
        period_id=day.isoformat(),
    )


def expected_day(lth=None, states=None):
    lth = lth or LTH
    states = states or {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal"}
    lth_group = {
        "L1_L3": (lth["L1"] + lth["L2"] + lth["L3"]) / 3,
        "L4_L5": (lth["L4"] + lth["L5"]) / 2,
        "L6_L7": (lth["L6"] + lth["L7"]) / 2,
    }
    ltp = {g: ref_corr(v) for g, v in INPUTS["vazao_ltp"].items()}
    corr = {
        g: ref_corr(lth_group[g] * INPUTS["fator_retirada_cond_corr_lth"][g])
        for g in lth_group
    }
    group = {
        "L1_L3": ref_group_l1_l3(
            *(INPUTS["valor_retirada"][l] for l in ("L1", "L2", "L3")),
            ltp["L1_L3"], corr["L1_L3"],
        ),
        # Referência = workbook tal como está (41c/41d: ver BLOCKING).
        "L4_L5": ref_group_conditional(
            states["L4"], states["L5"], ltp["L4_L5"], corr["L4_L5"],
            INPUTS["desconto_retirada_41d"]["L4_L5"],
            INPUTS["desconto_retirada_41c"]["L4_L5"],
        ),
        "L6_L7": ref_group_conditional(
            states["L6"], states["L7"], ltp["L6_L7"], corr["L6_L7"],
            INPUTS["desconto_retirada_41c"]["L6_L7"],
            INPUTS["desconto_retirada_41c"]["L6_L7"],
        ),
    }
    members = {g: sorted(SCOPE_RESOLVER.GROUP_MEMBERS[g]) for g in ("L1_L3", "L4_L5", "L6_L7")}
    line = {}
    for g, lines in members.items():
        if group[g] == FAILURE:
            continue
        denominator = sum(lth[l] for l in lines)
        for l in lines:
            line[l] = group[g] * lth[l] / denominator
    total = (
        None if FAILURE in group.values()
        else group["L1_L3"] + group["L4_L5"] + group["L6_L7"]
    )
    return {"lth_grupo": lth_group, "ltp": ltp, "corr": corr, "group": group, "line": line, "total": total}


# ============================================================
# Agregações (materialização genérica por escopo concreto)
# ============================================================


def build_aggregation_instances(model, variables):
    """
    Para cada linha de agregação do workbook: resolve a origem por
    escopo concreto do alvo (resolvedor central, consumer_scope =
    instância concreta) e gera uma AggregationRule por origem distinta;
    o ScopeResolver materializa as instâncias. Nenhum tratamento por
    nome de variável.
    """

    entities = model["entities"]
    index = reference_resolver.build_name_index(entities)
    result = []

    for target in entities:
        if target["role"] != "aggregation":
            continue
        aggregation_type, target_frequency, source_name = target["dsl"]
        assert target_frequency == target["frequency"], target["row"]

        by_source = defaultdict(list)
        for scope in SCOPE_RESOLVER.resolve_scopes(
            target["scope_type"], target["scope_value"]
        ):
            source_id = reference_resolver.resolve_reference(
                source_name, index, "diário", consumer_scope=scope
            )
            by_source[source_id].append(scope)

        for source_id, scopes in by_source.items():
            rule = AggregationRule(
                aggregation_rule_id=f"AGG99-{target['row']}-{source_id}",
                source_variable_id=source_id,
                source_frequency="diário",
                target_variable_id=target["entity_id"],
                target_frequency=target_frequency,
                aggregation_type=aggregation_type,
            )
            instances = SCOPE_RESOLVER.resolve_aggregation_rule(
                rule,
                target_definition=variables.get(target["entity_id"]),
                source_definition=variables.get(source_id),
            )
            assert [(i.scope_type, i.scope_value) for i in instances] == scopes
            result.append((target, rule, instances))

    return result


# ============================================================
# Fixtures
# ============================================================


@pytest.fixture(scope="module")
def model():
    return load_workbook(WORKBOOK)


@pytest.fixture(scope="module")
def registries(model):
    return build_registries(model)


# ============================================================
# E1/E2/E3 — identidade, integridade, descrições
# ============================================================


def test_workbook_identity(model):
    print(f"\nWORKBOOK {WORKBOOK}\nsha256 {sha256(WORKBOOK)}")
    assert model["headers"][: len(BASE_HEADERS)] == BASE_HEADERS
    assert set(model["headers"][len(BASE_HEADERS):]) <= set(OPTIONAL_HEADERS)
    assert model["formulas"] == []


def test_entity_counts(model):
    roles = defaultdict(int)
    for entity in model["entities"]:
        roles[entity["role"]] += 1
    print("\nentities", len(model["entities"]), dict(roles))
    assert len(model["entities"]) == 54


def test_identities_are_unique(model):
    identities = [identity(e) for e in model["entities"]]
    assert len(identities) == len(set(identities))


def test_every_entity_has_a_description(model):
    missing = [
        e["row"] for e in model["entities"]
        if not isinstance(e["description"], str) or not e["description"].strip()
    ]
    assert missing == []


def test_descriptions_match_frequency(model):
    wrong = []
    for e in model["entities"]:
        text = e["description"].lower()
        if e["frequency"] == "mensal" and e["role"] == "aggregation" and "mensal" not in text:
            wrong.append(e["row"])
        if e["frequency"] == "anual" and "anual" not in text:
            wrong.append(e["row"])
    assert wrong == []


def test_field_enums_and_scopes(model):
    errors = []
    for e in model["entities"]:
        if e["unit"] not in vsv.ALLOWED_UNITS:
            errors.append((e["row"], "unit", e["unit"]))
        if e["status"] not in vsv.ALLOWED_STATUSES:
            errors.append((e["row"], "status", e["status"]))
        if not e["source_reference"]:
            errors.append((e["row"], "source_reference", None))
        try:
            SCOPE_RESOLVER.resolve_scopes(e["scope_type"], e["scope_value"])
        except ValueError as exc:
            errors.append((e["row"], "scope", str(exc)))
        if e["kind"] == "variable":
            if e["variable_type"] not in vsv.ALLOWED_VARIABLE_TYPES:
                errors.append((e["row"], "variable_type", e["variable_type"]))
            if e["frequency"] not in vsv.ALLOWED_FREQUENCIES:
                errors.append((e["row"], "frequency", e["frequency"]))
            if e["value"] is not None or e["version"] is not None:
                errors.append((e["row"], "parameter fields on variable", None))
        else:
            if not isinstance(e["value"], (int, float)) or isinstance(e["value"], bool):
                errors.append((e["row"], "value", e["value"]))
            if e["version"] != 1:
                errors.append((e["row"], "version", e["version"]))
            if e["frequency"] or e["variable_type"] or e["expression"]:
                errors.append((e["row"], "variable fields on parameter", None))
    assert errors == []


def test_value_type_is_declared_for_every_variable(model):
    assert "value_type" in model["headers"], (
        "workbook não declara value_type: estados textuais não podem ser "
        "distinguidos de números sem regra por nome"
    )
    bad = [
        (e["row"], e.get("value_type")) for e in model["entities"]
        if (e["kind"] == "variable" and e.get("value_type") not in VALUE_TYPES)
        or (e["kind"] == "parameter" and e.get("value_type") is not None)
    ]
    assert bad == []


# ============================================================
# Expressões: parsing, funções, referências, grafo
# ============================================================


def test_every_expression_translates_and_parses(model, registries):
    _, _, translated = registries
    parser = ExpressionParser()
    extractor = DependencyExtractor()
    for entity_id, expression in translated.items():
        parser.parse(expression)
        extractor.extract(expression)
    assert len(translated) == sum(e["role"] == "equation" for e in model["entities"])


def test_dsl_features_present(model, registries):
    _, _, translated = registries
    text = " ".join(translated.values())
    for feature in ("ln(", "@L1_L3", "@L4_L5", "@L6_L7", "@L1", " and ", " or ", '"F"'):
        assert feature in text, feature
    for state in STATES:
        assert f'"{state}"' in text, state


def test_aggregation_rows_use_platform_dsl(model):
    aggregations = [e for e in model["entities"] if e["role"] == "aggregation"]
    assert len(aggregations) == 10
    assert {e["dsl"][0] for e in aggregations} == {"AVERAGE"}
    free_text = [
        e["row"] for e in model["entities"]
        if e["role"] == "equation" and "resultados diários" in e["expression"]
    ]
    assert free_text == []


def test_scope_matrix_all_references_reachable(model):
    rows = scope_matrix(model)
    failures = [r for r in rows if r["result"] != "PASS"]
    assert failures == [], json.dumps(failures, ensure_ascii=False, indent=1)


def test_line_to_group_bindings(model):
    """§20: L1/L2/L3 -> L1_L3, L4/L5 -> L4_L5, L6/L7 -> L6_L7."""

    expected = {
        "L1": "L1_L3", "L2": "L1_L3", "L3": "L1_L3",
        "L4": "L4_L5", "L5": "L4_L5", "L6": "L6_L7", "L7": "L6_L7",
    }
    rows = scope_matrix(model)
    for line, group in expected.items():
        (row,) = [
            r for r in rows
            if r["consumer"] == "retirada_condensado_linha"
            and r["consumer_scope"] == f"linha/{line}"
            and r["reference"].startswith("retirada_condensado_grupo")
        ]
        assert row["reference"] == f"retirada_condensado_grupo@{group}"
        assert row["reference_scope"].startswith(f"linha_grupo/{group} (diário")
        assert row["result"] == "PASS"


def test_rateio_denominators_do_not_use_lth_grupo(model):
    for e in model["entities"]:
        if e["name"] == "retirada_condensado_linha" and e["role"] == "equation":
            assert "lth_grupo" not in e["expression"], e["row"]


def test_dependency_graph_orders_group_before_lines(model, registries):
    """Registrar as equações em ordem inversa não altera o resultado."""

    variables, equations, translated = registries
    reversed_registry = EquationDefinitionRegistry()
    for definition in reversed(equations.all()):
        reversed_registry.add(definition)
    context = new_context(variables)
    day = date(2026, 7, 1)
    set_inputs(model, context, day)
    ForecastEngine().calculate_from_definition_registry(
        reversed_registry, context, variables, run_date=day
    )
    assert value_of(model, context, "retirada_condensado_total", "L1_L7", day) == pytest.approx(
        expected_day()["total"]
    )


# ============================================================
# E4 / E10 — lth_grupo e rateio
# ============================================================


def test_lth_grupo_is_average(model, registries):
    context = new_context(registries[0])
    day = date(2026, 7, 1)
    run_day(model, registries, context, day)
    got = {g: value_of(model, context, "lth_grupo", g, day) for g in ("L1_L3", "L4_L5", "L6_L7")}
    print("\nlth_grupo", got)
    assert got == {"L1_L3": 20.0, "L4_L5": 50.0, "L6_L7": 50.0}


def test_group_total_and_rateio_reconcile(model, registries):
    context = new_context(registries[0])
    day = date(2026, 7, 1)
    run_day(model, registries, context, day)
    expected = expected_day()
    for group, members in (("L1_L3", "L1 L2 L3"), ("L4_L5", "L4 L5"), ("L6_L7", "L6 L7")):
        group_value = value_of(model, context, "retirada_condensado_grupo", group, day)
        assert group_value == pytest.approx(expected["group"][group])
        lines = [value_of(model, context, "retirada_condensado_linha", l, day) for l in members.split()]
        for l, v in zip(members.split(), lines):
            assert v == pytest.approx(expected["line"][l])
        assert sum(lines) == pytest.approx(group_value), group
    total = value_of(model, context, "retirada_condensado_total", "L1_L7", day)
    assert total == pytest.approx(expected["total"])


# ============================================================
# E7 / E8 / E9 — ln, estados textuais, "F"
# ============================================================


def test_ln_terms_match_reference(model, registries):
    context = new_context(registries[0])
    day = date(2026, 7, 1)
    run_day(model, registries, context, day)
    expected = expected_day()
    for g in ("L1_L3", "L4_L5", "L6_L7"):
        assert value_of(model, context, "retirada_cond_corr_ltp", g, day) == pytest.approx(expected["ltp"][g])
        assert value_of(model, context, "retirada_cond_corr_lth", g, day) == pytest.approx(expected["corr"][g])


def test_states_are_categorical(model, registries):
    variables = registries[0]
    for line in ("l4", "l5", "l6", "l7"):
        definition = variables.get(find(model["entities"], f"hes_{line}", "diário", line.upper())["entity_id"])
        assert definition.is_categorical


STATE_CASES = [
    ("Normal", "Normal"),
    ("LC", "Normal"),
    ("Normal", "LC"),
    ("Overhaul/Parada", "Normal"),
    ("LC", "Overhaul/Parada"),          # ordem: LC antes de Overhaul
    ("1 By pass", "Normal"),
    ("Overhaul/Parada", "1 By pass"),   # ordem: Overhaul antes de By pass
    ("1 By pass e LC", "1 By pass e LC"),
]


@pytest.mark.parametrize("group,lines", [("L4_L5", ("L4", "L5")), ("L6_L7", ("L6", "L7"))])
@pytest.mark.parametrize("state_a,state_b", STATE_CASES)
def test_conditional_branches_match_reference(model, registries, group, lines, state_a, state_b):
    states = {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal"}
    states[lines[0]], states[lines[1]] = state_a, state_b
    context = new_context(registries[0])
    day = date(2026, 7, 1)
    run_day(model, registries, context, day, states=states)
    expected = expected_day(states=states)
    assert value_of(model, context, "retirada_condensado_grupo", group, day) == pytest.approx(
        expected["group"][group]
    )


@pytest.mark.parametrize("group,lines", [("L4_L5", ("L4", "L5")), ("L6_L7", ("L6", "L7"))])
@pytest.mark.parametrize(
    "state_a,state_b",
    [("Normal", "1 By pass e LC"), ("1 By pass e LC", "Normal"), ("1 By pass e LC", "Overhaul")],
)
def test_uncovered_combination_yields_F_and_fails_consumers(model, registries, group, lines, state_a, state_b):
    """
    entrada -> nenhuma condição atendida -> "F" (gravado sem conversão)
    -> consumidor falha explicitamente (ConditionalFailureError).
    """

    states = {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal"}
    states[lines[0]], states[lines[1]] = state_a, state_b
    assert expected_day(states=states)["group"][group] == FAILURE

    variables = registries[0]
    context = new_context(variables)
    day = date(2026, 7, 1)
    with pytest.raises(EquationEvaluationError) as error:
        run_day(model, registries, context, day, states=states)
    assert isinstance(error.value.original_error, ConditionalFailureError)
    assert value_of(model, context, "retirada_condensado_grupo", group, day) == FAILURE


# ============================================================
# E5 — agregações
# ============================================================


def test_aggregation_materialization(model, registries):
    variables = registries[0]
    materialized = build_aggregation_instances(model, variables)
    table = defaultdict(list)
    for target, rule, instances in materialized:
        for instance in instances:
            table[(target["name"], target["frequency"], target["scope_value"])].append(
                (rule.aggregation_type, instance.scope_value)
            )
    lines = [f"L{n}" for n in range(1, 8)]
    for frequency in ("mensal", "anual"):
        linha = table[("retirada_condensado_linha", frequency, "L1_L7")]
        assert [scope for _, scope in linha] == lines
        assert {kind for kind, _ in linha} == {"AVERAGE"}
        for group in ("L1_L3", "L4_L5", "L6_L7"):
            assert table[("retirada_condensado_grupo", frequency, group)] == [("AVERAGE", group)]
        assert table[("retirada_condensado_total", frequency, "L1_L7")] == [("AVERAGE", "L1_L7")]
    assert sum(len(v) for v in table.values()) == 2 * (7 + 3 + 1)


def test_aggregations_average_daily_results(model, registries):
    """33 dias (01/01 a 02/02/2026) com LTH variando por dia."""

    variables = registries[0]
    context = new_context(variables)
    days = [date(2026, 1, 1) + timedelta(days=i) for i in range(33)]
    daily_expected = {}
    for i, day in enumerate(days):
        lth = {l: v + i for l, v in LTH.items()}
        run_day(model, registries, context, day, lth=lth)
        daily_expected[day] = expected_day(lth=lth)

    run_date = days[-1]
    service = TemporalAggregationService()
    february = [d for d in days if d.month == 2]

    def mean(values):
        return sum(values) / len(values)

    for target, rule, instances in build_aggregation_instances(model, variables):
        for instance in instances:
            result = service.aggregate(
                rule, context, instance.scope_type, instance.scope_value, run_date
            )
            window = february if target["frequency"] == "mensal" else days
            name = target["name"]
            if name == "retirada_condensado_linha":
                expected = mean([daily_expected[d]["line"][instance.scope_value] for d in window])
            elif name == "retirada_condensado_grupo":
                expected = mean([daily_expected[d]["group"][instance.scope_value] for d in window])
            else:
                expected = mean([daily_expected[d]["total"] for d in window])
            assert result.value == pytest.approx(expected), (name, target["frequency"], instance.scope_value)
            assert (result.scope_type, result.scope_value) == (instance.scope_type, instance.scope_value)

    # §10: média do total == soma das médias dos grupos (mesmos dias).
    feb_total = mean([daily_expected[d]["total"] for d in february])
    feb_groups = sum(
        mean([daily_expected[d]["group"][g] for d in february]) for g in ("L1_L3", "L4_L5", "L6_L7")
    )
    assert feb_total == pytest.approx(feb_groups)


def test_no_implicit_temporal_conversion_in_a41(model, registries):
    """§16 aplicado ao A41: nenhuma SUM de /h para /mês|/ano."""

    variables = registries[0]
    rules = [rule for _, rule, _ in build_aggregation_instances(model, variables)]
    assert find_sum_dimension_issues(rules, variables) == []
    for rule in rules:
        assert variables.get(rule.source_variable_id).unit == variables.get(rule.target_variable_id).unit


# ============================================================
# 41c / 41d — consistência com a fonte declarada
# ============================================================

_CELL = re.compile(r"^Forecast A41!([A-Z]+)(\d+)$")


# Decisão semântica do usuário (fechamento da Etapa 2). Não reinterpretar.
DISCOUNT_DECISION = {
    ("desconto_retirada_41c", "linha_grupo", "L4_L5"): "Forecast A41!C26",
    ("desconto_retirada_41c", "linha_grupo", "L6_L7"): "Forecast A41!C28",
    ("desconto_retirada_41d", "linha_grupo", "L4_L5"): "Forecast A41!C27",
}


def discount_source_consistency(model) -> dict:
    """
    Compara as entidades desconto_retirada_41c/41d do workbook com a
    decisão do usuário: mesmo conjunto de entidades (nenhuma a mais,
    nenhuma a menos), mesmo source_reference, nenhuma célula
    compartilhada.
    """

    found = {
        (e["name"], e["scope_type"], e["scope_value"]): e
        for e in model["entities"]
        if re.fullmatch(r"desconto_retirada_41[a-z]", e["name"])
    }
    findings = []
    for key in sorted(set(found) | set(DISCOUNT_DECISION)):
        entity = found.get(key)
        expected = DISCOUNT_DECISION.get(key)
        actual = entity["source_reference"] if entity else None
        findings.append(
            {
                "entity": f"{key[0]}@{key[2]}",
                "scope_type": key[1],
                "workbook_row": entity["row"] if entity else None,
                "expected_source": expected,
                "actual_source": actual,
                "consistent": expected is not None and actual == expected,
            }
        )
    cells = defaultdict(list)
    for key, entity in found.items():
        cells[entity["source_reference"]].append(f"{key[0]}@{key[2]}")
    shared = {cell: names for cell, names in cells.items() if len(names) > 1}
    return {"findings": findings, "shared_cells": shared}


def test_41c_41d_source_consistency(model):
    report = discount_source_consistency(model)
    print("\n" + json.dumps(report, ensure_ascii=False, indent=1))
    assert [f for f in report["findings"] if not f["consistent"]] == []
    assert report["shared_cells"] == {}
    names = {f["entity"] for f in report["findings"]}
    assert "desconto_retirada_41d@L6_L7" not in names
    by_cell = {f["actual_source"]: f["entity"] for f in report["findings"]}
    assert not by_cell.get("Forecast A41!C26", "").startswith("desconto_retirada_41d")
    assert not by_cell.get("Forecast A41!C27", "").startswith("desconto_retirada_41c")
    assert not by_cell.get("Forecast A41!C28", "").startswith("desconto_retirada_41d")


def test_discounts_reach_the_intended_branches(model, registries):
    """
    O desconto de cada grupo chega ao ramo correto (valores distintos por
    entidade): L4_L5 "1 By pass" usa 41d, "1 By pass e LC" usa 41c;
    L6_L7 usa 41c@L6_L7 nos dois ramos (expressões do workbook).
    """

    day = date(2026, 7, 1)
    for states, group in (
        ({"L4": "1 By pass", "L5": "Normal", "L6": "Normal", "L7": "Normal"}, "L4_L5"),
        ({"L4": "1 By pass e LC", "L5": "1 By pass e LC", "L6": "Normal", "L7": "Normal"}, "L4_L5"),
        ({"L4": "Normal", "L5": "Normal", "L6": "1 By pass", "L7": "Normal"}, "L6_L7"),
        ({"L4": "Normal", "L5": "Normal", "L6": "1 By pass e LC", "L7": "1 By pass e LC"}, "L6_L7"),
    ):
        context = new_context(registries[0])
        run_day(model, registries, context, day, states=states)
        assert value_of(model, context, "retirada_condensado_grupo", group, day) == pytest.approx(
            expected_day(states=states)["group"][group]
        )


def test_F_propagates_through_aggregation(model, registries):
    """
    Um dia com "F" no grupo L4_L5 faz a agregação mensal do grupo falhar
    explicitamente (AggregationFailureError com o período afetado); o
    "F" nunca é somado nem tratado como 0.
    """

    from app.engine.exceptions import AggregationFailureError

    variables = registries[0]
    context = new_context(variables)
    days = [date(2026, 3, 1), date(2026, 3, 2), date(2026, 3, 3)]
    for day in days:
        states = {"L4": "Normal", "L5": "Normal", "L6": "Normal", "L7": "Normal"}
        if day.day == 2:
            states["L5"] = "1 By pass e LC"
            with pytest.raises(EquationEvaluationError):
                run_day(model, registries, context, day, states=states)
        else:
            run_day(model, registries, context, day, states=states)

    service = TemporalAggregationService()
    for target, rule, instances in build_aggregation_instances(model, variables):
        if (target["name"], target["frequency"], target["scope_value"]) != (
            "retirada_condensado_grupo", "mensal", "L4_L5"
        ):
            continue
        (instance,) = instances
        with pytest.raises(AggregationFailureError) as error:
            service.aggregate(rule, context, instance.scope_type, instance.scope_value, days[-1])
        assert error.value.failed_period_ids == ["2026-03-02"]
        return
    raise AssertionError("regra mensal de retirada_condensado_grupo@L4_L5 não encontrada")


# ============================================================
# Evidência gravada
# ============================================================


BLOCKED_ENTITIES: set[str] = set()


def traceability(model, materialized) -> list[dict]:
    index = reference_resolver.build_name_index(model["entities"])
    aggregation_by_row = defaultdict(list)
    for target, rule, instances in materialized:
        aggregation_by_row[target["row"]] += [i.scope_value for i in instances]
    rows = []
    for e in model["entities"]:
        key = f"{e['name']}@{e['scope_value']}"
        cell = _CELL.match(e["source_reference"] or "")
        sheet, _, ref = (e["source_reference"] or "").partition("!")
        deps = (
            sorted({n + (f"@{s}" if s else "") for n, s in references(e["expression"], index)})
            if e["role"] == "equation"
            else ([e["dsl"][2]] if e["role"] == "aggregation" else [])
        )
        rows.append(
            {
                "row": e["row"],
                "entity": key,
                "frequency": e["frequency"] or "",
                "scope": f"{e['scope_type']}/{e['scope_value']}",
                "value_type": e.get("value_type") or "",
                "semantic_role": e["role"],
                "source_reference": e["source_reference"],
                "source_sheet": sheet,
                "source_cell": ref,
                "dependencies": " ".join(deps),
                "aggregation": (
                    f"{e['dsl'][0]} {e['dsl'][1]} de {e['dsl'][2]} diário -> "
                    f"{len(aggregation_by_row[e['row']])} instância(s): "
                    + " ".join(aggregation_by_row[e["row"]])
                    if e["role"] == "aggregation" else ""
                ),
                "validation_status": (
                    "BLOCKING" if key in BLOCKED_ENTITIES else "PASS"
                ),
            }
        )
    return rows


def write_csv(path: Path, rows: list[dict]) -> None:
    import csv

    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


V5_VALIDATED = HERE.parent / "area_41_v5" / "descritivo_das_variáveis_A41_v5.xlsx"
V6_RECEIVED = HERE / "descritivo_das_variáveis_A41_v6_recebido.xlsx"
V6_VALIDATED = HERE / "descritivo_das_variáveis_A41_v6.xlsx"


def change_matrix() -> list[dict]:
    """
    Célula a célula: v5 validado -> v6 recebido -> v6 validado.
    Classificação por evidência (decisão do usuário 41c/41d; correções
    estruturais D1/D2 do v5).
    """

    sheets = [openpyxl.load_workbook(path)[SHEET] for path in (V5_VALIDATED, V6_RECEIVED, V6_VALIDATED)]
    max_row = max(ws.max_row for ws in sheets)
    max_col = max(ws.max_column for ws in sheets)
    rows = []
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v5, v6r, v6v = (ws.cell(r, c).value for ws in sheets)
            if v5 == v6r == v6v:
                continue
            coordinate = sheets[0].cell(r, c).coordinate
            entity = sheets[2].cell(r, 2).value
            scope = sheets[2].cell(r, 10).value
            if coordinate in ("K30", "K31"):
                nature, classification, reason = (
                    "source_reference", "esperada",
                    "Decisão semântica do usuário 41c/41d (C26->41c@L4_L5, "
                    "C28->41c@L6_L7, C27->41d@L4_L5).",
                )
            elif coordinate.startswith("P"):
                nature, classification, reason = (
                    "value_type (D1)", "inesperada no v6 recebido; reaplicada",
                    "O v6 recebido deriva do input do v5 e perdeu a coluna "
                    "value_type exigida pelo contrato da Etapa 1; reaplicada "
                    "sem alteração de valor (correção mínima, §4).",
                )
            elif coordinate in ("M39", "M42"):
                nature, classification, reason = (
                    "expression (D2)", "inesperada no v6 recebido; reaplicada",
                    "hes_lN sem @LN não é alcançável a partir do grupo "
                    "(referência implícita nunca desce do grupo para a "
                    "linha); reaplicado hes_lN@LN, única mudança na célula.",
                )
            else:
                nature, classification, reason = ("?", "inesperada", "sem justificativa")
            rows.append(
                {
                    "cell": coordinate,
                    "entity": f"{entity}@{scope}" if entity and r > 2 else "(cabeçalho)",
                    "v5": "" if v5 is None else str(v5),
                    "v6_recebido": "" if v6r is None else str(v6r),
                    "v6_validado": "" if v6v is None else str(v6v),
                    "natureza": nature,
                    "classificacao": classification,
                    "justificativa": reason,
                }
            )
    return rows


def main():
    model = load_workbook(WORKBOOK)
    registries = build_registries(model)
    out = HERE / "evidence"
    out.mkdir(exist_ok=True)
    context = new_context(registries[0])
    day = date(2026, 7, 1)
    run_day(model, registries, context, day)
    materialized = build_aggregation_instances(model, registries[0])
    evidence = {
        "workbook": WORKBOOK.name,
        "value_type_contract": {
            "field": "value_type",
            "allowed": sorted(VALUE_TYPES),
            "categorical": sorted(
                f"{e['name']}@{e['scope_value']}" for e in model["entities"]
                if e.get("value_type") == "categorical"
            ),
        },
        "sha256": sha256(WORKBOOK),
        "entities": len(model["entities"]),
        "roles": {r: sum(e["role"] == r for e in model["entities"]) for r in ("input", "equation", "aggregation", "parameter")},
        "scope_matrix": scope_matrix(model),
        "translated_expressions": {
            f"{e['name']}@{e['scope_value']}": registries[2][e["entity_id"]]
            for e in model["entities"] if e["role"] == "equation"
        },
        "day_2026_07_01": {
            "lth_grupo": {g: value_of(model, context, "lth_grupo", g, day) for g in ("L1_L3", "L4_L5", "L6_L7")},
            "grupo": {g: value_of(model, context, "retirada_condensado_grupo", g, day) for g in ("L1_L3", "L4_L5", "L6_L7")},
            "linha": {f"L{n}": value_of(model, context, "retirada_condensado_linha", f"L{n}", day) for n in range(1, 8)},
            "total": value_of(model, context, "retirada_condensado_total", "L1_L7", day),
            "reference": expected_day(),
        },
        "aggregation_instances": [
            {
                "target": f"{t['name']} {t['frequency']} {t['scope_type']}/{t['scope_value']} (linha {t['row']})",
                "type": r.aggregation_type,
                "instances": [f"{i.scope_type}/{i.scope_value}" for i in inst],
            }
            for t, r, inst in materialized
        ],
        "discount_41c_41d": discount_source_consistency(model),
    }
    write_csv(out / "scope_matrix.csv", evidence["scope_matrix"])
    write_csv(out / "traceability_matrix.csv", traceability(model, materialized))
    write_csv(out.parent / "change_matrix.csv", change_matrix())
    (out / "a41_v6_evidence.json").write_text(
        json.dumps(evidence, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8"
    )
    print(out / "a41_v6_evidence.json")


if __name__ == "__main__":
    main()
