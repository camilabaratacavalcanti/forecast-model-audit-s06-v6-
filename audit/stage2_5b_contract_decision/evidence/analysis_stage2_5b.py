"""
Etapa 2.5B — auditoria contratual (somente leitura).

    python audit/stage2_5b_contract_decision/evidence/analysis_stage2_5b.py

Workbooks: area_41 v8, energy v5, max_ht v9 (data/workbooks, aprovados na
2.4) + production v8 e yield v10 (candidatos corrigidos pelo dono,
copiados byte a byte em evidence/inputs/). Builders em memória; nenhum
arquivo de data/, app/, tools/ ou tests/ é escrito.

Saídas em audit/stage2_5b_contract_decision/ (CSV exigidos) e evidence/.
"""

from __future__ import annotations

import csv
import json
import re
import shutil
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
REPO = HERE.parents[2]
sys.path.insert(0, str(REPO))

import openpyxl  # noqa: E402

from app.engine import reference_resolver  # noqa: E402
from app.engine.scope_resolver import ScopeResolver  # noqa: E402
from app.engine.scoped_reference import scope_type_for  # noqa: E402
from app.engine.spatial_candidate_resolver import get_spatial_candidates  # noqa: E402
from tools.workbook_seed.blocks import BLOCKS, seed_json  # noqa: E402
from tools.workbook_seed.canonical import build_canonical_model  # noqa: E402
from tools.workbook_seed.reader import read_workbook, sha256_of  # noqa: E402
from tools.workbook_seed.seeds import build_seeds  # noqa: E402

SR = ScopeResolver()
FREQ_RANK = {"diário": 0, "mensal": 1, "anual": 2}
SOURCES = {
    "area_41": (BLOCKS["area_41"].workbook_path, "A41", "v8"),
    "energy": (BLOCKS["energy"].workbook_path, "energy", "v5"),
    "max_ht": (BLOCKS["max_ht"].workbook_path, "MaxHT", "v9"),
    "production": (HERE / "inputs" / "descritivo_das_variáveis_production_v8.xlsx", "production", "v8"),
    "yield": (HERE / "inputs" / "descritivo_das_variáveis_yield_v10.xlsx", "yield", "v10"),
}
PREVIOUS = {
    "production": REPO / "audit/stage2_5_contract_closure/evidence/inputs/descritivo_das_variáveis_production_v7.xlsx",
    "yield": BLOCKS["yield"].workbook_path,
}
summary: dict = {}


def write_csv(path, rows):
    rows = rows or [{"result": "0 ocorrências"}]
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


# ------------------------------------------------------------------
# 1. hashes e diffs
# ------------------------------------------------------------------
lines = []
for block, (path, _sheet, version) in sorted(SOURCES.items()):
    digest = sha256_of(path)
    if block in ("production", "yield"):
        state = "CANDIDATO CORRIGIDO PELO DONO (2.5B)"
    else:
        state = "APROVADO 2.4 (inalterado)" if digest == BLOCKS[block].sha256 else "DIVERGENTE!"
    lines.append(f"{digest}  {path.name}  {block} {version}  {state}")
(HERE / "workbooks_sha256.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
summary["approved_unchanged"] = all(
    sha256_of(SOURCES[b][0]) == BLOCKS[b].sha256 for b in ("area_41", "energy", "max_ht")
)

diffs = []
for block, old in PREVIOUS.items():
    new, sheet, version = SOURCES[block]
    wa = openpyxl.load_workbook(old, data_only=True)[sheet]
    wb = openpyxl.load_workbook(new, data_only=True)[sheet]
    header_row = 1 if block == "yield" else 2
    header = {c: wb.cell(header_row, c).value for c in range(1, wb.max_column + 1)}
    for r in range(1, max(wa.max_row, wb.max_row) + 1):
        for c in range(1, max(wa.max_column, wb.max_column) + 1):
            if wa.cell(r, c).value != wb.cell(r, c).value:
                diffs.append({
                    "block": block, "from": old.name, "to": new.name, "cell": wb.cell(r, c).coordinate,
                    "column": header.get(c), "name": wb.cell(r, 2).value,
                    "before": wa.cell(r, c).value, "after": wb.cell(r, c).value,
                })
write_csv(HERE / "workbook_diffs.csv", diffs)
summary["diffs"] = Counter(d["block"] for d in diffs)

# ------------------------------------------------------------------
# 2. build em memória + integridade
# ------------------------------------------------------------------
workbooks, models, seeds = {}, {}, {}
for block, (path, sheet, _v) in sorted(SOURCES.items()):
    workbooks[block] = read_workbook(path, sheet)
    models[block] = build_canonical_model(block, workbooks[block], BLOCKS[block].id_base)
    seeds[block] = build_seeds(models[block], BLOCKS[block].id_base)

integrity = []
for block in sorted(SOURCES):
    rows = [r.values for r in workbooks[block].rows]
    integrity.append({
        "block": block, "version": SOURCES[block][2], "semantic_rows": len(rows),
        "definitions": len(models[block].entities),
        "variables": len(seeds[block]["variables"]), "parameter_records": len(seeds[block]["parameters"]),
        "equations": len(seeds[block]["equations"]), "aggregations": len(seeds[block]["aggregation_rules"]),
        "value_type": dict(Counter(r["value_type"] for r in rows)),
        "allowed_values_rows": sum(r["allowed_values"] is not None for r in rows),
        "declared_result_states_rows": sum(r["declared_result_states"] is not None for r in rows),
        "value_filled": dict(Counter(r["Type"] for r in rows if r["value"] is not None)),
        "version_filled": dict(Counter(r["Type"] for r in rows if r["version"] is not None)),
        "pending_contract_decisions": len(seeds[block]["manifest"]["pending_contract_decisions"]),
    })
write_csv(HERE / "workbook_integrity.csv", integrity)
summary["totals"] = {
    "rows": sum(i["semantic_rows"] for i in integrity),
    "equations": sum(i["equations"] for i in integrity),
    "aggregations": sum(i["aggregations"] for i in integrity),
}

tmp = Path(tempfile.mkdtemp())
try:
    shutil.copytree(REPO / "data" / "seed", tmp / "seed")
    for block in ("production", "yield"):
        for n in ("variables", "parameters", "equations", "aggregation_rules", "manifest"):
            (tmp / "seed" / block / f"{n}.json").write_text(seed_json(seeds[block][n]), encoding="utf-8")
    from app.validation import equation_seed_validator, parameter_seed_validator, variable_seed_validator  # noqa: E402
    from app.repositories.seed_loader import SeedLoader  # noqa: E402
    from app.engine.registry_validator import RegistryIntegrityValidator  # noqa: E402
    vres = variable_seed_validator.validate_seed(tmp / "seed")
    perr, _ = parameter_seed_validator.validate_seed(tmp / "seed")
    eerr, _ = equation_seed_validator.validate_seed(tmp / "seed")
    vd, _vi, pd, _pi, ed, _ei = SeedLoader(tmp / "seed").load_all_definitions_and_instances()
    RegistryIntegrityValidator().validate_definition_registry(ed, vd, pd)
    summary["validators_with_candidates"] = {
        "variable_errors": len(vres["errors"]), "variable_warnings": len(vres["warnings"]),
        "parameter_errors": len(perr), "equation_errors": len(eerr), "registry_integrity": "OK",
    }
finally:
    shutil.rmtree(tmp)

# ------------------------------------------------------------------
# 3. value / version (D24-12)
# ------------------------------------------------------------------
vv = []
for block in sorted(SOURCES):
    for r in workbooks[block].rows:
        v = r.values
        if v["value"] is None and v["version"] is None:
            continue
        kind = (
            "parameter" if v["Type"] == "parameter"
            else "aggregation" if v["Type"] == "variable / equation" and v["expression"] and re.search(r"resultados diários|Média móvel|média_ponderada", v["expression"])
            else "variable"
        )
        vv.append({
            "block": block, "workbook": workbooks[block].file_name, "row": r.row, "Type": v["Type"],
            "classified_as": kind, "name": v["name"], "frequency": v["frequency"],
            "scope_type": v["scope_type"], "scope_value": v["scope_value"],
            "value": v["value"], "value_physical_type": type(v["value"]).__name__, "version": v["version"],
            "value_type": v["value_type"], "status": v["status"],
            "verdict": "OK (parameter)" if kind == "parameter" else "CONTRACT_VIOLATION (D24-12)",
        })
write_csv(ROOT / "value_version_evidence.csv", vv)
summary["value_version"] = Counter(x["verdict"] for x in vv)

# ------------------------------------------------------------------
# 4. unidades (D25-02) — oee/oee_total e toda identidade repetida entre blocos
# ------------------------------------------------------------------
units = []
for block in sorted(SOURCES):
    for r in workbooks[block].rows:
        v = r.values
        if v["name"] in ("oee", "oee_total") or (v["expression"] and re.search(r"(?<![\w@])oee(_total)?(?![\w])", v["expression"])):
            units.append({
                "check": "oee/oee_total", "block": block, "row": r.row, "Type": v["Type"], "name": v["name"],
                "frequency": v["frequency"], "scope": f"{v['scope_type']}/{v['scope_value']}", "unit": v["unit"],
                "variable_type": v["variable_type"], "expression": v["expression"],
                "verdict": (
                    "OK unit='-'" if v["name"] in ("oee", "oee_total") and v["unit"] == "-"
                    else "CONTRACT_VIOLATION (D25-02)" if v["name"] in ("oee", "oee_total")
                    else "CONSUMIDOR (verificar escala na fórmula)"
                ),
            })
by_identity = defaultdict(list)
for block in sorted(SOURCES):
    for e in models[block].entities:
        by_identity[(e.name, e.frequency, e.scope_type, e.scope_value)].append((block, e))
for key, occ in sorted(by_identity.items()):
    if len({b for b, _ in occ}) < 2:
        continue
    unit_set = sorted({e.unit for _, e in occ})
    units.append({
        "check": "identidade repetida entre blocos", "block": " ".join(b for b, _ in occ), "row": "",
        "Type": " ".join(e.kind for _, e in occ), "name": key[0], "frequency": key[1],
        "scope": f"{key[2]}/{key[3]}", "unit": " | ".join(f"{b}:{e.unit}" for b, e in occ),
        "variable_type": " | ".join(f"{b}:{e.variable_type}" for b, e in occ), "expression": "",
        "verdict": "OK (unidades iguais)" if len(unit_set) == 1 else "DIVERGÊNCIA DE UNIDADE",
    })
write_csv(ROOT / "unit_evidence.csv", units)
summary["units"] = Counter(u["verdict"] for u in units)

# ------------------------------------------------------------------
# 5. referências (A019) — fallback temporal/espacial, inalcançáveis
# ------------------------------------------------------------------
REF = re.compile(r"\b(VAR\d+|PARAM\d+)(?:@(L[1-7](?:_L[1-7])?))?")
refs, verdicts = [], Counter()
for block in sorted(SOURCES):
    ents = {e.entity_id: e for e in models[block].entities}
    rows = {e["equation_id"]: e["row"] for e in seeds[block]["manifest"]["equations"]}
    index = reference_resolver.build_name_index([e.as_index_entry() for e in models[block].entities])
    for eq in seeds[block]["equations"]:
        target = ents[eq["target_variable_id"]]
        cscopes = SR.resolve_scopes(eq["scope_type"], eq["scope_value"])
        for m in REF.finditer(eq["expression"]):
            ref, explicit = ents[m.group(1)], m.group(2)
            cands = index[ref.name]
            inst = set(SR.resolve_scopes(ref.scope_type, ref.scope_value))
            spatial = set()
            for cs in cscopes:
                if explicit:
                    spatial.add("EXPLICIT_OK" if (scope_type_for(explicit), explicit) in inst else "EXPLICIT_MISSING")
                else:
                    chain = [tuple(x) for x in get_spatial_candidates(*cs)]
                    hit = next((i for i, x in enumerate(chain) if x in inst), None)
                    spatial.add("UNREACHABLE" if hit is None else ("OWN" if hit == 0 else "PARENT"))
            if ref.kind == "parameter":
                temporal, ffb = "PARAMETER", False
            else:
                a, b = FREQ_RANK[target.frequency], FREQ_RANK[ref.frequency]
                temporal = "SAME" if a == b else ("COARSER" if b > a else "FINER")
                ffb = not any(c["kind"] == "parameter" for c in cands) and not any(c.get("frequency") == target.frequency for c in cands)
            if spatial & {"UNREACHABLE", "EXPLICIT_MISSING"} or temporal == "FINER":
                verdict = "UNREACHABLE"
            elif ffb:
                verdict = "TEMPORAL_FALLBACK_COARSER (válido)"
            elif "PARENT" in spatial:
                verdict = "SPATIAL_PARENT (válido)"
            else:
                verdict = "DIRECT"
            verdicts[verdict] += 1
            refs.append({
                "block": block, "row": rows[eq["equation_id"]], "equation_id": eq["equation_id"],
                "consumer": f"{target.name}|{target.frequency}|{eq['scope_type']}/{eq['scope_value']}",
                "reference": f"{ref.name}|{ref.frequency}|{ref.scope_type}/{ref.scope_value}",
                "kind": ref.kind, "candidates_by_name": len(cands), "temporal": temporal,
                "spatial": "|".join(sorted(spatial)), "verdict": verdict,
            })
write_csv(HERE / "reference_inventory.csv", refs)
summary["references"] = {"total": len(refs), **verdicts}

# ------------------------------------------------------------------
# 6. cross-workbook: fonte, identidades repetidas (D25-01, D25-03, D25-04)
# ------------------------------------------------------------------
FONTE = re.compile(r"bloco\s+(\w+)", re.I)


def canonical_formula(block, e):
    """Expressão por nome (linha do workbook) ou regra de agregação."""
    out = set()
    for r in e.rows:
        if r.aggregation:
            a = r.aggregation
            out.add(f"AGG {a.aggregation_type} {a.source_name}" + (f" w={a.weight_name}" if a.weight_name else ""))
        elif r.expression:
            out.add("EQ " + re.sub(r"\s+", "", r.expression))
        else:
            out.add("INPUT")
    return " || ".join(sorted(out))


def producer_kind(block, e):
    if e.kind == "parameter":
        return "parameter"
    if any(q["target_variable_id"] == e.entity_id for q in seeds[block]["equations"]):
        return "equation"
    if e.entity_id in {r["target_variable_id"] for r in seeds[block]["aggregation_rules"]}:
        return "aggregation"
    return "input"


cross = []
for key, occ in sorted(by_identity.items()):
    blocks = {b for b, _ in occ}
    fontes_rows = []
    for b, e in occ:
        raw = {r.row: r.values for r in workbooks[b].rows}
        fontes = sorted({raw[r.row]["fonte"] for r in e.rows} - {None})
        fontes_rows.append((b, e, fontes))
    if len(blocks) < 2 and not any(f for *_x, f in fontes_rows):
        continue
    kinds = {b: producer_kind(b, e) for b, e in occ}
    formulas = {b: canonical_formula(b, e) for b, e in occ}
    computed = [b for b in kinds if kinds[b] in ("equation", "aggregation")]
    inputs = [b for b in kinds if kinds[b] == "input"]
    if len(blocks) < 2:
        continue
    if len(computed) >= 2:
        same = len({formulas[b] for b in computed}) == 1
        cls = "D25-01 cálculo local em vários blocos — " + ("fórmula idêntica (réplica textual)" if same else "fórmulas diferentes (cálculo independente)")
    elif computed and inputs:
        cls = "entrada local de identidade calculada em outro bloco"
    else:
        cls = "somente entradas/parâmetros"
    for b, e, fontes in fontes_rows:
        m = [FONTE.search(f) for f in fontes]
        cited = sorted({x.group(1).lower() for x in m if x})
        cross.append({
            "identity": "|".join(str(k) for k in key), "block": b, "definition_id": e.entity_id,
            "Type": e.kind, "variable_type": e.variable_type or "", "producer_kind_in_block": kinds[b],
            "unit": e.unit, "value_type": e.value_type,
            "instances": " ".join(s for _t, s in SR.resolve_scopes(e.scope_type, e.scope_value)),
            "fonte": " ; ".join(fontes), "fonte_cites_block": " ".join(cited),
            "fonte_cited_block_kind": " ".join(
                f"{c}:{kinds.get(c, 'ausente no bloco citado')}" for c in cited
            ),
            "formula": formulas[b], "classification": cls,
        })

# relações com fonte fora de repetições (ex.: fonte para blocos fora dos cinco) e casos D25-04
fonte_rows = []
for block in sorted(SOURCES):
    for e in models[block].entities:
        raw = {r.row: r.values for r in workbooks[block].rows}
        for r in e.rows:
            f = raw[r.row]["fonte"]
            if f is None:
                continue
            m = FONTE.search(f)
            cited = m.group(1).lower() if m else ""
            fonte_rows.append({
                "block": block, "row": r.row, "name": e.name, "frequency": e.frequency,
                "scope": f"{e.scope_type}/{r.scope_value}", "Type": raw[r.row]["Type"],
                "variable_type": e.variable_type or "", "has_expression": raw[r.row]["expression"] is not None,
                "fonte": f, "cited_block": cited, "cited_block_is_platform_block": cited in SOURCES,
                "cited_block_defines_identity": bool(cited in SOURCES and by_identity.get((e.name, e.frequency, e.scope_type, e.scope_value)) and any(b == cited for b, _ in by_identity[(e.name, e.frequency, e.scope_type, e.scope_value)])),
                "d25_04_case": (e.variable_type != "entrada"),
            })
write_csv(ROOT / "cross_workbook_evidence.csv", cross)
write_csv(HERE / "fonte_inventory.csv", fonte_rows)
summary["cross_identity_classes"] = Counter(r["classification"] for r in cross if True)
summary["fonte_rows"] = len(fonte_rows)
summary["fonte_by_variable_type"] = Counter(f"{r['variable_type']}|plataforma={r['cited_block_is_platform_block']}" for r in fonte_rows)
summary["d25_04_cases"] = [(r["block"], r["row"], r["name"], r["variable_type"], r["fonte"]) for r in fonte_rows if r["d25_04_case"]]

(HERE / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
print(json.dumps(summary, ensure_ascii=False, indent=1, default=str))
