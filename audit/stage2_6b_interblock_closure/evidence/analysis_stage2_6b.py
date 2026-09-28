"""
Etapa 2.6B — fechamento do contrato interbloco (evidência independente).

    python audit/stage2_6b_interblock_closure/evidence/analysis_stage2_6b.py

Independente do builder: NÃO importa nada de tools/ nem de app/. Lê os
cinco workbooks oficiais com openpyxl, declara a taxonomia D26-01 por
conta própria, reconstrói a classificação de cada `fonte` pelas regras
normativas e compara com a saída do builder gravada no repositório
(`data/seed/interblock_links.json`, cuja igualdade com o build é
verificada pela suíte de testes). Qualquer divergência reprova.

Também: diff célula a célula production v9 -> v10, SHA-256, investigação
mx_ht/max_ht e estabilidade de IDs (manifestos em a126e02 x atuais).
Nada fora de audit/stage2_6b_interblock_closure/ é escrito.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
REPO = HERE.parents[2]
WB = REPO / "data" / "workbooks"
BASELINE = "a126e02"

# Workbooks oficiais (identificador de código -> arquivo, aba, SHA-256 esperado).
OFFICIAL = {
    "area_41": ("descritivo_das_variáveis_A41_v9.xlsx", "A41",
                "36bbae135285ad9b3b88d21ef94e2bc9c92f2ef74736412b66bc62de4dcf3f50"),
    "energy": ("descritivo_das_variáveis_energy_v6.xlsx", "energy",
               "cfc46031fc2f3f375b87ebbba58676d23ad92f3ee43fb81b353fe5a8247e1df8"),
    "max_ht": ("descritivo_das_variáveis_MaxHT_v10.xlsx", "MaxHT",
               "3e40aab12ec0ec07918c5e358145b09f138663f635ca9e7fc0990d1b765b683b"),
    "production": ("descritivo_das_variáveis_production_v10.xlsx", "production",
                   "302cf18b92ac22c690b3add9bdd1ded0617bfa60636809696d010b8858d72487"),
    "yield": ("descritivo_das_variáveis_yield_v11.xlsx", "yield",
              "639e8b980f19bc20429cf9763157ac1828d9067bdbabea3bf3db0badbc2df931"),
}
PRODUCTION_V9 = ("descritivo_das_variáveis_production_v9.xlsx",
                 "3e3024e564ee6f18b82e9771957e61f4940becd1e80db5705143123377956a28")

# D26-01 — lista do proprietário, na ordem recebida.
TAXONOMY = (
    "maintenance", "area_04_13", "forecast_volume", "acido", "yield", "energy",
    "meta_volume_cheio", "custo_budget", "production", "boilers",
    "controle_espaco_vazio_meta", "custo_forecast_bdgt", "mx_ht", "volume",
    "lime_dia", "custo_forecast_real", "alumina", "soda",
    "floculante_hidrato_2026", "budget", "temperature_lp", "fator_residuo",
    "floculante_lama_dia", "forecast", "area_41", "vazao_condensado",
    "premissas_ppt_mensal", "shared",
)
LINES = [f"L{i}" for i in range(1, 8)]
CONTRACT_DIMENSIONS = ("kind", "unit", "value_type", "allowed_values", "declared_result_states")


def write_csv(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ------------------------------------------------------------------
# 0. integridade
# ------------------------------------------------------------------
sha_lines = []
for block, (name, _sheet, expected) in sorted(OFFICIAL.items()):
    digest = sha(WB / name)
    assert digest == expected, (block, digest)
    sha_lines.append(f"{digest}  data/workbooks/{name}  (oficial: {block})")
digest_v9 = sha(WB / PRODUCTION_V9[0])
assert digest_v9 == PRODUCTION_V9[1]
sha_lines.append(f"{digest_v9}  data/workbooks/{PRODUCTION_V9[0]}  (referência do diff; substituído por v10)")
(HERE / "workbook_sha256.txt").write_text("\n".join(sha_lines) + "\n", encoding="utf-8")


# ------------------------------------------------------------------
# 1. diff production v9 -> v10
# ------------------------------------------------------------------
def load_sheet(path, sheet):
    return openpyxl.load_workbook(path)[sheet]


v9 = load_sheet(WB / PRODUCTION_V9[0], "production")
v10 = load_sheet(WB / OFFICIAL["production"][0], "production")
header = [c.value for c in v9[2]]
diff_rows = []
for r in range(1, max(v9.max_row, v10.max_row) + 1):
    for c in range(1, max(v9.max_column, v10.max_column) + 1):
        a, b = v9.cell(r, c).value, v10.cell(r, c).value
        if a == b:
            continue
        column = header[c - 1] if c <= len(header) else f"col{c}"
        name = v10.cell(r, 2).value
        if name in ("fator_mpsa", "fator_mrn") and column == "fonte" and a == "forecast" and b is None:
            classification = "D26-03"
        elif name == "lth_meta" and column in ("Type", "value", "version", "variable_type") and (
            (column == "Type" and (a, b) == ("parameter", "variable"))
            or (column in ("value", "version") and b is None)
            or (column == "variable_type" and (a, b) == (None, "entrada_externa"))
        ):
            classification = "D26-02"
        elif column == "OBS":
            classification = "UNEXPECTED_OBS_CHANGE"
        else:
            classification = "UNEXPECTED"
        diff_rows.append({
            "cell": v10.cell(r, c).coordinate, "row": r, "column": column, "name": name,
            "scope_value": v10.cell(r, 10).value, "v9": a, "v10": b, "classification": classification,
        })
write_csv(HERE / "production_v9_v10_diff.csv", diff_rows)
diff_summary = Counter(d["classification"] for d in diff_rows)

# Valores antigos de lth_meta (value) x novo OBS: alinhamento linha a linha.
lth_meta_rows = [r for r in range(3, v9.max_row + 1) if v9.cell(r, 2).value == "lth_meta"]
obs_alignment = []
for r in sorted(set(lth_meta_rows) | {min(lth_meta_rows) - 1, max(lth_meta_rows) + 1}):
    obs_alignment.append({
        "row": r, "name": v10.cell(r, 2).value, "frequency": v10.cell(r, 8).value,
        "scope_value": v10.cell(r, 10).value,
        "v9_value": v9.cell(r, 5).value, "v10_OBS": v10.cell(r, 18).value,
        "aligned": v9.cell(r, 5).value == v10.cell(r, 18).value,
    })


# ------------------------------------------------------------------
# 2. leitura crua e definições
# ------------------------------------------------------------------
def raw_rows(block):
    name, sheet, _ = OFFICIAL[block]
    ws = load_sheet(WB / name, sheet)
    header_row, cols = None, None
    for r in range(1, 11):
        values = [c.value for c in ws[r]]
        if "name" in values and "fonte" in values:
            header_row, cols = r, values
            break
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        values = dict(zip(cols, [c.value for c in ws[r]]))
        if values.get("name") is None:
            continue
        values["_row"] = r
        values["_kind"] = "parameter" if values["Type"] == "parameter" else "variable"
        rows.append(values)
    return rows


def instances(scope_type, scope_value):
    if scope_type == "linha" and "_" in str(scope_value):
        a, b = str(scope_value).split("_")
        return {(scope_type, l) for l in LINES[LINES.index(a): LINES.index(b) + 1]}
    return {(scope_type, scope_value)}


RAW = {block: raw_rows(block) for block in OFFICIAL}


def definitions(block):
    groups = defaultdict(list)
    for r in RAW[block]:
        per_line = r["scope_type"] == "linha" and r["scope_value"] in LINES
        key = (r["_kind"], r["name"], r["frequency"], r["scope_type"], "*" if per_line else r["scope_value"])
        groups[key].append(r)
    out = []
    for key, rows in groups.items():
        inst = set()
        for r in rows:
            inst |= instances(r["scope_type"], r["scope_value"])
        out.append({"key": key, "rows": rows, "instances": inst, "first": rows[0], "block": block})
    return out


DEFS = {block: definitions(block) for block in OFFICIAL}


# ------------------------------------------------------------------
# 3. classificação independente
# ------------------------------------------------------------------
def classify(block, d):
    row = d["first"]
    source = row["fonte"]
    findings = []
    if row["_kind"] != "variable" or any(r["expression"] is not None for r in d["rows"]):
        findings.append(("CONTRACT_MISMATCH", "consumer_context"))
    if source not in TAXONOMY:
        return None, findings + [("SOURCE_BLOCK_UNKNOWN", "source_block")]
    if source not in OFFICIAL:
        return None, findings + [("SOURCE_BLOCK_NOT_LOADED", "source_block")]
    wanted = d["instances"]
    same_axes = [
        p for p in DEFS[source]
        if (p["key"][1], p["key"][2], p["key"][3]) == (row["name"], row["frequency"], row["scope_type"])
    ]
    covering = [p for p in same_axes if wanted <= p["instances"]]
    if len(covering) > 1:
        return None, findings + [("AMBIGUOUS", "identity")]
    if not covering:
        if not [p for p in DEFS[source] if p["key"][1] == row["name"]]:
            return None, findings + [("SOURCE_NOT_FOUND", "name")]
        if same_axes:
            union = set().union(*(p["instances"] for p in same_axes))
            if wanted <= union:
                return None, findings + [("AMBIGUOUS", "identity")]
            return None, findings + [("INSTANCE_MISMATCH", "instances")]
        return None, findings + [("SOURCE_NOT_FOUND", "frequency/scope")]
    producer = covering[0]
    p = producer["first"]
    for dimension in CONTRACT_DIMENSIONS:
        mine = row["_kind"] if dimension == "kind" else row[dimension]
        theirs = p["_kind"] if dimension == "kind" else p[dimension]
        if mine != theirs:
            findings.append(("CONTRACT_MISMATCH", dimension))
    return producer, findings


links = []
for block in sorted(OFFICIAL):
    for d in DEFS[block]:
        if d["first"]["fonte"] is None:
            continue
        # fonte deve ser igual em todas as linhas da definição
        assert len({r["fonte"] for r in d["rows"]}) == 1
        producer, findings = classify(block, d)
        links.append({"block": block, "def": d, "producer": producer, "findings": findings})

# Ciclos: vínculos + dependências intrabloco sobre-aproximadas (todo
# nome citado numa expressão liga a TODAS as definições desse nome no
# bloco). Sem ciclo na sobre-aproximação => sem ciclo real.
node = lambda d: (d["block"], d["key"])
graph = defaultdict(set)
for block, defs in DEFS.items():
    by_name = defaultdict(list)
    for d in defs:
        by_name[d["key"][1]].append(d)
    for d in defs:
        for r in d["rows"]:
            expression = r["expression"] or ""
            for token in set(re.findall(r"[A-Za-z_À-ÿ][\wÀ-ÿ]*", expression)):
                for target in by_name.get(token, []):
                    if target is not d:
                        graph[node(d)].add(node(target))
for l in links:
    if l["producer"]:
        graph[node(l["def"])].add(node(l["producer"]))

color, cycles = {}, []


def dfs(start):
    stack = [(start, iter(sorted(graph.get(start, ()), key=str)))]
    path = [start]
    color[start] = 1
    while stack:
        n, it = stack[-1]
        nxt = next(it, None)
        if nxt is None:
            stack.pop(); path.pop(); color[n] = 2
            continue
        if color.get(nxt) == 1:
            cycles.append(path[path.index(nxt):] + [nxt])
        elif not color.get(nxt):
            color[nxt] = 1; path.append(nxt)
            stack.append((nxt, iter(sorted(graph.get(nxt, ()), key=str))))


for n in sorted(graph, key=str):
    if not color.get(n):
        dfs(n)
interblock_cycles = [
    c for c in cycles
    if any(a[0] != b[0] for a, b in zip(c, c[1:]))
]
for l in links:
    if any(node(l["def"]) in c for c in interblock_cycles) and l["producer"]:
        l["findings"].append(("CYCLE", "cycle"))

for l in links:
    l["status"] = l["findings"][0][0] if l["findings"] else "VALID"

# Cadeias
by_consumer = {node(l["def"]): l for l in links}
for l in links:
    chain, status, cur, seen = [f"{l['block']}.{l['def']['key'][1]}"], "VALID", l, set()
    while True:
        if node(cur["def"]) in seen:
            status = "CYCLE"; break
        seen.add(node(cur["def"]))
        src = cur["def"]["first"]["fonte"]
        if cur["status"] != "VALID":
            status = ("PENDING_AT:" if cur["status"] == "SOURCE_BLOCK_NOT_LOADED" else "BROKEN_AT:") + \
                f"{cur['block']}.{cur['def']['key'][1]}"
            chain.append(f"{src}.{cur['producer']['key'][1] if cur['producer'] else '?'}")
            break
        chain.append(f"{src}.{cur['producer']['key'][1]}")
        nxt = by_consumer.get(node(cur["producer"]))
        if nxt is None:
            break
        cur = nxt
    l["chain"], l["chain_status"] = " → ".join(chain), status


# ------------------------------------------------------------------
# 4. comparação com o builder (seed gravado)
# ------------------------------------------------------------------
seed = json.loads((REPO / "data/seed/interblock_links.json").read_text(encoding="utf-8"))
builder = {}
for rec in seed["links"]:
    builder[(rec["consumer_block"], tuple(rec["consumer_rows"]))] = ("VALID", rec["source_block"])
for group in ("pending", "rejected"):
    for rec in seed[group]:
        builder[(rec["consumer_block"], tuple(rec["consumer_rows"]))] = (rec["validation_status"], rec["source_block"])

divergences = []
independent = {}
for l in links:
    key = (l["block"], tuple(r["_row"] for r in l["def"]["rows"]))
    independent[key] = (l["status"], l["def"]["first"]["fonte"])
for key in sorted(set(independent) | set(builder)):
    if independent.get(key) != builder.get(key):
        divergences.append({"key": key, "independent": independent.get(key), "builder": builder.get(key)})

taxonomy_divergences = []
if tuple(seed["taxonomy"]["official_blocks"]) != TAXONOMY:
    taxonomy_divergences.append("seed.taxonomy.official_blocks != D26-01")


# ------------------------------------------------------------------
# 5. CSVs
# ------------------------------------------------------------------
link_rows, contract_rows = [], []
for l in links:
    d, p = l["def"], l["producer"]
    c = d["first"]
    pf = p["first"] if p else None
    for r in d["rows"]:
        link_rows.append({
            "consumer_block": l["block"],
            "consumer_row": r["_row"],
            "consumer_name": c["name"],
            "consumer_unit": c["unit"],
            "consumer_frequency": c["frequency"],
            "consumer_scope": f"{r['scope_type']}/{r['scope_value']}",
            "source_block": c["fonte"],
            "source_in_taxonomy": c["fonte"] in TAXONOMY,
            "source_block_loaded": c["fonte"] in OFFICIAL,
            "producer_name": pf["name"] if pf else "",
            "producer_kind": pf["_kind"] if pf else "",
            "producer_unit": pf["unit"] if pf else "",
            "producer_frequency": pf["frequency"] if pf else "",
            "producer_scope": f"{p['key'][3]}/{'/'.join(sorted(v for _t, v in p['instances']))}" if p else "",
            "consumer_variable_type": c["variable_type"],
            "producer_variable_type": pf["variable_type"] if pf else "",
            "classification": l["status"],
            "state": "RESOLVED" if l["status"] == "VALID" else (
                "PENDING_LOAD" if l["status"] == "SOURCE_BLOCK_NOT_LOADED" else "REJECTED"),
            "findings": ";".join(f"{k}:{dim}" for k, dim in l["findings"]),
            "chain": l["chain"],
            "chain_status": l["chain_status"],
            "builder_classification": builder.get((l["block"], tuple(x["_row"] for x in d["rows"])), ("?",))[0],
            "agreement": "AGREE" if independent[(l["block"], tuple(x["_row"] for x in d["rows"]))]
            == builder.get((l["block"], tuple(x["_row"] for x in d["rows"]))) else "DIVERGE",
        })
    if not p:
        contract_rows.append({
            "consumer_block": l["block"], "consumer_rows": ",".join(str(x["_row"]) for x in d["rows"]),
            "consumer_name": c["name"], "source_block": c["fonte"], "dimension": "source_block",
            "consumer_value": c["fonte"],
            "producer_value": "fora da taxonomia" if c["fonte"] not in TAXONOMY else "workbook não carregado",
            "result": l["status"], "is_contract_dimension": True,
        })
        continue
    values = {
        "kind": (c["_kind"], pf["_kind"]),
        "name": (c["name"], pf["name"]),
        "unit": (c["unit"], pf["unit"]),
        "value_type": (c["value_type"], pf["value_type"]),
        "allowed_values": (c["allowed_values"], pf["allowed_values"]),
        "declared_result_states": (c["declared_result_states"], pf["declared_result_states"]),
        "frequency": (c["frequency"], pf["frequency"]),
        "scope_type": (c["scope_type"], pf["scope_type"]),
        "instances": (sorted(v for _t, v in d["instances"]), sorted(v for _t, v in p["instances"])),
        "variable_type": (c["variable_type"], pf["variable_type"]),
    }
    for dimension, (mine, theirs) in values.items():
        if dimension == "instances":
            ok = set(mine) <= set(theirs)
        else:
            ok = mine == theirs
        contract_rows.append({
            "consumer_block": l["block"], "consumer_rows": ",".join(str(x["_row"]) for x in d["rows"]),
            "consumer_name": c["name"], "source_block": c["fonte"], "dimension": dimension,
            "consumer_value": mine, "producer_value": theirs,
            "result": ("OK" if ok else "DIFFERENT (não é dimensão do contrato)")
            if dimension == "variable_type" else ("OK" if ok else "MISMATCH"),
            "is_contract_dimension": dimension != "variable_type",
        })
write_csv(ROOT / "interblock_links.csv", link_rows)
write_csv(ROOT / "contract_validation.csv", contract_rows)

graph_rows = []
block_edges = Counter()
for l in links:
    block_edges[(l["block"], l["def"]["first"]["fonte"], l["status"])] += 1
    graph_rows.append({
        "level": "definition",
        "consumer": f"{l['block']}.{l['def']['key'][1]}[{l['def']['key'][2]} {l['def']['key'][3]}]",
        "producer": (f"{l['def']['first']['fonte']}.{l['producer']['key'][1]}"
                     f"[{l['producer']['key'][2]} {l['producer']['key'][3]}]")
        if l["producer"] else f"{l['def']['first']['fonte']}.?",
        "edges": 1, "status": l["status"], "note": f"{l['chain']} ({l['chain_status']})",
    })
pairs = {(a, b) for (a, b, _s) in block_edges}
for (a, b, status), n in sorted(block_edges.items()):
    graph_rows.append({
        "level": "block", "consumer": a, "producer": b, "edges": n, "status": status,
        "note": ("dependência mútua entre blocos; sem ciclo no nível de definição"
                 if (b, a) in pairs else ""),
    })
write_csv(ROOT / "interblock_graph.csv", graph_rows)


# ------------------------------------------------------------------
# 6. mx_ht / max_ht e taxonomia x código
# ------------------------------------------------------------------
fontes = sorted({l["def"]["first"]["fonte"] for l in links})
naming = {
    "taxonomy_name": "mx_ht" if "mx_ht" in TAXONOMY else None,
    "max_ht_in_taxonomy": "max_ht" in TAXONOMY,
    "workbook_file": OFFICIAL["max_ht"][0],
    "workbook_sheet": OFFICIAL["max_ht"][1],
    "workbook_declared_block_title": None,  # nenhuma célula/propriedade declara título de bloco
    "code_block_id": "max_ht",
    "fonte_using_mx_ht": [l["block"] for l in links if l["def"]["first"]["fonte"] == "mx_ht"],
    "fonte_using_max_ht": [l["block"] for l in links if l["def"]["first"]["fonte"] == "max_ht"],
    "divergent": True,
}
wb_titles = {}
for block, (name, sheet, _) in OFFICIAL.items():
    wb = openpyxl.load_workbook(WB / name)
    wb_titles[block] = {"sheets": wb.sheetnames, "properties_title": wb.properties.title}
assert wb_titles["max_ht"]["properties_title"] is None
loaded_vs_taxonomy = {block: block in TAXONOMY for block in OFFICIAL}


# ------------------------------------------------------------------
# 7. estabilidade de IDs (manifestos baseline x atuais)
# ------------------------------------------------------------------
def manifest_ids(text):
    return {
        (e["kind"], e["name"], e["frequency"], e["scope_type"], e["scope_value"]): e["entity_id"]
        for e in json.loads(text)["entities"]
    }


id_report = {}
for block in OFFICIAL:
    before = manifest_ids(subprocess.check_output(
        ["git", "show", f"{BASELINE}:data/seed/{block}/manifest.json"], cwd=REPO).decode("utf-8"))
    after = manifest_ids((REPO / f"data/seed/{block}/manifest.json").read_text(encoding="utf-8"))
    id_report[block] = {
        "renumbered": sorted(
            [list(k), before[k], after[k]] for k in before.keys() & after.keys() if before[k] != after[k]
        ),
        "retired": sorted([list(k), before[k]] for k in before.keys() - after.keys()),
        "new": sorted([list(k), after[k]] for k in after.keys() - before.keys()),
    }


# ------------------------------------------------------------------
# 8. resumo
# ------------------------------------------------------------------
by_class_defs = Counter(l["status"] for l in links)
by_class_rows = Counter(r["classification"] for r in link_rows)
summary = {
    "baseline": BASELINE,
    "workbooks": {b: {"file": v[0], "sha256": v[2]} for b, v in sorted(OFFICIAL.items())},
    "production_v9_v10_diff": {"cells": len(diff_rows), "by_classification": dict(diff_summary)},
    "lth_meta_obs_alignment": obs_alignment,
    "taxonomy": list(TAXONOMY),
    "fonte_names": {f: {"in_taxonomy": f in TAXONOMY, "loaded": f in OFFICIAL} for f in fontes},
    "loaded_blocks_in_taxonomy": loaded_vs_taxonomy,
    "mx_ht_max_ht": naming,
    "workbook_titles": wb_titles,
    "fonte_rows": len(link_rows),
    "declared_links_definitions": len(links),
    "links_by_class_definitions": dict(sorted(by_class_defs.items())),
    "links_by_class_rows": dict(sorted(by_class_rows.items())),
    "cycles_overapproximated": [[str(n) for n in c] for c in interblock_cycles],
    "independent_vs_builder_divergences": [
        {k: (list(v) if isinstance(v, tuple) else v) for k, v in d.items()} for d in divergences
    ],
    "taxonomy_divergences": taxonomy_divergences,
    "id_stability": id_report,
}
(HERE / "validation_results.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2, default=list) + "\n", encoding="utf-8")

print(json.dumps({k: summary[k] for k in (
    "production_v9_v10_diff", "fonte_names", "loaded_blocks_in_taxonomy", "fonte_rows",
    "declared_links_definitions", "links_by_class_definitions", "links_by_class_rows",
    "cycles_overapproximated", "independent_vs_builder_divergences", "taxonomy_divergences",
)}, ensure_ascii=False, indent=1))
print("renumbered:", {b: len(v["renumbered"]) for b, v in id_report.items()},
      "retired:", {b: v["retired"] for b, v in id_report.items() if v["retired"]},
      "new:", {b: v["new"] for b, v in id_report.items() if v["new"]})
assert not divergences and not taxonomy_divergences
assert all(not v["renumbered"] for v in id_report.values())
