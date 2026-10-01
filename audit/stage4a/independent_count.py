"""
Stage 4A — recálculo INDEPENDENTE do universo integrado de 5 blocos.

    python -I audit/stage4a/independent_count.py

Não importa `app/`, `tools/` nem os harnesses da Stage 3: lê só os JSON dos
seeds, `interblock_links.json`, `plan_evidence.csv` (3.2) e o workbook A41 v9
(openpyxl). As regras de escopo, de fecho de dependências e de status do plano
são reescritas aqui a partir dos contratos documentados (3.1/3.2/3.4A), para
que as cardinalidades do contrato 4A não dependam do código sob teste.

Saída: JSON em stdout. `--check` confronta com `contract_expectations.json`
e retorna 1 se houver divergência. Falha também se `app` ou `tools` tiverem
sido importados (prova de independência).
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
SEED = REPO / "data" / "seed"
WORKBOOK = REPO / "data" / "workbooks" / "descritivo_das_variáveis_A41_v9.xlsx"
PLAN_EVIDENCE = REPO / "audit" / "stage3_2_execution_orchestration" / "evidence" / "plan_evidence.csv"
EXPECTATIONS = HERE / "contract_expectations.json"
BLOCKS4 = ("production", "yield", "energy", "max_ht")
BLOCKS5 = BLOCKS4 + ("area_41",)
ACTIVE = {"PUBLISHED", "ativo"}
REFERENCE = re.compile(r"\b(VAR\d{5})")
LINES = [f"L{i}" for i in range(1, 8)]
HES_VALUES = ("Normal", "LC", "Overhaul/Parada", "1 By pass", "1 By pass e LC")


# ------------------------------------------------------------------ leitura
def load(block: str, name: str) -> list[dict]:
    return json.loads((SEED / block / f"{name}.json").read_text(encoding="utf-8"))


def scope_instances(scope_type: str, scope_value: str | None) -> list[str]:
    """linha/Lx_Ly -> uma instância por linha do intervalo; demais -> uma instância."""
    if scope_type == "linha":
        if "_" in scope_value:
            a, b = (int(x[1:]) for x in scope_value.split("_"))
            return [f"L{i}" for i in range(a, b + 1)]
        return [scope_value]
    return [scope_value]


def catalog() -> dict:
    variables, equations, rules, block_of = {}, {}, {}, {}
    for block in BLOCKS5:
        for v in load(block, "variables"):
            variables[v["variable_id"]] = v
            block_of[v["variable_id"]] = block
        for e in load(block, "equations"):
            if e["status"] in ACTIVE:
                equations[e["equation_id"]] = e
                block_of[e["equation_id"]] = block
        for r in load(block, "aggregation_rules"):
            rules[r["aggregation_rule_id"]] = r
            block_of[r["aggregation_rule_id"]] = block
    links_payload = json.loads((SEED / "interblock_links.json").read_text(encoding="utf-8"))
    return {"variables": variables, "equations": equations, "rules": rules, "block_of": block_of,
            "links": links_payload["links"], "pending": links_payload["pending"]}


def graph(cat: dict, with_links: bool = True) -> tuple[dict, dict, dict]:
    """nós {key: (kind, block, produces)}, deps {key: {var}}, producers {var: {key}}."""
    nodes, deps = {}, {}
    for eid, e in cat["equations"].items():
        nodes["EQUATION:" + eid] = ("EQUATION", cat["block_of"][eid], e["target_variable_id"])
        deps["EQUATION:" + eid] = set(REFERENCE.findall(e["expression"]))
    for rid, r in cat["rules"].items():
        nodes["AGGREGATION:" + rid] = ("AGGREGATION", cat["block_of"][rid], r["target_variable_id"])
        deps["AGGREGATION:" + rid] = {r["source_variable_id"]} | ({r["weight_variable_id"]}
                                                                   if r.get("weight_variable_id") else set())
    if with_links:
        for link in cat["links"]:
            nodes["TRANSFER:" + link["consumer_definition"]] = ("TRANSFER", link["consumer_block"],
                                                                link["consumer_definition"])
            deps["TRANSFER:" + link["consumer_definition"]] = {link["source_definition"]}
    producers = defaultdict(set)
    for key, (_k, _b, produced) in nodes.items():
        producers[produced].add(key)
    return nodes, deps, producers


def closure(targets, deps, producers) -> tuple[set, set]:
    """(nós, variáveis visitadas) do fecho de dependências dos alvos."""
    seen_nodes, seen_vars, stack = set(), set(targets), list(targets)
    while stack:
        variable = stack.pop()
        for key in producers.get(variable, ()):
            if key in seen_nodes:
                continue
            seen_nodes.add(key)
            for dep in deps[key]:
                if dep not in seen_vars:
                    seen_vars.add(dep)
                    stack.append(dep)
    return seen_nodes, seen_vars


def targets_of(nodes, blocks) -> list[str]:
    return sorted({produced for (_k, block, produced) in nodes.values() if block in blocks})


def instances_of_node(cat, key) -> int:
    kind = key.split(":", 1)[0]
    node_id = key.split(":", 1)[1]
    if kind == "EQUATION":
        e = cat["equations"][node_id]
        return len(scope_instances(e["scope_type"], e["scope_value"]))
    if kind == "AGGREGATION":
        target = cat["variables"][cat["rules"][node_id]["target_variable_id"]]
        return len(scope_instances(target["scope_type"], target["scope_value"]))
    link = next(link for link in cat["links"] if link["consumer_definition"] == node_id)
    return len(link["instances"])


def universe(cat, blocks) -> dict:
    nodes, deps, producers = graph(cat)
    targets = targets_of(nodes, blocks)
    selected, visited = closure(targets, deps, producers)
    produced = {nodes[k][2] for k in selected}
    kinds = Counter(nodes[k][0] for k in selected)
    by_block = defaultdict(Counter)
    for k in selected:
        by_block[nodes[k][1]][nodes[k][0]] += 1
    events = sum(instances_of_node(cat, k) for k in selected)
    return {"targets": len(targets), "targets_by_block": dict(Counter(cat["block_of"][t] for t in targets)),
            "nodes": len(selected), "nodes_by_kind": dict(sorted(kinds.items())),
            "nodes_by_block": {b: dict(sorted(c.items())) for b, c in sorted(by_block.items())},
            "required_inputs": len(visited - produced),
            "equation_instances": sum(instances_of_node(cat, k) for k in selected if k.startswith("EQUATION")),
            "aggregation_instances": sum(instances_of_node(cat, k) for k in selected if k.startswith("AGGREGATION")),
            "transfer_instances_per_date": sum(instances_of_node(cat, k) for k in selected if k.startswith("TRANSFER")),
            "events_per_date": events,
            "transfers": sorted(k for k in selected if k.startswith("TRANSFER")),
            "_selected": selected}


# ------------------------------------------------------------------ plano oficial (sem fixture)
def official_plan(cat) -> dict:
    """
    Plano OFICIAL: os consumidores dos 16 vínculos pendentes não têm produtor;
    um alvo cujo fecho alcança um consumidor pendente é INTERBLOCK_SOURCE_NOT_LOADED
    (contrato 3.2), com os blocos-fonte pendentes alcançados.
    """
    nodes, deps, producers = graph(cat)
    pending_source = {p["consumer_definition"]: p["source_block"] for p in cat["pending"]}
    rows = {}
    for target in targets_of(nodes, BLOCKS5):
        selected, visited = closure([target], deps, producers)
        blocked = sorted({pending_source[v] for v in visited | {target} if v in pending_source})
        rows[target] = {"block": cat["block_of"][target],
                        "status": "INTERBLOCK_SOURCE_NOT_LOADED" if blocked else "OK",
                        "pending_blocks": ";".join(blocked), "steps": "" if blocked else str(len(selected))}
    evidence = {r["target"]: r for r in csv.DictReader(PLAN_EVIDENCE.open(encoding="utf-8"))}
    diffs = []
    for target in sorted(set(rows) | set(evidence)):
        a, b = rows.get(target), evidence.get(target)
        if a is None or b is None:
            diffs.append(f"{target}: presente só em {'evidência' if a is None else 'recálculo'}")
            continue
        got = (a["block"], a["status"], a["pending_blocks"], a["steps"])
        want = (b["block"], b["observed"], b["pending_blocks"], b["steps"])
        if got != want:
            diffs.append(f"{target}: recálculo {got} != plan_evidence {want}")
    area = [r for r in rows.values() if r["block"] == "area_41"]
    return {"targets": len(rows),
            "area_41": dict(Counter(r["status"] for r in area)),
            "area_41_pending_blocks": sorted({r["pending_blocks"] for r in area if r["pending_blocks"]}),
            "pending_links": len(cat["pending"]),
            "pending_by_source_block": dict(sorted(Counter(p["source_block"] for p in cat["pending"]).items())),
            "pending_by_consumer_block": dict(sorted(Counter(p["consumer_block"] for p in cat["pending"]).items())),
            "plan_evidence_rows": len(evidence),
            "differences_vs_plan_evidence": diffs}


# ------------------------------------------------------------------ hes (enumeração independente)
def hes_matrix() -> list[dict]:
    """
    Avalia a expressão LITERAL do workbook (colunas expression das linhas
    `retirada_condensado_grupo` diário L4_L5 e L6_L7) com o `eval` do Python,
    após trocar `nome@Lx` por identificadores. Entradas numéricas escolhidas
    para que cada ramo tenha valor distinto (identificação do ramo).
    """
    import openpyxl
    sheet = openpyxl.load_workbook(WORKBOOK, read_only=True)["A41"]
    expressions = {}
    for row in sheet.iter_rows(min_row=3, values_only=True):
        if row[1] == "retirada_condensado_grupo" and row[7] == "diário" and row[9] in ("L4_L5", "L6_L7"):
            expressions[row[9]] = row[15]
    env_numbers = {"retirada_cond_corr_ltp": 3.0, "retirada_cond_corr_lth": 2.0,
                   "desconto_retirada_41c": 3.0, "desconto_retirada_41d": 5.0}
    d = 1.0
    base = 100 - d * 2
    branch_values = {
        "Normal": base,
        "LC": (base * 14 + (base / 2) * 10) / 24,
        "Overhaul/Parada": 50 - d * 2,
        "1 By pass[41d]": base - 5.0,
        "1 By pass[41c]": base - 3.0,
        "1 By pass e LC[41c]": (base * 14 + ((base - 3.0) / 2) * 10) / 24,
        "F": "F",
    }
    rows = []
    for group, (la, lb) in (("L4_L5", ("L4", "L5")), ("L6_L7", ("L6", "L7"))):
        text = re.sub(r"\bhes@(L\d)\b", r"hes_\1", expressions[group])
        for a in HES_VALUES:
            for b in HES_VALUES:
                env = dict(env_numbers, **{f"hes_{la}": a, f"hes_{lb}": b})
                value = eval(compile(text, f"<A41 {group}>", "eval"), {"__builtins__": {}}, env)  # noqa: S307
                branch = [name for name, v in branch_values.items() if v == value]
                rows.append({"group": group, "hes_first": a, "hes_second": b, "lines": f"{la},{lb}",
                             "branch": "|".join(branch) or "UNIDENTIFIED", "value": value})
    return rows


# ------------------------------------------------------------------ OBS
OBS_CLASSIFICATION = (
    # (padrão no texto da observação, classificação, motivo)
    ("Dados não utilizados na planilha", "REQUIRES_FOLLOWUP",
     "parâmetro sem uso em nenhuma equação; a inclusão na plataforma é decisão do cliente (não afeta cálculo)"),
    ("Yield 2026!", "DOCUMENTATION_ONLY", "rastreabilidade da origem de lth (vínculo yield.VAR11031 -> area_41.VAR16007)"),
    ("irão receber strings", "DOCUMENTATION_ONLY", "descreve o domínio categórico de hes (allowed_values)"),
)


def obs_register() -> list[dict]:
    import openpyxl
    sheet = openpyxl.load_workbook(WORKBOOK, read_only=True)["A41"]
    rows = []
    for number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        obs = row[17]
        if not obs:
            continue
        cls, why = "REQUIRES_FOLLOWUP", "confirmação de negócio pendente com o cliente (não bloqueia, §10.1)"
        for pattern, c, w in OBS_CLASSIFICATION:
            if pattern in obs:
                cls, why = c, w
        rows.append({"row": number, "type": row[0], "name": row[1], "frequency": row[7], "scope_type": row[8],
                     "scope_value": row[9], "obs": obs, "classification": cls, "reason": why})
    return rows


# ------------------------------------------------------------------ escopos das referências
def scope_references(cat) -> dict:
    out = Counter()
    for eid, e in cat["equations"].items():
        if cat["block_of"][eid] != "area_41":
            continue
        for _var, suffix in re.findall(r"\b(VAR\d{5})@(\w+)", e["expression"]):
            out[(e["scope_type"], suffix)] += 1
    return {f"{st}@{sfx}": n for (st, sfx), n in sorted(out.items())}


def main() -> int:
    cat = catalog()
    u5, u4 = universe(cat, BLOCKS5), universe(cat, BLOCKS4)
    ua = universe(cat, ("area_41",))
    shared = sorted(u4["_selected"] & ua["_selected"])
    links_raw = (SEED / "interblock_links.json").read_bytes()
    result = {
        "universe_5": {k: v for k, v in u5.items() if not k.startswith("_")},
        "universe_4": {k: v for k, v in u4.items() if not k.startswith("_")},
        "universe_area_41_only": {k: v for k, v in ua.items() if not k.startswith("_")},
        "union": {"nodes_4": u4["nodes"], "nodes_area_41_only": ua["nodes"], "shared": shared,
                  "union": len(u4["_selected"] | ua["_selected"]),
                  "union_equals_5": (u4["_selected"] | ua["_selected"]) == u5["_selected"]},
        "link_VAR16007_used_in_5": "TRANSFER:VAR16007" in u5["_selected"],
        "link_VAR16007_used_in_4": "TRANSFER:VAR16007" in u4["_selected"],
        "official_plan": official_plan(cat),
        "interblock_links_sha256": hashlib.sha256(links_raw).hexdigest(),
        "hes_f_combinations": sorted([(r["group"], r["hes_first"], r["hes_second"])
                                      for r in hes_matrix() if r["branch"] == "F"]),
        "hes_unidentified": [r for r in hes_matrix() if r["branch"] == "UNIDENTIFIED"],
        "obs_rows": len(obs_register()),
        "scope_references_area_41": scope_references(cat),
        "area_41_aggregation_rules": len([r for r in cat["rules"] if cat["block_of"][r] == "area_41"]),
        "area_41_equations": len([e for e in cat["equations"] if cat["block_of"][e] == "area_41"]),
    }
    imported = sorted(m for m in sys.modules if m == "app" or m.startswith(("app.", "tools")))
    result["imports_app_or_tools"] = imported
    problems = [f"INDEPENDENCE_FAILURE módulos importados {imported}"] if imported else []
    problems += [f"PLAN_EVIDENCE_DIVERGENCE {d}" for d in result["official_plan"]["differences_vs_plan_evidence"]]
    if "--check" in sys.argv[1:]:
        expected = json.loads(EXPECTATIONS.read_text(encoding="utf-8"))["independent"]
        for path, want in expected.items():
            got = result
            for part in path.split("."):
                got = got[part]
            if json.loads(json.dumps(got)) != want:
                problems.append(f"EXPECTATION_DIVERGENCE {path}: {got!r} != {want!r}")
    result["problems"] = problems
    result["result"] = "PASS" if not problems else "FAIL"
    print(json.dumps(result, indent=1, ensure_ascii=False, default=list))
    return 0 if not problems else 1


if __name__ == "__main__":
    sys.exit(main())
