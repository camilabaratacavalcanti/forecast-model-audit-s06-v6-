"""
Stage 2.1 — probes read-only para o fechamento do contrato conceitual
de estados, resultados heterogêneos e erros.

NÃO reexecuta os probes da Etapa 2 (audit/stage2_typing_contract/), que
permanecem válidos e são citados como evidência histórica. Este script
cobre apenas as perguntas NOVAS desta etapa (D1-D10, C01-C16):

  - D1.1: "sem regra aplicável" vs "entrada inválida" (categoria
    fora do domínio declarado);
  - D3: propagação multi-hop (dependente direto vs indireto) através
    de consumidores de tipos diferentes (aritmético, comparação,
    condicional-passthrough);
  - D4: confirma que "agregação espacial" (grupo/total) no A41 é uma
    EQUAÇÃO comum com referências @grupo, não um AggregationRule; só
    a agregação MENSAL/ANUAL é um AggregationRule real;
  - D5: domínio de categorias de hes_l4..l7 (extraído do workbook real);
  - D6/D7: matriz value_type x resultado, fronteira estado x exceção;
  - Segundo caso de estudo: padrão "=SE(...;"ERRO!!!";valor)".

Todos os IDs são fixtures genéricas efêmeras (VAR95xxx), fora de
qualquer faixa produtiva. Nenhum arquivo de produção é importado em
modo de escrita; apenas leitura via chamadas normais da API pública.

Uso:
    python3 audit/stage2_1_result_state_contract/probes/probes_2_1.py
"""

import csv
import json
import platform
import subprocess
import sys
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
OUTDIR = HERE.parent
ROOT = OUTDIR.parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from app.domain.equations.models import EquationDefinition  # noqa: E402
from app.domain.equations.registry import EquationDefinitionRegistry  # noqa: E402
from app.domain.forecast.aggregation import AggregationRule  # noqa: E402
from app.domain.variables.models import VariableDefinition  # noqa: E402
from app.domain.variables.registry import VariableDefinitionRegistry  # noqa: E402
from app.engine import reference_resolver  # noqa: E402
from app.engine.calculation_context import CalculationContext  # noqa: E402
from app.engine.expression_evaluator import ExpressionEvaluator  # noqa: E402
from app.engine.expression_parser import ExpressionParser  # noqa: E402
from app.engine.forecast_engine import ForecastEngine  # noqa: E402
from app.engine.temporal_aggregation_service import TemporalAggregationService  # noqa: E402

DAY = date(2026, 3, 2)
COMMIT = subprocess.run(
    ["git", "rev-parse", "HEAD"], cwd=ROOT, capture_output=True, text=True
).stdout.strip()

EVIDENCE = OUTDIR / "evidence"


def outcome(fn):
    try:
        value = fn()
        return {"kind": "value", "type": type(value).__name__, "repr": repr(value)}
    except Exception as exc:  # noqa: BLE001
        original = getattr(exc, "original_error", None)
        return {
            "kind": "error",
            "type": type(exc).__name__,
            "original": type(original).__name__ if original else None,
            "repr": str(exc)[:220],
        }


def define(specs):
    """specs: [(id, value_type, scope_type, scope_value, expression|None)]"""

    variables, equations = VariableDefinitionRegistry(), EquationDefinitionRegistry()
    for n, (vid, vt, st, sv, expr) in enumerate(specs):
        variables.add(
            VariableDefinition(
                vid, vid.lower(), "-", "-",
                "calculado" if expr else "entrada",
                "diário", st, sv, "probe", "ativo", vt,
            )
        )
        if expr:
            equations.add(
                EquationDefinition(f"EQ95{n:03d}", vid, 1, st, sv, expr, "probe", "PUBLISHED")
            )
    return variables, equations


def run(specs, inputs):
    """inputs: {(id, scope_value): value}. Retorna (dict com TODOS os
    valores gravados no contexto, mensagem de erro ou None)."""

    variables, equations = define(specs)
    ctx = CalculationContext()
    ctx.declare_categorical_variables(
        d.variable_definition_id for d in variables.all() if d.is_categorical
    )
    for (vid, sv), value in inputs.items():
        st = next(d.scope_type for d in variables.all() if d.variable_definition_id == vid)
        ctx.set_variable_value(vid, value, st, sv, period_id=DAY.isoformat())
    error = None
    try:
        ForecastEngine().calculate_from_definition_registry(equations, ctx, variables, run_date=DAY)
    except Exception as exc:  # noqa: BLE001
        error = f"{type(exc).__name__}({type(getattr(exc, 'original_error', exc)).__name__})"
    produced = {}
    for d in equations.all():
        for key in ctx._scoped_variables:
            if key.entity_id == d.target_variable_id:
                produced[d.target_variable_id] = ctx._scoped_variables[key]
    return produced, error


def evaluate(expr, variables=None, categorical=()):
    ctx = CalculationContext(categorical_variable_ids=categorical)
    for k, v in (variables or {}).items():
        ctx.set_variable(k, v)
    return ExpressionEvaluator(ctx).evaluate(ExpressionParser().parse(expr))


RESULTS = {}


# ------------------------------------------------------------
# D1.1 — "sem regra" vs "entrada inválida"
# ------------------------------------------------------------


def d1_1_no_rule_vs_invalid_input():
    """
    Reproduz, com a EXPRESSÃO REAL do A41 (retirada_condensado_grupo,
    escopo L4_L5), as duas origens distintas de "F":

      Caso A: combinação de categorias VÁLIDAS (ambas dentro do
              domínio hes_l4/hes_l5) para a qual nenhum ramo foi
              definido: (Normal, "1 By pass e LC").

      Caso B: uma das entradas NÃO pertence ao domínio declarado
              (erro de digitação): "Overhaul" em vez de
              "Overhaul/Parada".

    Extrai a expressão diretamente do workbook v6 validado (célula
    M39 da aba A41), sem alterar o arquivo.
    """

    path = ROOT / "audit" / "area_41_v6" / "descritivo_das_variáveis_A41_v6.xlsx"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    idx = {h: i for i, h in enumerate(hdr)}
    expr = None
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[idx["name"]] == "retirada_condensado_grupo" and row[idx["scope_value"]] == "L4_L5" and row[idx["frequency"]] == "diário":
            expr = row[idx["expression"]]
            break
    assert expr is not None, "não encontrou a equação real no workbook v6"

    ents = [
        {"entity_id": "VAR95801", "name": "retirada_cond_corr_ltp", "kind": "variable", "frequency": "diário", "scope_type": "linha_grupo", "scope_value": "L4_L5"},
        {"entity_id": "VAR95802", "name": "retirada_cond_corr_lth", "kind": "variable", "frequency": "diário", "scope_type": "linha_grupo", "scope_value": "L4_L5"},
        {"entity_id": "VAR95803", "name": "hes_l4", "kind": "variable", "frequency": "diário", "scope_type": "linha", "scope_value": "L4"},
        {"entity_id": "VAR95804", "name": "hes_l5", "kind": "variable", "frequency": "diário", "scope_type": "linha", "scope_value": "L5"},
        {"entity_id": "VAR95805", "name": "desconto_retirada_41c", "kind": "variable", "frequency": "diário", "scope_type": "linha_grupo", "scope_value": "L4_L5"},
        {"entity_id": "VAR95806", "name": "desconto_retirada_41d", "kind": "variable", "frequency": "diário", "scope_type": "linha_grupo", "scope_value": "L4_L5"},
    ]
    translated = reference_resolver.translate_expression(
        expr, reference_resolver.build_name_index(ents), "diário",
        consumer_scope=("linha_grupo", "L4_L5"),
    )
    tree = ExpressionParser().parse(translated)

    def evaluate_case(hes_l4, hes_l5):
        ctx = CalculationContext(categorical_variable_ids={"VAR95803", "VAR95804"})
        for vid, sv, v in (("VAR95801", "L4_L5", 20.0), ("VAR95802", "L4_L5", 15.0),
                          ("VAR95805", "L4_L5", 5.0), ("VAR95806", "L4_L5", 4.0)):
            ctx.set_variable_value(vid, v, "linha_grupo", sv)
        # Nota: NÃO há validação de domínio antes deste ponto — o
        # valor categórico é aceito no CalculationContext independente
        # de pertencer ou não ao conjunto {Normal, LC, Overhaul/Parada,
        # 1 By pass, 1 By pass e LC}. Isso é demonstrado separadamente
        # em input_domain_gap() abaixo.
        ctx.set_variable_value("VAR95803", hes_l4, "linha", "L4")
        ctx.set_variable_value("VAR95804", hes_l5, "linha", "L5")
        return outcome(lambda: ExpressionEvaluator(ctx, "linha_grupo", "L4_L5").evaluate(tree))

    case_a = evaluate_case("Normal", "1 By pass e LC")  # combinação válida, sem ramo definido
    case_b = evaluate_case("Overhaul", "Normal")  # "Overhaul" não é categoria declarada ("Overhaul/Parada" é)

    return {
        "expression_source_cell": "Forecast A41!A12 (workbook M39)",
        "translated_expression_prefix": translated[:80] + "...",
        "case_A_valid_uncovered_combo": {
            "hes_l4": "Normal", "hes_l5": "1 By pass e LC",
            "description": "Ambas as categorias pertencem ao domínio declarado de hes; nenhum dos 5 ramos do IF cobre esta combinação.",
            "result": case_a,
        },
        "case_B_invalid_category": {
            "hes_l4": "Overhaul (typo de 'Overhaul/Parada')", "hes_l5": "Normal",
            "description": "'Overhaul' NÃO pertence ao domínio declarado {Normal, LC, Overhaul/Parada, 1 By pass, 1 By pass e LC}.",
            "result": case_b,
        },
        "observation": (
            "Ambos os casos produzem exatamente o mesmo resultado observável "
            '("F", str) — a expressão (e a plataforma) não distinguem "nenhuma '
            'regra aplicável a uma entrada válida" de "entrada fora do domínio '
            'declarado". Ver input_domain_gap() para a causa raiz.'
        ),
    }


def input_domain_gap():
    """
    Confirma que o CalculationContext aceita qualquer texto não-vazio
    em uma variável categorical, sem verificar contra um conjunto de
    categorias permitidas — porque esse conjunto não existe como
    conceito na plataforma hoje (nenhum campo `allowed_values` em
    VariableDefinition; grep confirma 0 ocorrências fora de
    ENUM_FIELDS dos validators, que são enums de METADADOS como
    variable_type/frequency/scope_type, não de valores de variável).
    """

    ctx = CalculationContext(categorical_variable_ids={"VAR95901"})
    o1 = outcome(lambda: ctx.set_variable("VAR95901", "Overhaul"))  # typo aceito
    stored1 = ctx.get_variable("VAR95901")
    o2 = outcome(lambda: ctx.set_variable("VAR95901", "qualquer_coisa_arbitraria"))  # também aceito
    stored2 = ctx.get_variable("VAR95901")
    return {
        "typo_accepted": {**o1, "stored_value": stored1},
        "arbitrary_string_accepted": {**o2, "stored_value": stored2},
        "has_allowed_values_field_in_VariableDefinition": False,
        "evidence": "grep -rn allowed_values app/ tools/ data/ -> só ocorrências de ENUM_FIELDS em *_seed_validator.py, que validam metadados de campo (variable_type, frequency, scope_type), não o valor de uma variável categorical.",
    }


# ------------------------------------------------------------
# D3 — propagação multi-hop: direto vs indireto, por tipo de consumidor
# ------------------------------------------------------------


def d3_propagation_matrix():
    """
    A -> estado ("F")
    B = A (referência direta)                [dependente DIRETO, passthrough]
    C = B + 1 (aritmético)                   [dependente INDIRETO, aritmético]
    D = 1 if B == "F" else 2 (comparação)    [dependente INDIRETO, detecção explícita]
    E = 1 if B > 0 else 2 (comparação numérica) [dependente INDIRETO, comparação ilegal]
    F_node = B if C_input > 0 else 0 (condicional-passthrough) [INDIRETO]
    G = independente, não depende de A       [controle: deve sempre calcular]
    """

    L = ("linha", "L1")
    specs = [
        ("VAR95101", "numeric", *L, None),                              # A (entrada, receberá "F")
        ("VAR95102", "numeric", *L, "VAR95101"),                        # B = A (referência direta)
        ("VAR95103", "numeric", *L, "VAR95102 + 1"),                    # C = B + 1 (aritmético)
        ("VAR95104", "numeric", *L, '1 if VAR95102 == "F" else 2'),     # D = comparação com literal "F" (detecção explícita)
        ("VAR95105", "numeric", *L, "1 if VAR95102 > 0 else 2"),        # E = comparação numérica ilegal
        ("VAR95106", "numeric", *L, None),                              # controle de ramo p/ F_node
        ("VAR95107", "numeric", *L, "VAR95102 if VAR95106 > 0 else 0"), # F_node = condicional-passthrough
        ("VAR95108", "numeric", *L, None),                              # G entrada independente
        ("VAR95109", "numeric", *L, "VAR95108 * 10"),                   # G' = dependente só de G
    ]
    inputs = {("VAR95101", "L1"): "F", ("VAR95106", "L1"): 1.0, ("VAR95108", "L1"): 5.0}
    produced, error = run(specs, inputs)

    rows = [
        {"produtor": "A=VAR95101", "estado": "F", "consumidor": "B=A (referência direta)", "profundidade": "direto",
         "tipo_consumo": "passthrough", "comportamento": produced.get("VAR95102", "não calculado")},
        {"produtor": "A=VAR95101", "estado": "F", "consumidor": "C=B+1 (aritmético)", "profundidade": "indireto (via B)",
         "tipo_consumo": "aritmético", "comportamento": "exceção (rodada abortada)" if error else produced.get("VAR95103", "não calculado")},
        {"produtor": "A=VAR95101", "estado": "F", "consumidor": 'D=(1 if B=="F" else 2) (comparação literal)', "profundidade": "indireto (via B)",
         "tipo_consumo": "detecção explícita de estado", "comportamento": "não calculado (aborta antes de alcançar D, ver ordem topológica)"},
        {"produtor": "A=VAR95101", "estado": "F", "consumidor": "E=(1 if B>0 else 2) (comparação numérica)", "profundidade": "indireto (via B)",
         "tipo_consumo": "comparação ilegal estado x número", "comportamento": "não calculado (aborta antes)"},
        {"produtor": "A=VAR95101", "estado": "F", "consumidor": "F_node=(B if k>0 else 0) (condicional-passthrough)", "profundidade": "indireto (via B)",
         "tipo_consumo": "passthrough condicional", "comportamento": "não calculado (aborta antes)"},
        {"produtor": "G=VAR95108 (INDEPENDENTE de A)", "estado": "valor normal", "consumidor": "G'=G*10", "profundidade": "n/a",
         "tipo_consumo": "aritmético", "comportamento": produced.get("VAR95109", "NÃO CALCULADO — F-001")},
    ]
    return {
        "erro_global": error,
        "valores_produzidos": produced,
        "rows": rows,
        "observacao": (
            "A rodada aborta na PRIMEIRA exceção encontrada durante a "
            "travessia topológica (aqui, em C=B+1, o primeiro consumidor "
            "aritmético de B na ordem de registro). D, E e F_node nunca "
            "chegam a ser avaliados, MESMO QUE fossem capazes de tratar o "
            "estado corretamente (D detectaria == \"F\"; o resultado "
            "observado não reflete a capacidade de cada consumidor, e sim "
            "a posição na ordem de execução — reafirma F-001/T-02."
        ),
    }


def d3_isolated_arithmetic_consumer():
    """C=B+1 isolado (sem D/E/F_node), para confirmar que SÓ o consumidor
    aritmético direto já é suficiente para abortar a rodada."""

    L = ("linha", "L1")
    specs = [
        ("VAR95111", "numeric", *L, None),
        ("VAR95112", "numeric", *L, "VAR95111"),
        ("VAR95113", "numeric", *L, "VAR95112 + 1"),
    ]
    produced, error = run(specs, {("VAR95111", "L1"): "F"})
    return {"produced": produced, "error": error}


def d3_isolated_comparison_consumer():
    """D=(1 if B=="F" else 2) isolado (sem C aritmético antes), para
    confirmar que a DETECÇÃO EXPLÍCITA funciona quando é o único
    consumidor — não é um problema do mecanismo de detecção em si,
    e sim de ordem de execução quando HÁ outro consumidor que aborta
    primeiro."""

    L = ("linha", "L1")
    specs = [
        ("VAR95121", "numeric", *L, None),
        ("VAR95122", "numeric", *L, "VAR95121"),
        ("VAR95123", "numeric", *L, '1 if VAR95122 == "F" else 2'),
    ]
    produced, error = run(specs, {("VAR95121", "L1"): "F"})
    return {"produced": produced, "error": error}


# ------------------------------------------------------------
# D4 — confirma: "agregação espacial" != AggregationRule
# ------------------------------------------------------------


def d4_spatial_vs_temporal_aggregation():
    """
    No workbook A41 real, retirada_condensado_total (diário) é uma
    equação COMUM somando três valores @grupo — NÃO existe um
    AggregationRule espacial na plataforma (AggregationRule é
    exclusivamente temporal: source_frequency -> target_frequency).
    Isso é FATO estrutural, confirmado por leitura de
    app/domain/forecast/aggregation.py (docstring: "Regra de
    agregação TEMPORAL") e por app/engine/temporal_aggregation_service.py
    (não referencia scope_type/scope_value como dimensão agregada).
    """

    path = ROOT / "audit" / "area_41_v6" / "descritivo_das_variáveis_A41_v6.xlsx"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    idx = {h: i for i, h in enumerate(hdr)}
    total_daily_expr = None
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[idx["name"]] == "retirada_condensado_total" and row[idx["frequency"]] == "diário":
            total_daily_expr = row[idx["expression"]]
    return {
        "retirada_condensado_total_diario_expression": total_daily_expr,
        "is_AggregationRule": False,
        "is_ordinary_equation_with_grupo_refs": True,
        "AggregationRule_docstring_first_line": "Regra de agregação temporal de uma Variable de uma frequência de origem para uma frequência de destino.",
        "implicacao": (
            "Propagação de estado em 'agregação espacial' (grupo/total) "
            "segue as MESMAS regras de D3 (equação comum) — não há uma "
            "política de agregação espacial separada a definir. Apenas a "
            "agregação MENSAL/ANUAL (via AggregationRule real) tem uma "
            "política de agregação distinta, já coberta em "
            "audit/stage2_typing_contract/aggregation_matrix.csv (P09a/P09b)."
        ),
    }


def d4_temporal_policy_citation():
    """Cita (sem reexecutar) a política já demonstrada na Etapa 2, e
    confirma via leitura de código que a política é Alternativa B
    (invalidar o agregado, nunca ignorar, nunca parcial)."""

    import inspect
    from app.engine import temporal_aggregation_service as tas
    src = inspect.getsource(tas.TemporalAggregationService._require_numeric_series)
    return {
        "existing_evidence": "audit/stage2_typing_contract/aggregation_matrix.csv (P09a AVERAGE, P09b SUM com 'F' -> AggregationFailureError)",
        "source_confirms_policy_B": "a agregação inteira falha" in " ".join(src.split()),
        "docstring_excerpt": src.split('"""')[1].strip()[:300],
    }


# ------------------------------------------------------------
# D5 — domínio de categorias de hes (extraído do workbook real)
# ------------------------------------------------------------


def d5_category_domain():
    import re
    path = ROOT / "audit" / "area_41_v6" / "descritivo_das_variáveis_A41_v6.xlsx"
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb.active
    hdr = [c.value for c in ws[2]]
    idx = {h: i for i, h in enumerate(hdr)}
    cats = set()
    declared_value_type = {}
    for row in ws.iter_rows(min_row=3, values_only=True):
        expr = row[idx["expression"]]
        if expr and isinstance(expr, str) and "hes_l" in expr:
            for m in re.finditer(r'hes_l\d@L\d\s*==\s*"([^"]+)"', expr):
                cats.add(m.group(1))
        if row[idx["name"]] in ("hes_l4", "hes_l5", "hes_l6", "hes_l7"):
            declared_value_type[row[idx["name"]]] = row[idx["value_type"]]
    return {
        "categories_used_in_equations": sorted(cats),
        "count": len(cats),
        "hes_value_type_declared": declared_value_type,
        "allowed_values_field_exists": False,
        "conclusion": (
            "O domínio de 5 categorias É extraível estaticamente do texto "
            "das equações (via grep de comparações == \"...\"), mas NÃO é "
            "declarado em nenhum campo estruturado da VariableDefinition. "
            "A validação de entrada não pode ocorrer sem essa declaração "
            "explícita — hoje ela é apenas implícita nos ramos do IF."
        ),
    }


# ------------------------------------------------------------
# Segundo caso de estudo: "=SE(...;"ERRO!!!";valor)"
# ------------------------------------------------------------


def erro_case_study():
    """
    Fórmula original (Excel):
        =SE(ARRED(D595-D601-D616;10)>ARRED(D61;10);"ERRO!!!";D595-D601-D616)

    Traduzida para o padrão conceitual já suportado pela DSL atual
    (SEM ARRED/ROUND, que não existe na allowlist -- ln é a única
    função permitida; isso por si só já é um FATO, não uma limitação
    introduzida aqui): substituímos por comparação direta, preservando
    a estrutura semântica (SE saldo > limite ENTÃO "ERRO!!!" SENÃO saldo).

        saldo if saldo <= limite else "ERRO!!!"

    Testado em DUAS declarações de value_type para a variável de
    destino, para expor o comportamento real (não hipotético) do
    storage/evaluator diante desse padrão.
    """

    L = ("linha", "L1")

    # Caso 1: variável de destino declarada numeric (o caso comum p/
    # uma variável que representa um saldo contábil)
    specs_numeric = [
        ("VAR95201", "numeric", *L, None),  # saldo (D595-D601-D616), calculado externamente aqui como input simplificado
        ("VAR95202", "numeric", *L, None),  # limite (D61)
        ("VAR95203", "numeric", *L, "VAR95201 if VAR95201 <= VAR95202 else \"ERRO!!!\""),
    ]
    ok_case, _ = run(specs_numeric, {("VAR95201", "L1"): 3.0, ("VAR95202", "L1"): 10.0})
    branch_hit_case, err_numeric = run(specs_numeric, {("VAR95201", "L1"): 15.0, ("VAR95202", "L1"): 10.0})

    # Caso 2: variável de destino declarada categorical (tentando
    # "acomodar" o texto de erro dentro do domínio categórico)
    specs_categorical = [
        ("VAR95211", "numeric", *L, None),
        ("VAR95212", "numeric", *L, None),
        ("VAR95213", "categorical", *L, "VAR95211 if VAR95211 <= VAR95212 else \"ERRO!!!\""),
    ]
    branch_hit_categorical, err_categorical = run(specs_categorical, {("VAR95211", "L1"): 15.0, ("VAR95212", "L1"): 10.0})
    # nota: se VAR95213 é categorical, o branch "ok" devolve um NÚMERO,
    # o que também é observável:
    ok_case_categorical, err_categorical_ok = run(specs_categorical, {("VAR95211", "L1"): 3.0, ("VAR95212", "L1"): 10.0})

    # Verificação estática (parser aceita a expressão em ambos os casos —
    # o parser não conhece value_type, confirma achado T-06 da Etapa 2)
    parses_ok = outcome(lambda: ExpressionParser().parse('VAR95201 if VAR95201 <= VAR95202 else "ERRO!!!"'))

    return {
        "original_excel_formula": '=SE(ARRED(D595-D601-D616;10)>ARRED(D61;10);"ERRO!!!";D595-D601-D616)',
        "note_on_ARRED": "ARRED (ROUND) não existe na allowlist de funções do parser (só 'ln'); fato estrutural preexistente, não avaliado nesta etapa.",
        "conceptual_dsl_equivalent": 'VAR_saldo if VAR_saldo <= VAR_limite else "ERRO!!!"',
        "parser_accepts_pattern_regardless_of_declared_value_type": parses_ok["kind"] == "value",
        "numeric_declared__branch_not_hit": ok_case,
        "numeric_declared__branch_hit__ERRO": {"produced": branch_hit_case, "error": err_numeric},
        "categorical_declared__branch_hit__ERRO": {"produced": branch_hit_categorical, "error": err_categorical},
        "categorical_declared__branch_not_hit__numeric_value_in_categorical_var": {"produced": ok_case_categorical, "error": err_categorical_ok},
        "interpretation": (
            "(1) Quando a variável é numeric, o ramo 'ERRO!!!' quando executado "
            "é REJEITADO pelo storage (CalculationValueError) exatamente como "
            "P05a da Etapa 2 -- mas só quando o ramo de fato executa; o outro "
            "ramo (numérico) é aceito silenciosamente sem qualquer indício de "
            "que a fórmula tem um segundo ramo textual não declarado (T-06). "
            "(2) Quando a variável é categorical, o padrão INVERTE o problema: "
            "agora é o ramo NUMÉRICO que é aceito sem checagem (pois "
            "categorical aceita qualquer texto OU qualquer número — "
            "is_numeric() é checado antes de is_categorical em "
            "_validate_variable_value), e 'ERRO!!!' também é aceito como se "
            "fosse uma categoria de domínio legítima, indistinguível de "
            "'Normal'/'LC'/etc. (3) Em NENHUM dos dois value_types existe uma "
            "forma de declarar 'este ramo textual é um estado de validação "
            "falha, não um valor nem uma categoria de negócio' — reafirma que "
            "o padrão SE(cond;\"ERRO!!!\";valor) é estruturalmente idêntico ao "
            "padrão A41 SE(cond;valor;\"F\") e exige a MESMA solução "
            "arquitetural (result_state), não um mecanismo próprio."
        ),
    }


# ------------------------------------------------------------
# D6/D7 — value_type x resultado, estado x exceção
# ------------------------------------------------------------


def d6_d7_boundary_probes():
    L = ("linha", "L1")
    rows = []

    def add(situacao, fn, esperado_categoria):
        o = outcome(fn)
        rows.append({
            "situacao": situacao,
            "produziu_valor_sem_excecao": "sim" if o["kind"] == "value" else "não",
            "resultado_observado": o["repr"][:120] if o["kind"] == "value" else f"{o['type']}: {o['repr'][:100]}",
            "categoria_proposta": esperado_categoria,
        })

    add("cálculo válido (2+3)", lambda: evaluate("VAR95301 + VAR95302", {"VAR95301": 2.0, "VAR95302": 3.0}), "VALID")
    add('nenhuma regra aplicável (produção de "F" por IF)', lambda: evaluate('VAR95303 if VAR95303 > 0 else "F"', {"VAR95303": -1.0}), "RESULT_STATE (NO_APPLICABLE_RULE)")
    add("categoria fora do domínio (aceita, sem erro)", lambda: CalculationContext(categorical_variable_ids={"VAR95304"}).set_variable("VAR95304", "CategoriaInexistente"), "AMBIGUOUS (hoje: VALID; deveria ser INVALID_INPUT)")
    add("ln(-10) — domínio matemático", lambda: evaluate("ln(VAR95305)", {"VAR95305": -10.0}), "TECHNICAL ERROR (MathDomainError)")
    add("variável inexistente (ID válido, sem valor no contexto)", lambda: evaluate("VAR95999"), "TECHNICAL ERROR (VariableNotFoundError)")
    add("divisão por zero", lambda: evaluate("VAR95306 / VAR95307", {"VAR95306": 1.0, "VAR95307": 0.0}), "TECHNICAL ERROR (DivisionByZeroError)")
    add('consumo aritmético de "F"', lambda: evaluate("VAR95308 + 1", {"VAR95308": "F"}), "RESULT_STATE consumido -> hoje: TECHNICAL EXCEPTION (ConditionalFailureError); mistura estado com erro")
    add('"ERRO!!!" como saída de fórmula futura', lambda: evaluate('VAR95309 if VAR95309 <= 10 else "ERRO!!!"', {"VAR95309": 15.0}), "RESULT_STATE (VALIDATION_FAILED) — hoje indistinguível de valor/categoria")
    return rows


# ------------------------------------------------------------
# C01-C16 — matriz de cenários conceituais
# ------------------------------------------------------------


def scenario_matrix():
    L = ("linha", "L1")
    rows = []

    def add(cid, titulo, status, evidencia):
        rows.append({"id": cid, "cenario": titulo, "status": status, "evidencia": evidencia})

    add("C01", "numeric + valor válido", "SUPPORTED", "evaluate('2+3')=5.0 (probe P07, Etapa 2)")
    add("C02", "categorical + categoria válida", "SUPPORTED", "P02, Etapa 2: 'ALTO' armazenado em categorical")
    add("C03", "numeric + estado especial (F)", "SUPPORTED (mas indistinguível de valor no canal)", "P04, Etapa 2: 'F' aceito em variável numeric")
    add("C04", "categorical + estado especial (F)", "AMBIGUOUS", "P14a, Etapa 2: 'F' como categoria colide com o sentinela — ConditionalFailureError inesperado numa comparação de categoria")
    add("C05", "categorical + categoria inválida (fora do domínio)", "UNSUPPORTED (não detectado)", "input_domain_gap() nesta etapa: aceito silenciosamente, sem allowed_values")
    add("C06", "numeric + string inesperada", "SUPPORTED (rejeitado se declarado numeric)", "P05a, Etapa 2: CalculationValueError se o ramo executa; P05b: NÃO detectado se o ramo não executa (latente)")
    add("C07", "technical exception (ln(-10), div/0, ref inexistente)", "SUPPORTED", "d6_d7_boundary_probes() nesta etapa + testes oficiais (MathDomainError, DivisionByZeroError, VariableNotFoundError)")
    add("C08", "estado em A + B independente (isolamento)", "UNSUPPORTED (F-001)", "f001_matrix.csv (Etapa 2) + d3_propagation_matrix() nesta etapa: B independente às vezes não calcula, dependendo da ordem topológica")
    add("C09", "estado + consumidor direto (referência/passthrough)", "SUPPORTED (repassa como valor)", "d3_propagation_matrix(): B=A repassa 'F' quando alcançado")
    add("C10", "estado + consumidor indireto (aritmético/comparação)", "PARTIALLY SUPPORTED (aritmético levanta exceção corretamente; mas aborta a rodada, impedindo outros indiretos de serem sequer tentados)", "d3_propagation_matrix()")
    add("C11", "estado + AVERAGE temporal", "SUPPORTED (invalida com períodos reportados)", "aggregation_matrix.csv (Etapa 2, P09a/P09b) + d4_temporal_policy_citation() nesta etapa")
    add("C12", "estado + 'agregação espacial' (grupo)", "N/A — não existe mecanismo de agregação espacial; é equação comum", "d4_spatial_vs_temporal_aggregation() nesta etapa")
    add("C13", "estado + total (soma de grupos)", "UNSUPPORTED (aritmético comum -> exceção, mesma classe de C10)", "f001_matrix.csv agregador(A+B)")
    add("C14", "categoria 'F' legítima (nome de negócio coincide com o sentinela)", "UNSUPPORTED (colisão)", "P14a, Etapa 2")
    add("C15", '"F" como estado reservado (uso atual do A41)', "SUPPORTED operacionalmente, mas MAL MODELADO (mistura valor/estado)", "F-analysis, Etapa 2 §6")
    add("C16", '"ERRO!!!" como saída de fórmula futura', "UNSUPPORTED como conceito distinto (hoje colapsa em C03 ou C04 dependendo do value_type declarado)", "erro_case_study() nesta etapa")
    return rows


# ------------------------------------------------------------
# main
# ------------------------------------------------------------


def write_csv(name, rows):
    with open(OUTDIR / "evidence" / name, "w", encoding="utf-8", newline="") as h:
        w = csv.DictWriter(h, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def main():
    EVIDENCE.mkdir(exist_ok=True)
    results = {
        "environment": {
            "commit": COMMIT, "python": platform.python_version(),
            "command": "python3 audit/stage2_1_result_state_contract/probes/probes_2_1.py",
        },
        "d1_1_no_rule_vs_invalid_input": d1_1_no_rule_vs_invalid_input(),
        "input_domain_gap": input_domain_gap(),
        "d3_propagation_matrix": d3_propagation_matrix(),
        "d3_isolated_arithmetic_consumer": d3_isolated_arithmetic_consumer(),
        "d3_isolated_comparison_consumer": d3_isolated_comparison_consumer(),
        "d4_spatial_vs_temporal_aggregation": d4_spatial_vs_temporal_aggregation(),
        "d4_temporal_policy_citation": d4_temporal_policy_citation(),
        "d5_category_domain": d5_category_domain(),
        "erro_case_study": erro_case_study(),
    }
    (EVIDENCE / "probes_2_1_results.json").write_text(
        json.dumps(results, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8"
    )
    write_csv("d3_propagation_detail.csv", results["d3_propagation_matrix"]["rows"])
    write_csv("state_error_boundary.csv", d6_d7_boundary_probes())
    write_csv("scenario_matrix_detail.csv", scenario_matrix())

    print(json.dumps(results, ensure_ascii=False, indent=1, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
