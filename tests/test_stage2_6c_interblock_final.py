"""
Etapa 2.6C — fechamento final do contrato interbloco.

    D26B-01  production v11: correção exclusiva da coluna OBS
    D26B-02  nome canônico `max_ht` (taxonomia, código, seeds, fonte)
    D26-01..03 revalidadas; livro de IDs auditado (estabilidade,
    preservação, aposentadoria, independência de ordem, determinismo)

Os testes 01..17 seguem a lista da etapa.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from tests.test_stage2_6_interblock_contract import entity, model, only_link, producer
from tools.workbook_seed.blocks import (
    BLOCKS,
    ID_LEDGER_ROOT,
    REPO_ROOT,
    SEED_ROOT,
    WORKBOOK_DIR,
    build_all,
    read_interblock_seed,
    read_seed_file,
)
from tools.workbook_seed.interblock import (
    SOURCE_BLOCK_NOT_LOADED,
    SOURCE_BLOCK_UNKNOWN,
    SOURCE_BLOCK_UNKNOWN_CODE,
    VALID,
    validate_interblock,
)
from tools.workbook_seed.taxonomy import BLOCK_TAXONOMY, OFFICIAL_BLOCKS


V10 = "descritivo_das_variáveis_production_v10.xlsx"
V11 = "descritivo_das_variáveis_production_v11.xlsx"
V11_SHA256 = "d946dfd522c8cbb301c0c63caa3c62263b74a849e043d7b2f4bcb0bdb262c922"
ANALYSIS = REPO_ROOT / "audit" / "stage2_6c_interblock_final" / "evidence" / "analysis_stage2_6c.py"


@pytest.fixture(scope="module")
def official():
    return build_all()


def _link(official, block, name):
    return next(
        l for l in official.interblock.links
        if (l.consumer_block, l.consumer.name) == (block, name)
    )


def _entity(official, block, name, frequency="diário"):
    return next(
        e for e in official.blocks[block].model.entities
        if e.name == name and e.frequency == frequency
    )


# ------------------------------------------------------------
# 01 — production v10 -> v11 (D26B-01)
# ------------------------------------------------------------

def test_01_production_v11_differs_from_v10_only_in_the_obs_correction():
    spec = BLOCKS["production"]
    assert (spec.version, spec.file_name, spec.sha256) == ("v11", V11, V11_SHA256)
    assert hashlib.sha256((WORKBOOK_DIR / V11).read_bytes()).hexdigest() == V11_SHA256

    v10 = openpyxl.load_workbook(WORKBOOK_DIR / V10)["production"]
    v11 = openpyxl.load_workbook(WORKBOOK_DIR / V11)["production"]
    assert (v10.max_row, v10.max_column) == (v11.max_row, v11.max_column)
    header = [c.value for c in v10[2]]

    changed = {
        (v11.cell(r, c).coordinate, header[c - 1], v10.cell(r, c).value, v11.cell(r, c).value)
        for r in range(1, v10.max_row + 1)
        for c in range(1, v10.max_column + 1)
        if v10.cell(r, c).value != v11.cell(r, c).value
    }
    assert changed == {
        ("R36", "OBS", 1050, None),
        ("R38", "OBS", 1100, 1050),
        ("R43", "OBS", None, 1100),
    }

    # OBS de lth_meta alinhado linha a linha com os antigos `value` (v9).
    obs = {
        v11.cell(r, 10).value: v11.cell(r, 18).value
        for r in range(3, v11.max_row + 1)
        if v11.cell(r, 2).value == "lth_meta"
    }
    assert obs == {"L1": 1050, "L2": 1050, "L3": 1100, "L4": 1100, "L5": 1100, "L6": 1100, "L7": 1100}
    assert v11.cell(36, 2).value == "lth" and v11.cell(36, 18).value is None


def test_01b_obs_is_not_part_of_the_seed():
    for block in BLOCKS:
        for name in ("variables", "parameters", "equations", "aggregation_rules", "manifest"):
            assert '"OBS"' not in json.dumps(read_seed_file(block, name), ensure_ascii=False)


# ------------------------------------------------------------
# 02..05 — max_ht (D26B-02)
# ------------------------------------------------------------

def test_02_taxonomy_contains_max_ht():
    assert "max_ht" in OFFICIAL_BLOCKS
    assert len(BLOCK_TAXONOMY) == len(set(BLOCK_TAXONOMY)) == 28
    assert set(BLOCKS) <= OFFICIAL_BLOCKS
    assert read_seed_file("max_ht", "manifest")["block"] == "max_ht"


def test_03_taxonomy_does_not_contain_mx_ht():
    assert "mx_ht" not in OFFICIAL_BLOCKS
    assert "mx_ht" not in BLOCKS
    assert "mx_ht" not in json.dumps(read_interblock_seed(), ensure_ascii=False)
    for path in list((REPO_ROOT / "data").rglob("*.json")):
        assert "mx_ht" not in path.read_text(encoding="utf-8"), path


def test_04_fonte_max_ht_resolves_when_the_producer_exists(official):
    models = {b: r.model for b, r in official.blocks.items()}
    result = validate_interblock({
        **models,
        "energy": model("energy", entity("VAR18999", "lth", fonte="max_ht")),
    })
    link = next(l for l in result.links if l.consumer.entity_id == "VAR18999")
    assert link.validation_status == VALID
    assert link.producer is _entity(official, "max_ht", "lth")


def test_05_fonte_mx_ht_is_rejected_as_unknown_block(official):
    models = {b: r.model for b, r in official.blocks.items()}
    result = validate_interblock({
        **models,
        "energy": model("energy", entity("VAR18999", "lth", fonte="mx_ht")),
    })
    link = next(l for l in result.links if l.consumer.entity_id == "VAR18999")
    assert link.validation_status == SOURCE_BLOCK_UNKNOWN
    assert link.findings[0].code == SOURCE_BLOCK_UNKNOWN_CODE
    assert link.is_rejected and link.producer is None


# ------------------------------------------------------------
# 06..11 — D26-02, D26-03
# ------------------------------------------------------------

def test_06_lth_meta_variable_to_variable(official):
    link = _link(official, "energy", "lth_meta")
    assert link.validation_status == VALID
    assert (link.consumer.kind, link.producer.kind) == ("variable", "variable")
    assert link.producer.entity_id == "VAR12066"
    for dimension in ("unit", "value_type", "allowed_values", "declared_result_states",
                      "frequency", "scope_type", "scope_value"):
        assert getattr(link.consumer, dimension) == getattr(link.producer, dimension)


def test_07_entrada_to_entrada_externa_is_not_a_mismatch(official):
    link = _link(official, "energy", "lth_meta")
    assert (link.consumer.variable_type, link.producer.variable_type) == ("entrada", "entrada_externa")
    assert link.findings == []


@pytest.mark.parametrize("name", ["fator_mpsa", "fator_mrn"])
def test_08_09_calculated_factor_without_fonte(official, name):
    factor = _entity(official, "production", name)
    assert factor.fonte is None and factor.variable_type == "calculado"
    assert [r.expression for r in factor.rows] == [f"{name}_kg_t / 1000"]
    assert not [l for l in official.interblock.links if l.consumer is factor]


@pytest.mark.parametrize("name", ["fator_mpsa_kg_t", "fator_mrn_kg_t"])
def test_10_11_kg_t_input_links_to_forecast(official, name):
    link = _link(official, "production", name)
    assert (link.source_block, link.consumer.variable_type) == ("forecast", "entrada")
    assert all(r.expression is None for r in link.consumer.rows)
    assert link.validation_status == SOURCE_BLOCK_NOT_LOADED and link.producer is None


# ------------------------------------------------------------
# 12..14 — livro de IDs
# ------------------------------------------------------------

def _manifest_ids(text):
    return {
        (e["kind"], e["name"], e["frequency"], e["scope_type"], e["scope_value"]): e["entity_id"]
        for e in json.loads(text)["entities"]
    }


@pytest.mark.parametrize("baseline", ["a126e02", "eceffd4"])
def test_12_ids_are_stable_against_baselines(baseline):
    for block in BLOCKS:
        before = _manifest_ids(subprocess.check_output(
            ["git", "show", f"{baseline}:data/seed/{block}/manifest.json"], cwd=REPO_ROOT,
        ).decode("utf-8"))
        after = _manifest_ids((SEED_ROOT / block / "manifest.json").read_text(encoding="utf-8"))
        for key in before.keys() & after.keys():
            assert before[key] == after[key], (block, key)
        if baseline == "eceffd4":
            assert before == after, block


def test_13_retired_ids_are_recorded_and_never_reused():
    ledger = json.loads((ID_LEDGER_ROOT / "production.json").read_text(encoding="utf-8"))
    assert [(r["entity_id"], r["kind"], r["name"]) for r in ledger["retired"]] == [
        ("PARAM12003", "parameter", "lth_meta"),
    ]
    for block in BLOCKS:
        ledger = json.loads((ID_LEDGER_ROOT / f"{block}.json").read_text(encoding="utf-8"))
        active = [e["entity_id"] for e in ledger["entries"]]
        retired = {r["entity_id"] for r in ledger["retired"]}
        assert len(active) == len(set(active))
        assert not retired & set(active)
        for name in ("variables", "parameters", "equations", "aggregation_rules"):
            assert not any(r in json.dumps(read_seed_file(block, name)) for r in retired)


_BUILD_SNIPPET = """
import hashlib, json, sys
from pathlib import Path
sys.path.insert(0, {repo!r})
from tools.workbook_seed.blocks import BLOCKS, build_all, write_id_ledger, write_interblock_seed, write_seeds
order = list(BLOCKS)
if {reverse!r}:
    order.reverse()
out = Path({out!r})
result = build_all(order=order)
for block, built in result.blocks.items():
    write_seeds(built, out / "seed" / block)
    write_id_ledger(built, out / "id_ledger")
write_interblock_seed(result, out / "seed")
digests = {{
    str(p.relative_to(out)): hashlib.sha256(p.read_bytes()).hexdigest()
    for p in sorted(out.rglob("*.json"))
}}
print(json.dumps(digests, sort_keys=True))
"""


def _isolated_build(tmp_path, name, hash_seed, reverse):
    out = tmp_path / name
    env = {**os.environ, "PYTHONHASHSEED": str(hash_seed)}
    code = _BUILD_SNIPPET.format(repo=str(REPO_ROOT), out=str(out), reverse=reverse)
    return json.loads(subprocess.check_output([sys.executable, "-c", code], env=env, cwd=REPO_ROOT))


def test_14_repeated_builds_are_byte_identical_and_order_independent(tmp_path):
    first = _isolated_build(tmp_path, "a", 0, False)
    second = _isolated_build(tmp_path, "b", 0, False)
    other_hash = _isolated_build(tmp_path, "c", 987654321, False)
    reversed_order = _isolated_build(tmp_path, "d", 12345, True)

    assert first == second == other_hash == reversed_order
    committed = {
        name: hashlib.sha256(
            ((SEED_ROOT if name.startswith("seed/") else ID_LEDGER_ROOT)
             / name.split("/", 1)[1]).read_bytes()
        ).hexdigest()
        for name in first
    }
    assert first == committed
    assert len(first) == 5 * 5 + 1 + 5


# ------------------------------------------------------------
# 15..17 — ciclos, órfãos, análise independente
# ------------------------------------------------------------

def test_15_no_cycle(official):
    assert official.interblock.cycles == []
    for link in official.interblock.valid:
        assert link.chain_status == VALID


def test_16_no_orphan_link_or_reference():
    ids = {
        block: {e["entity_id"] for e in read_seed_file(block, "manifest")["entities"]}
        for block in BLOCKS
    }
    seed = read_interblock_seed()
    for record in seed["links"]:
        assert record["consumer_definition"] in ids[record["consumer_block"]]
        assert record["source_definition"] in ids[record["source_block"]]
    for record in seed["pending"] + seed["rejected"]:
        assert record["consumer_definition"] in ids[record["consumer_block"]]
    for record in seed["pending"]:
        assert record["source_block"] not in BLOCKS and record["source_definition"] is None
    for block in BLOCKS:
        for eq in read_seed_file(block, "equations"):
            for ref in re.findall(r"\b(VAR\d+|PARAM\d+)", eq["expression"]):
                assert ref in ids[block], (block, eq["equation_id"], ref)
        for rule in read_seed_file(block, "aggregation_rules"):
            assert {rule["source_variable_id"], rule["target_variable_id"]} <= ids[block]
    lth_meta_equations = [
        eq["equation_id"] for eq in read_seed_file("production", "equations")
        if "VAR12066" in eq["expression"]
    ]
    assert lth_meta_equations == ["EQ12012", "EQ12014", "EQ12015"]


def test_17_independent_analysis_matches_the_builder():
    completed = subprocess.run(
        [sys.executable, str(ANALYSIS), "--no-write"],
        cwd=REPO_ROOT, capture_output=True, text=True,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
    summary = json.loads(completed.stdout)
    assert summary["independent_vs_builder_divergences"] == []
    assert summary["failures"] == []
    assert summary["links_by_class_definitions"] == {"SOURCE_BLOCK_NOT_LOADED": 16, "VALID": 13}
    assert summary["cycles_overapproximated"] == []


def test_official_classification_counts(official):
    result = official.interblock
    assert len(result.links) == 29
    assert len(result.valid) == 13
    assert len(result.pending) == 16
    assert result.rejected == []
