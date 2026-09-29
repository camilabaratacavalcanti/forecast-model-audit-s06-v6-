"""
Etapa 2.4 (rodada 2) — contrato dos cinco workbooks aprovados, do
workbook ao runtime (T24-01 .. T24-18).

    workbook aprovado (data/workbooks, SHA-256 conferido)
        -> reader (tools.workbook_seed.reader)
        -> modelo canônico (definições + instâncias por escopo)
        -> seed (data/seed/<bloco>)
        -> validators -> SeedLoader -> domínio
        -> ReferenceResolver / ScopeResolver -> ForecastEngine

Nenhum teste aqui depende de snapshot histórico: o workbook aprovado
é lido de verdade (openpyxl) e comparado com o seed versionado.
"""

import json
import re
from pathlib import Path

import openpyxl
import pytest

from app.domain.equations.models import EquationInstance
from app.domain.equations.registry import EquationDefinitionRegistry
from app.domain.variables.models import (
    DeclaredResultState,
    Variable,
    VariableDefinition,
)
from app.engine import reference_resolver
from app.engine.calculation_context import CalculationContext
from app.engine.equation_engine import EquationEngine
from app.engine.exceptions import (
    CalculationValueError,
    ConditionalFailureError,
    EquationEvaluationError,
    ParameterNotFoundError,
    VariableNotFoundError,
)
from app.engine.forecast_engine import ForecastEngine
from app.engine.temporal_aggregation_service import TemporalAggregationService
from app.repositories.seed_loader import SeedLoader
from app.validation import parameter_seed_validator, variable_seed_validator
from app.domain.results import Result, StatefulResultOnScalarApiError
from tools.workbook_seed.blocks import (
    BLOCKS,
    SEED_FILES,
    UnapprovedWorkbookError,
    build_block,
    read_seed_file,
    read_approved_workbook,
)
from tools.workbook_seed.canonical import (
    CanonicalModelError,
    IdentityCollisionError,
    build_canonical_model,
)
from tools.workbook_seed.reader import (
    REQUIRED_COLUMNS,
    WorkbookFormatError,
    read_workbook,
    sha256_of,
)
from tools.workbook_seed.seeds import build_seeds

from seed_ids import aggregation_rule, equation, parameter_id, variable_id

REPO = Path(__file__).resolve().parent.parent
SEED_ROOT = REPO / "data" / "seed"
LINES = ["L1", "L2", "L3", "L4", "L5", "L6", "L7"]
BLOCK_NAMES = sorted(BLOCKS)

EXPECTED_ROWS = {"area_41": 54, "energy": 56, "max_ht": 120, "production": 88, "yield": 254}
EXPECTED_HEADER_ROW = {"area_41": 2, "energy": 2, "max_ht": 2, "production": 2, "yield": 1}
EXPECTED_RULES = {"area_41": 10, "energy": 11, "max_ht": 78, "production": 28, "yield": 80}


@pytest.fixture(scope="module")
def builds():
    return {block: build_block(block) for block in BLOCK_NAMES}


@pytest.fixture(scope="module")
def loaded():
    loader = SeedLoader(SEED_ROOT)
    return loader.load_all_definitions_and_instances()


# ------------------------------------------------------------
# Leitura independente do workbook (não reutiliza o builder)
# ------------------------------------------------------------

def _raw_rows(block):
    spec = BLOCKS[block]
    ws = openpyxl.load_workbook(spec.workbook_path, data_only=True)[spec.sheet]
    header_row = next(
        r for r in range(1, 11)
        if "name" in [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    )
    header = {
        c: ws.cell(header_row, c).value
        for c in range(1, ws.max_column + 1)
        if ws.cell(header_row, c).value
    }
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        values = {name: ws.cell(r, c).value for c, name in header.items()}
        if any(v not in (None, "") for v in values.values()):
            values["_row"] = r
            rows.append(values)
    return rows


def _decode(expression, names):
    return re.sub(r"\b(VAR\d+|PARAM\d+)", lambda m: names[m.group(1)], expression)


def _squash(text):
    return re.sub(r"\s+", "", text)


# ============================================================
# T24-01 — builder(workbook aprovado) == seed versionado
# ============================================================


@pytest.mark.parametrize("block", BLOCK_NAMES)
def test_t24_01_builder_output_equals_versioned_seed(builds, block):
    seeds = builds[block].seeds

    for name in SEED_FILES:
        assert seeds[name] == read_seed_file(block, name), (block, name)


@pytest.mark.parametrize("block", BLOCK_NAMES)
def test_t24_01_every_workbook_row_is_represented_in_the_seed(block):
    """
    Reconciliação independente, linha a linha: cada linha do workbook
    aprovado corresponde a exatamente um registro do seed, com os
    mesmos campos; as expressões do seed, decodificadas de IDs para
    nomes, reproduzem o texto do workbook.
    """

    spec = BLOCKS[block]
    variables = read_seed_file(block, "variables")
    parameters = read_seed_file(block, "parameters")
    equations = read_seed_file(block, "equations")
    rules = read_seed_file(block, "aggregation_rules")

    names = {v["variable_id"]: v["variable_name"] for v in variables}
    names.update({p["parameter_id"]: p["parameter_name"] for p in parameters})

    used_parameters, used_equations, used_rules = set(), set(), set()

    for row in _raw_rows(block):
        where = f"{block} linha {row['_row']} ({row['name']})"
        source_reference = row["source_reference"] or spec.file_name

        if row["Type"] == "parameter":
            matches = [
                i for i, p in enumerate(parameters)
                if p["parameter_name"] == row["name"]
                and p["scope_type"] == row["scope_type"]
                and p["scope_value"] == row["scope_value"]
            ]
            assert len(matches) == 1, where
            used_parameters.add(matches[0])
            p = parameters[matches[0]]
            assert type(p["value"]) is type(row["value"]), where
            assert (p["value"], p["unit"], p["version"], p["frequency"],
                    p["status"], p["value_type"], p["description"],
                    p["source_reference"]) == (
                row["value"], row["unit"], row["version"], row["frequency"],
                row["status"], row["value_type"], row["description"],
                source_reference,
            ), where
            continue

        candidates = [
            v for v in variables
            if v["variable_name"] == row["name"]
            and v["frequency"] == row["frequency"]
            and v["scope_type"] == row["scope_type"]
            and (
                v["scope_value"] == row["scope_value"]
                or any(i["scope_value"] == row["scope_value"] for i in v.get("instances") or [])
            )
        ]
        assert len(candidates) == 1, where
        v = candidates[0]
        instance = next(
            (i for i in v.get("instances") or [] if i["scope_value"] == row["scope_value"]),
            None,
        )
        declared = instance or v
        assert (v["unit"], v["variable_type"], v["status"], v["value_type"]) == (
            row["unit"], row["variable_type"], row["status"], row["value_type"],
        ), where
        assert declared["description"] == row["description"] or (
            instance is None and v["description"] is None
        ), where
        assert declared["source_reference"] == source_reference or (
            instance is None and v["source_reference"] == spec.file_name
        ), where
        assert v.get("allowed_values") == (
            row["allowed_values"].split("\n") if row["allowed_values"] else None
        ), where
        assert (
            [f'{s["state"]} → "{s["literal"]}"' for s in v["declared_result_states"]]
            if v.get("declared_result_states") else None
        ) == (row["declared_result_states"].split("\n") if row["declared_result_states"] else None), where

        expression = row["expression"]

        if expression is None:
            continue

        row_rules = [
            i for i, r in enumerate(rules)
            if r["target_variable_id"] == v["variable_id"]
        ]
        row_equations = [
            i for i, e in enumerate(equations)
            if e["target_variable_id"] == v["variable_id"]
            and e["scope_value"] == row["scope_value"]
        ]

        if row_rules:
            assert len(row_rules) == 1 and not row_equations, where
            used_rules.add(row_rules[0])
            rule = rules[row_rules[0]]
            assert f"'{names[rule['source_variable_id']]}'" in expression or (
                names[rule["source_variable_id"]] in expression
            ), where
            if "weight_variable_id" in rule:
                assert names[rule["weight_variable_id"]] in expression, where
        else:
            assert len(row_equations) == 1, where
            used_equations.add(row_equations[0])
            decoded = _decode(equations[row_equations[0]]["expression"], names)
            assert _squash(decoded) == _squash(expression), where

    assert used_parameters == set(range(len(parameters)))
    assert used_equations == set(range(len(equations)))
    assert used_rules == set(range(len(rules)))


# ============================================================
# T24-02 — os cinco workbooks são consumíveis pelos builders
# ============================================================


@pytest.mark.parametrize("block", BLOCK_NAMES)
def test_t24_02_each_block_builder_consumes_its_approved_workbook(block):
    import importlib

    module = importlib.import_module(f"tools.{block}_seed_builder")
    result = module.build()

    assert result.workbook.file_name == BLOCKS[block].file_name
    assert result.workbook.header_row == EXPECTED_HEADER_ROW[block]
    assert len(result.workbook.rows) == EXPECTED_ROWS[block]
    assert sum(len(e.rows) for e in result.model.entities) == EXPECTED_ROWS[block]
    assert all(e.name for e in result.model.entities)


def test_t24_02_total_rows_and_equations_match_stage_2_3(builds):
    assert sum(len(b.workbook.rows) for b in builds.values()) == 572
    # 234 expressões executáveis da 2.3 + as 4 linhas antes recusadas
    # pela A019 (lth r35, oee r48, producao r71, n_ppt r107).
    assert sum(len(b.seeds["equations"]) for b in builds.values()) == 238


# ============================================================
# T24-03 — A019: lth r35 / oee r48 -> lth_meta por instância
# ============================================================


# LEGACY_TEST_EXPECTATION (Etapa 2.6B, D26-02): até production v9 os
# testes T24-03 fixavam `lth_meta` como Parameter anual por linha com
# value 1050/1050/1100x5 lido do workbook. Desde production v10,
# `lth_meta` é Variable `entrada_externa` anual linha/L1_L7 (sem value
# no workbook). A regra testada — UMA definição com sete instâncias,
# resolução A019 por instância, instância ausente é erro — não muda; os
# valores 1050/1100 passam a ser entradas do teste.
LTH_META_VALUES = {
    "L1": 1050, "L2": 1050, "L3": 1100, "L4": 1100,
    "L5": 1100, "L6": 1100, "L7": 1100,
}


def test_t24_03_lth_meta_is_one_definition_with_seven_instances():
    assert not [
        p for p in read_seed_file("production", "parameters")
        if p["parameter_name"] == "lth_meta"
    ]
    records = [
        v for v in read_seed_file("production", "variables")
        if v["variable_name"] == "lth_meta"
    ]

    assert [r["variable_id"] for r in records] == [variable_id("production", "lth_meta", "anual")]
    [record] = records
    assert (record["variable_type"], record["frequency"]) == ("entrada_externa", "anual")
    assert [i["scope_value"] for i in record["instances"]] == LINES
    assert "value" not in record


def test_t24_03_lth_and_oee_reference_the_single_lth_meta_definition():
    lth_meta = variable_id("production", "lth_meta", "anual")

    lth = equation("production", variable_id("production", "lth", "diário"))
    oee = equation("production", variable_id("production", "oee", "diário"))

    assert (lth["scope_type"], lth["scope_value"]) == ("linha", "L1_L7")
    assert (oee["scope_type"], oee["scope_value"]) == ("linha", "L1_L7")
    assert re.findall(r"PARAM\d+", oee["expression"]) == []
    assert re.findall(r"VAR\d+", oee["expression"]) == [
        variable_id("production", "lth", "diário"), lth_meta,
    ]
    assert lth_meta in lth["expression"]


def _run_oee(lth_values, lth_meta_values):
    loader = SeedLoader(SEED_ROOT)
    eq_defs = loader.load_equation_definitions()
    oee_eq = equation("production", variable_id("production", "oee", "diário"))
    subset = EquationDefinitionRegistry()
    subset.add(eq_defs.get(oee_eq["equation_id"]))

    lth = variable_id("production", "lth", "diário")
    lth_meta = variable_id("production", "lth_meta", "anual")
    context = CalculationContext()
    for line, value in lth_values.items():
        context.set_variable_value(lth, value, "linha", line)
    for line, value in lth_meta_values.items():
        context.set_variable_value(lth_meta, value, "linha", line)

    ForecastEngine().calculate_from_definition_registry(
        equation_definition_registry=subset, calculation_context=context,
    )
    oee = variable_id("production", "oee", "diário")
    return {line: context.get_variable_value(oee, "linha", line) for line in LINES}


def test_t24_03_runtime_each_instance_uses_its_own_line():
    lth_meta = dict(LTH_META_VALUES)
    lth = {line: 900.0 + 13 * i for i, line in enumerate(LINES, start=1)}

    oee = _run_oee(lth, lth_meta)

    for line in ("L1", "L3", "L7"):
        assert oee[line] == pytest.approx(lth[line] / lth_meta[line])
    # Associação cruzada detectável: L1 (1050) x L3 (1100).
    assert oee["L1"] != pytest.approx(lth["L1"] / lth_meta["L3"])
    assert oee["L3"] != pytest.approx(lth["L3"] / lth_meta["L1"])


def test_t24_03_missing_instance_is_an_error_not_a_fallback():
    lth_meta = dict(LTH_META_VALUES)
    del lth_meta["L3"]

    with pytest.raises((VariableNotFoundError, EquationEvaluationError)):
        _run_oee({line: 1000.0 for line in LINES}, lth_meta)


# ============================================================
# T24-04 — n_ppt r107 = tanque_base - tanque, por instância
# ============================================================


def test_t24_04_n_ppt_resolves_tanque_and_tanque_base_per_instance():
    n_ppt_eq = equation("yield", variable_id("yield", "n_ppt", "diário"))
    tanque = variable_id("yield", "tanque", "diário")
    tanque_base = parameter_id("yield", "tanque_base")

    assert n_ppt_eq["expression"] == f"{tanque_base} - {tanque}"
    # Uma definição de tanque (7 instâncias declaradas) e um parâmetro
    # tanque_base com 7 valores por linha — nada duplicado.
    tanque_def = next(v for v in read_seed_file("yield", "variables") if v["variable_id"] == tanque)
    assert [i["scope_value"] for i in tanque_def["instances"]] == LINES
    base = {p["scope_value"]: p["value"] for p in read_seed_file("yield", "parameters")
            if p["parameter_id"] == tanque_base}
    assert base == {"L1": 14, "L2": 14, "L3": 16, "L4": 18, "L5": 18, "L6": 18, "L7": 18}

    loader = SeedLoader(SEED_ROOT)
    definition = loader.load_equation_definitions().get(n_ppt_eq["equation_id"])
    engine = EquationEngine()
    context = CalculationContext()
    for i, line in enumerate(LINES, start=1):
        context.set_variable_value(tanque, 1.0 * i, "linha", line)
        context.set_parameter_value(tanque_base, base[line], "linha", line)

    for i, line in enumerate(LINES, start=1):
        instance = EquationInstance.create(definition=definition, scope_type="linha", scope_value=line)
        result = engine.calculate_instance(instance=instance, definition=definition, calculation_context=context)
        assert result == pytest.approx(base[line] - i), line

    # Sem tanque@L3: erro explícito, nunca o valor de outra linha.
    context_missing = CalculationContext()
    for i, line in enumerate(LINES, start=1):
        if line != "L3":
            context_missing.set_variable_value(tanque, 1.0 * i, "linha", line)
        context_missing.set_parameter_value(tanque_base, base[line], "linha", line)
    instance_l3 = EquationInstance.create(definition=definition, scope_type="linha", scope_value="L3")
    with pytest.raises((VariableNotFoundError, EquationEvaluationError)):
        engine.calculate_instance(instance=instance_l3, definition=definition, calculation_context=context_missing)


# ============================================================
# T24-05 — A019 20/20 e consumidores/agregações sobre definições
# por linha (retirada_condensado_linha, pick_up, producao)
# ============================================================

A019_CASES = [
    ("area_41", 22, "valor_retirada", None), ("area_41", 39, "hes", None),
    ("area_41", 42, "hes", None), ("area_41", 52, "retirada_condensado_linha", "diário"),
    ("area_41", 53, "retirada_condensado_linha", "diário"), ("production", 35, "lth_meta", None),
    ("production", 48, "lth_meta", None), ("production", 50, "lth_meta", None),
    ("production", 60, "pick_up", "diário"), ("production", 61, "pick_up", None),
    ("production", 71, "pick_up", None),
    *[("production", row, "pick_up_yield", None) for row in range(53, 60)],
    ("yield", 107, "tanque", None), ("yield", 107, "tanque_base", None),
]


def test_t24_05_a019_twenty_cases_resolve(builds):
    assert len(A019_CASES) == 20

    for block, row, name, lookup_frequency in A019_CASES:
        model = builds[block].model
        index = reference_resolver.build_name_index([e.as_index_entry() for e in model.entities])
        consumer = next(e for e in model.entities if any(r.row == row for r in e.rows))
        consumer_row = next(r for r in consumer.rows if r.row == row)
        explicit = re.findall(rf"(?<![\w@]){name}@(L\d(?:_L\d)?)", consumer_row.expression)

        for scope in explicit or [None]:
            resolved = reference_resolver.resolve_reference(
                name, index, lookup_frequency or consumer.frequency,
                consumer_scope=(consumer.scope_type, consumer_row.scope_value),
                explicit_scope_value=scope,
            )
            target = model.by_id()[resolved]
            assert target.name == name, (block, row, name)


def test_t24_05_formerly_blocked_cases_bind_one_definition(builds):
    retirada = variable_id("area_41", "retirada_condensado_linha", "diário")
    for frequency in ("mensal", "anual"):
        rule = aggregation_rule("area_41", variable_id("area_41", "retirada_condensado_linha", frequency))
        assert rule["source_variable_id"] == retirada
        assert rule["aggregation_type"] == "AVERAGE"

    pick_up = variable_id("production", "pick_up", "diário")
    assert aggregation_rule("production", variable_id("production", "pick_up", "mensal"))["source_variable_id"] == pick_up

    producao = equation("production", variable_id("production", "producao", "diário"))
    assert (producao["scope_type"], producao["scope_value"]) == ("linha", "L1_L7")
    assert pick_up in producao["expression"]

    # Nenhuma definição foi duplicada por linha: um ID por conceito.
    for block, name in (("area_41", "retirada_condensado_linha"), ("production", "pick_up"),
                        ("production", "producao"), ("yield", "tanque")):
        daily = [v for v in read_seed_file(block, "variables")
                 if v["variable_name"] == name and v["frequency"] == "diário"]
        assert len(daily) == 1, (block, name)
    # 7 equações de pick_up (uma por linha do workbook) produzem a MESMA definição.
    pick_up_equations = [e for e in read_seed_file("production", "equations") if e["target_variable_id"] == pick_up]
    assert sorted(e["scope_value"] for e in pick_up_equations) == LINES


def test_t24_05_aggregation_over_per_line_definition_keeps_each_line(loaded):
    retirada = variable_id("area_41", "retirada_condensado_linha", "diário")
    rule_data = aggregation_rule("area_41", variable_id("area_41", "retirada_condensado_linha", "mensal"))

    loader = SeedLoader(SEED_ROOT)
    rules = loader.load_aggregation_rules()
    instances = loader.load_aggregation_rule_instances(variable_definition_registry=loaded[0])
    rule = rules.get(rule_data["aggregation_rule_id"])
    assert sorted(
        i.scope_value for i in instances.all()
        if i.rule.aggregation_rule_id == rule.aggregation_rule_id
    ) == LINES

    from datetime import date

    context = CalculationContext()
    for day in range(1, 4):
        for i, line in enumerate(LINES, start=1):
            context.set_variable_value(retirada, 10.0 * i + day, "linha", line, f"2026-09-{day:02d}")

    service = TemporalAggregationService()
    for i, line in enumerate(LINES, start=1):
        result = service.aggregate(rule, context, "linha", line, date(2026, 9, 3))
        assert result.value == pytest.approx(10.0 * i + 2)


def test_t24_05_producao_instances_consume_pick_up_of_the_same_line(loaded):
    loader = SeedLoader(SEED_ROOT)
    eq_defs = loader.load_equation_definitions()
    producao_eq = equation("production", variable_id("production", "producao", "diário"))
    definition = eq_defs.get(producao_eq["equation_id"])

    lth = variable_id("production", "lth", "diário")
    pick_up = variable_id("production", "pick_up", "diário")
    fator = parameter_id("production", "fator_producao")
    context = CalculationContext()
    for i, line in enumerate(LINES, start=1):
        context.set_variable_value(lth, 1000.0, "linha", line)
        context.set_variable_value(pick_up, 10.0 * i, "linha", line)
        context.set_parameter_value(fator, 0.88, "linha", line)

    engine = EquationEngine()
    for i, line in enumerate(LINES, start=1):
        instance = EquationInstance.create(definition=definition, scope_type="linha", scope_value=line)
        result = engine.calculate_instance(instance=instance, definition=definition, calculation_context=context)
        assert result == pytest.approx(1000.0 * 10.0 * i * 0.88 * 1.0902 * 24 / 1000)


# ============================================================
# T24-06 — janela de linhas derivada do conteúdo
# ============================================================


def _synthetic(tmp_path, rows, header_row=2, extra=None, drop=(), name="wb.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    columns = [c for c in REQUIRED_COLUMNS if c not in drop]
    for c, column in enumerate(columns, start=1):
        ws.cell(header_row, c, column)
    for r, row in enumerate(rows, start=header_row + 1):
        for c, column in enumerate(columns, start=1):
            if column in row:
                ws.cell(r, c, row[column])
    for (r, c), value in (extra or {}).items():
        ws.cell(r, c, value)
    path = tmp_path / name
    wb.save(path)
    return path


def _row(**overrides):
    row = {
        "Type": "variable", "name": "x", "unit": "t", "variable_type": "entrada",
        "frequency": "diário", "scope_type": "linha", "scope_value": "L1_L7",
        "status": "ativo", "value_type": "numerico",
    }
    row.update(overrides)
    return row


def _build_synthetic(path, block="yield"):
    workbook = read_workbook(path, "S")
    model = build_canonical_model(block, workbook, BLOCKS[block].id_base)
    return workbook, model, build_seeds(model, BLOCKS[block].id_base)


def test_t24_06_max_ht_v9_ignores_trailing_empty_rows(builds):
    workbook = builds["max_ht"].workbook
    assert len(workbook.rows) == 120
    assert workbook.rows[-1].row == 122
    assert len(workbook.empty_rows) == workbook.last_sheet_row - 122


def test_t24_06_trailing_empty_rows_never_become_entities(tmp_path):
    path = _synthetic(tmp_path, [_row(), _row(name="y")], extra={(40, 1): None, (60, 2): ""})
    workbook, model, _ = _build_synthetic(path)
    assert [r.row for r in workbook.rows] == [3, 4]
    assert [e.name for e in model.entities] == ["x", "y"]


def test_t24_06_row_without_name_is_an_explicit_error(tmp_path):
    path = _synthetic(tmp_path, [_row(), _row(name=None)])
    with pytest.raises(WorkbookFormatError, match="não tem 'name'"):
        read_workbook(path, "S")


def test_t24_06_residual_content_outside_header_is_an_error(tmp_path):
    path = _synthetic(tmp_path, [_row()], extra={(10, 30): "sobra"})
    with pytest.raises(WorkbookFormatError, match="fora das colunas"):
        read_workbook(path, "S")


def test_t24_06_missing_required_column_is_an_error(tmp_path):
    path = _synthetic(tmp_path, [_row()], drop=("value_type",))
    with pytest.raises(WorkbookFormatError, match="value_type"):
        read_workbook(path, "S")


def test_t24_06_header_is_located_not_assumed(tmp_path):
    path = _synthetic(tmp_path, [_row()], header_row=1)
    assert read_workbook(path, "S").header_row == 1


# ============================================================
# T24-07 — value_type: workbook -> seed -> domínio, sem padrão
# ============================================================


def test_t24_07_value_type_reaches_the_domain_for_every_definition(loaded):
    variable_definitions, _vi, parameter_definitions, *_ = loaded
    seeds = {
        v["variable_id"]: v["value_type"]
        for block in BLOCK_NAMES for v in read_seed_file(block, "variables")
    }
    assert {d.variable_definition_id: d.value_type for d in variable_definitions.all()} == seeds
    assert {d.value_type for d in parameter_definitions.all()} == {"numerico"}


def test_t24_07_absent_value_type_is_an_error_everywhere(tmp_path):
    with pytest.raises(TypeError, match="value_type"):
        VariableDefinition("VAR1", "x", None, "t", "entrada", "diário", "linha", "L1", "s", "ativo")

    path = _synthetic(tmp_path, [_row(value_type=None)])
    with pytest.raises(CanonicalModelError, match="value_type"):
        _build_synthetic(path)

    seed = tmp_path / "seed" / "yield"
    seed.mkdir(parents=True)
    record = dict(read_seed_file("yield", "variables")[0])
    del record["value_type"]
    (seed / "variables.json").write_text(json.dumps([record]), encoding="utf-8")
    result = variable_seed_validator.validate_seed(tmp_path / "seed")
    assert any("Missing required field 'value_type'" in e for e in result["errors"])


@pytest.mark.parametrize("alias", ["numeric", "categorical", "Numerico", "numerico "])
def test_t24_07_aliases_are_rejected(tmp_path, alias):
    path = _synthetic(tmp_path, [_row(value_type=alias)])
    with pytest.raises(CanonicalModelError, match="value_type"):
        _build_synthetic(path)


# ============================================================
# T24-08 — A41 hes (categorico) ponta a ponta
# ============================================================


def _retirada_grupo_l4_l5_context(loaded, hes_l4, hes_l5):
    hes = variable_id("area_41", "hes", "diário", "linha", "L4_L7")
    context = CalculationContext()
    # As categóricas são declaradas a partir do value_type do domínio
    # (mesma regra que o ForecastEngine aplica), antes das entradas.
    context.declare_categorical_variables(
        d.variable_definition_id for d in loaded[0].all() if d.is_categorical
    )
    context.set_variable_value(hes, hes_l4, "linha", "L4")
    context.set_variable_value(hes, hes_l5, "linha", "L5")
    for name, value in (("retirada_cond_corr_ltp", 30.0), ("retirada_cond_corr_lth", 10.0),
                        ("desconto_retirada_41c", 5.0), ("desconto_retirada_41d", 3.0)):
        context.set_variable_value(
            variable_id("area_41", name, "diário", "linha_grupo", "L4_L5"), value, "linha_grupo", "L4_L5",
        )
    return context


def _run_retirada_grupo_l4_l5(loaded, context):
    variable_definitions = loaded[0]
    target = variable_id("area_41", "retirada_condensado_grupo", "diário", "linha_grupo", "L4_L5")
    eq_defs = SeedLoader(SEED_ROOT).load_equation_definitions()
    subset = EquationDefinitionRegistry()
    subset.add(eq_defs.get(equation("area_41", target)["equation_id"]))
    ForecastEngine().calculate_from_definition_registry(
        equation_definition_registry=subset, calculation_context=context,
        variable_definition_registry=variable_definitions,
    )
    return context.get_variable_value(target, "linha_grupo", "L4_L5")


def test_t24_08_hes_is_categorical_from_workbook_to_runtime(loaded, builds):
    entity = next(e for e in builds["area_41"].model.entities if e.name == "hes")
    assert entity.value_type == "categorico"
    assert [r.row for r in entity.rows] == [25, 26, 27, 28]

    definition = loaded[0].get(variable_id("area_41", "hes", "diário", "linha", "L4_L7"))
    assert definition.is_categorical
    assert [i.scope_value for i in definition.instances] == ["L4", "L5", "L6", "L7"]

    # O runtime aceita texto em hes porque a VariableDefinition declara
    # categorico (ForecastEngine declara as categóricas a partir dela).
    result = _run_retirada_grupo_l4_l5(loaded, _retirada_grupo_l4_l5_context(loaded, "Normal", "Normal"))
    assert result == pytest.approx(100 - (30.0 - 10.0) * 2)


def test_t24_08_text_in_a_numeric_variable_is_rejected():
    context = CalculationContext()
    with pytest.raises(CalculationValueError, match="não é declarada categórica"):
        context.set_variable_value(variable_id("area_41", "lth", "diário"), "Normal", "linha", "L1")


# ============================================================
# T24-09 — allowed_values
# ============================================================


def test_t24_09_allowed_values_reach_the_domain_and_are_validatable(loaded):
    definition = loaded[0].get(variable_id("area_41", "hes", "diário", "linha", "L4_L7"))

    assert definition.allowed_values == ("Normal", "LC", "Overhaul/Parada", "1 By pass", "1 By pass e LC")
    assert definition.is_value_in_domain("LC")
    assert not definition.is_value_in_domain("Parado")
    assert not definition.is_value_in_domain(1.0)
    assert not definition.is_value_in_domain("F")

    numeric = loaded[0].get(variable_id("area_41", "lth", "diário"))
    assert numeric.allowed_values is None
    assert numeric.is_value_in_domain(1.5) and not numeric.is_value_in_domain("1.5")


def test_t24_09_allowed_values_only_for_categorical(tmp_path):
    path = _synthetic(tmp_path, [_row(allowed_values="a\nb")])
    with pytest.raises(CanonicalModelError, match="allowed_values"):
        _build_synthetic(path)

    with pytest.raises(ValueError, match="allowed_values"):
        VariableDefinition("VAR1", "x", None, "t", "entrada", "diário", "linha", "L1", "s", "ativo",
                           "numerico", allowed_values=("a",))

    errors = variable_seed_validator.validate_allowed_values(
        [{"variable_id": "VAR1", "_block": "yield", "value_type": "categorico", "allowed_values": ["a", "a"]}]
    )
    assert errors


def test_t24_09_allowed_values_is_generic_not_an_a41_rule(tmp_path):
    path = _synthetic(tmp_path, [_row(value_type="categorico", allowed_values="Sim\nNão")])
    _wb, _model, seeds = _build_synthetic(path)
    assert seeds["variables"][0]["allowed_values"] == ["Sim", "Não"]


# ============================================================
# T24-10 — declared_result_states
# ============================================================


def test_t24_10_declared_result_states_reach_the_domain(loaded):
    for group in ("L4_L5", "L6_L7"):
        definition = loaded[0].get(
            variable_id("area_41", "retirada_condensado_grupo", "diário", "linha_grupo", group)
        )
        assert definition.declared_result_states == (DeclaredResultState("NO_APPLICABLE_RULE", "F"),)

    others = [
        d for d in loaded[0].all()
        if d.declared_result_states is not None
    ]
    assert len(others) == 2  # R1: só os produtores locais (A41 r39, r42)


def test_t24_10_no_applicable_rule_produces_the_text_literal_f(loaded):
    # LEGACY_TEST_EXPECTATION (Etapa 3.3B, contrato D1 da Etapa 2.2 §6.3):
    # até a 3.3A o runtime gravava o literal "F" no canal de valor. A
    # variável declara NO_APPLICABLE_RULE -> "F", então o EquationEngine
    # (com a definição alvo) produz Result(value=None,
    # state=NO_APPLICABLE_RULE); a fórmula continua devolvendo "F" (a
    # expressão avaliada sem tradução continua sendo texto, nunca número).
    context = _retirada_grupo_l4_l5_context(loaded, "1 By pass e LC", "Normal")
    with pytest.raises(StatefulResultOnScalarApiError):
        _run_retirada_grupo_l4_l5(loaded, context)
    target = variable_id("area_41", "retirada_condensado_grupo", "diário", "linha_grupo", "L4_L5")
    assert context.get_variable_result(target, "linha_grupo", "L4_L5") == Result(None, "NO_APPLICABLE_RULE")

    engine = EquationEngine()
    consumer = EquationInstance.create(
        definition=_consumer_of("VAR99001"), scope_type="linha", scope_value="L1",
    )
    context = CalculationContext()
    context.set_variable_value("VAR99001", "F", "linha", "L1")
    with pytest.raises((ConditionalFailureError, EquationEvaluationError)):
        engine.calculate_instance(instance=consumer, definition=_consumer_of("VAR99001"), calculation_context=context)


def _consumer_of(variable):
    from app.domain.equations.models import EquationDefinition

    return EquationDefinition("EQ99001", "VAR99002", 1, "linha", "L1", f"{variable} + 1", "t", "PUBLISHED")


@pytest.mark.parametrize(
    "cell",
    ['NO_APPLICABLE_RULE → F', 'UNKNOWN_STATE → "F"', 'NO_APPLICABLE_RULE -> "F"', 'NO_APPLICABLE_RULE → ""'],
)
def test_t24_10_malformed_states_are_rejected(tmp_path, cell):
    path = _synthetic(tmp_path, [_row(Type="variable / equation", expression='"F"', declared_result_states=cell)])
    with pytest.raises(CanonicalModelError, match="declared_result_states"):
        _build_synthetic(path)


def test_t24_10_declared_literal_must_appear_in_the_own_expression(tmp_path):
    path = _synthetic(tmp_path, [_row(Type="variable / equation", expression="1 + 2",
                                      declared_result_states='NO_APPLICABLE_RULE → "F"')])
    with pytest.raises(CanonicalModelError, match="R1"):
        _build_synthetic(path)


# ============================================================
# T24-11 — ltp_tc_base = 274 (yield v9 r102)
# ============================================================


def test_t24_11_ltp_tc_base_is_274(loaded):
    records = [p for p in read_seed_file("yield", "parameters") if p["parameter_name"] == "ltp_tc_base"]
    assert [(p["value"], p["scope_type"], p["scope_value"]) for p in records] == [(274, "linha", "L1_L7")]
    assert {d.value for d in loaded[2].all() if d.parameter_name == "ltp_tc_base"} == {274}


# ============================================================
# T24-12 — identidade duplicada é erro
# ============================================================


def test_t24_12_duplicate_identity_rows_are_rejected(tmp_path):
    path = _synthetic(tmp_path, [_row(), _row(unit="kg")])
    with pytest.raises(IdentityCollisionError):
        _build_synthetic(path)


def test_t24_12_instance_overlap_is_rejected(tmp_path):
    path = _synthetic(tmp_path, [_row(scope_value="L1_L7"), _row(scope_value="L3")])
    with pytest.raises(IdentityCollisionError):
        _build_synthetic(path)


def test_t24_12_per_line_rows_must_share_definition_attributes(tmp_path):
    path = _synthetic(tmp_path, [_row(scope_value="L1"), _row(scope_value="L2", unit="kg")])
    with pytest.raises(CanonicalModelError, match="divergem"):
        _build_synthetic(path)


def test_t24_12_seeds_have_no_intra_block_identity_collision():
    result = variable_seed_validator.validate_seed(SEED_ROOT)
    errors, _warnings = parameter_seed_validator.validate_seed(SEED_ROOT)
    assert result["errors"] == [] and errors == []


# ============================================================
# T24-13 — agregações derivadas da expressão (sem nº de linha)
# ============================================================


def test_t24_13_aggregation_counts_and_types(builds):
    counts = {block: len(builds[block].seeds["aggregation_rules"]) for block in BLOCK_NAMES}
    assert counts == EXPECTED_RULES
    assert sum(counts.values()) == 207

    types = {}
    for block in BLOCK_NAMES:
        for rule in builds[block].seeds["aggregation_rules"]:
            types[rule["aggregation_type"]] = types.get(rule["aggregation_type"], 0) + 1
    assert types == {"AVERAGE": 166, "SUM": 30, "WEIGHTED_AVERAGE": 9, "MOVING_AVERAGE": 2}


def test_t24_13_energy_rules_follow_the_expression_even_if_rows_move(tmp_path):
    """Reordenar as linhas do workbook não muda o significado das regras."""

    rows = [r for r in _raw_rows("energy")]
    reordered = list(reversed(rows))
    path = _synthetic(
        tmp_path,
        [{k: v for k, v in r.items() if k in REQUIRED_COLUMNS} for r in reordered],
    )
    _wb, _model, seeds = _build_synthetic(path, block="energy")

    def semantic(rules, variables):
        names = {v["variable_id"]: (v["variable_name"], v["frequency"]) for v in variables}
        return sorted(
            (names[r["target_variable_id"]], names[r["source_variable_id"]],
             names.get(r.get("weight_variable_id")), r["aggregation_type"])
            for r in rules
        )

    assert semantic(seeds["aggregation_rules"], seeds["variables"]) == semantic(
        read_seed_file("energy", "aggregation_rules"), read_seed_file("energy", "variables")
    )


def test_t24_13_unknown_aggregation_text_is_an_error(tmp_path):
    path = _synthetic(tmp_path, [
        _row(),
        _row(Type="variable / equation", name="x", frequency="mensal",
             expression="Média trimestral de todos os resultados diários de 'x'"),
    ])
    with pytest.raises(Exception, match="gramática"):
        _build_synthetic(path)


# ============================================================
# T24-14 — source_reference / proveniência
# ============================================================


@pytest.mark.parametrize("block", BLOCK_NAMES)
def test_t24_14_source_reference_is_the_workbook_actually_read(block):
    spec = BLOCKS[block]
    manifest = read_seed_file(block, "manifest")
    assert manifest["workbook"]["file_name"] == spec.file_name
    assert manifest["workbook"]["sha256"] == spec.sha256 == sha256_of(spec.workbook_path)

    cells = {r["source_reference"] for r in _raw_rows(block) if r["source_reference"]}
    allowed = cells | {spec.file_name}
    for name in ("variables", "parameters", "equations"):
        for record in read_seed_file(block, name):
            assert record["source_reference"] in allowed, (block, name, record)
            for instance in record.get("instances") or []:
                assert instance["source_reference"] in allowed

    text = "".join((SEED_ROOT / block / f"{n}.json").read_text(encoding="utf-8") for n in SEED_FILES)
    assert "energy_v2" not in text and "MaxHT_v5" not in text and "yield_v4" not in text


def test_t24_14_builder_refuses_an_unapproved_workbook(tmp_path):
    spec = BLOCKS["energy"]
    tampered = tmp_path / spec.file_name
    wb = openpyxl.load_workbook(spec.workbook_path)
    wb.active.cell(3, 5, 1)
    wb.save(tampered)
    with pytest.raises(UnapprovedWorkbookError):
        read_approved_workbook(spec, tampered)


# ============================================================
# T24-15 — cross-workbook: decisão contratual pendente (D24-11)
# ============================================================

CROSS_WORKBOOK_LINKS = [
    ("area_41", "lth", "diário", "linha", "yield"),
    ("energy", "producao", "diário", "linha", "production"),
    ("energy", "pick_up", "diário", "linha", "production"),
    ("energy", "pick_up_total", "diário", "linha_grupo", "production"),
    ("energy", "pick_up_total", "mensal", "linha_grupo", "production"),
    ("energy", "lth", "diário", "linha", "production"),
    ("energy", "lth_total", "diário", "linha_grupo", "production"),
    ("energy", "lth_total", "mensal", "linha_grupo", "production"),
    ("energy", "lth_meta", "anual", "linha", "production"),
    ("max_ht", "lth", "diário", "linha", "production"),
    ("max_ht", "producao", "diário", "linha", "production"),
    ("production", "yield", "diário", "linha", "yield"),
    ("yield", "lth", "diário", "linha", "production"),
]


def test_t24_15_cross_workbook_links_are_documented_not_decided():
    """
    DECISÃO PENDENTE (D24-11): o contrato não define se um consumidor
    referencia a definição do workbook produtor (identidade global) ou
    mantém a sua própria entrada (identidade local), nem como tratar
    produtor-parâmetro x consumidor-variável (lth_meta). O builder
    representa cada workbook como escrito: o consumidor declara a sua
    própria definição (sem valor calculado); nada liga os IDs.
    """

    assert len(CROSS_WORKBOOK_LINKS) == 13

    for consumer, name, frequency, scope_type, producer in CROSS_WORKBOOK_LINKS:
        local = variable_id(consumer, name, frequency, scope_type, "L1_L7")
        local_def = next(v for v in read_seed_file(consumer, "variables") if v["variable_id"] == local)
        assert local_def["variable_type"] in {"entrada", "entrada_externa"}, (consumer, name)

        produced_by = [
            v["variable_id"] for v in read_seed_file(producer, "variables")
            if v["variable_name"] == name and v["frequency"] == frequency and v["scope_type"] == scope_type
        ] + [
            p["parameter_id"] for p in read_seed_file(producer, "parameters") if p["parameter_name"] == name
        ]
        assert produced_by and local not in produced_by, (consumer, name)

    warnings = variable_seed_validator.validate_seed(SEED_ROOT)["warnings"]
    assert warnings and all("D24-11" in w for w in warnings)


# ============================================================
# T24-16 — value em Type=variable: decisão contratual pendente (D24-12)
# ============================================================


def test_t24_16_value_on_a_variable_row_is_recorded_not_dropped():
    """
    D24-12 (fechada na 2.5B: value/version exclusivos de parameter).
    LEGACY_TEST_EXPECTATION atualizada na Etapa 2.6: o yield v9 r92
    (`ltp_tc`, variable) trazia value=273/version=1 e o teste fixava a
    pendência no manifesto; desde o yield v10 a célula está vazia e
    nenhum workbook oficial preenche value/version numa variável. O
    mecanismo de registro continua o mesmo (não há rejeição nova nesta
    etapa); o valor nunca entra no domínio.
    """

    for block in BLOCKS:
        assert [
            d for d in read_seed_file(block, "manifest")["pending_contract_decisions"]
            if d["decision_id"] == "D24-12"
        ] == [], block
    assert "value" not in Variable.__dataclass_fields__
    ltp_tc = variable_id("yield", "ltp_tc", "diário")
    assert "value" not in next(v for v in read_seed_file("yield", "variables") if v["variable_id"] == ltp_tc)


# ============================================================
# T24-17 — isolamento de falha por nó (Etapa 3)
# ============================================================


def test_t24_17_first_failure_still_aborts_the_run_explicitly():
    """
    FORA DO ESCOPO DA 2.4 (D24-13 / F-001, Etapa 3): hoje a primeira
    falha interrompe a execução com erro explícito — nenhum valor
    parcial silencioso. Este teste fixa o comportamento atual até a
    arquitetura de resultados por nó da Etapa 3.
    """

    from app.domain.equations.models import EquationDefinition

    registry = EquationDefinitionRegistry()
    registry.add(EquationDefinition("EQ99011", "VAR99011", 1, "linha", "L1", "VAR99001 + 1", "t", "PUBLISHED"))
    registry.add(EquationDefinition("EQ99012", "VAR99012", 1, "linha", "L1", "VAR99002 * 2", "t", "PUBLISHED"))
    context = CalculationContext()
    context.set_variable_value("VAR99001", "F", "linha", "L1")
    context.set_variable_value("VAR99002", 21.0, "linha", "L1")

    with pytest.raises(EquationEvaluationError):
        ForecastEngine().calculate_from_definition_registry(
            equation_definition_registry=registry, calculation_context=context,
        )


# ============================================================
# T24-18 — sem coerção numérica
# ============================================================


@pytest.mark.parametrize("value", ["1.05", "47,5", True, None])
def test_t24_18_non_numeric_parameter_values_are_rejected(tmp_path, value):
    path = _synthetic(tmp_path, [_row(Type="parameter", value=value, version=1)])
    with pytest.raises(CanonicalModelError, match="value"):
        _build_synthetic(path)


def test_t24_18_physical_types_are_preserved():
    values = {
        (p["parameter_name"], p["scope_value"]): p["value"]
        for p in read_seed_file("production", "parameters")
    }
    # LEGACY_TEST_EXPECTATION (Etapa 2.6B, D26-02): o inteiro de
    # referência era lth_meta@L1 (1050), que deixou de ser parâmetro.
    assert type(values[("n_dias_ano", "L1_L7")]) is int
    assert type(values[("desaguamento_produtividade", "PLANTA")]) is int
    assert type(values[("fator_producao", "L1_L7")]) is float
    assert {type(p["value"]) for p in read_seed_file("max_ht", "parameters")} == {float}

    for module in ("workbook_seed/reader.py", "workbook_seed/canonical.py", "workbook_seed/seeds.py",
                   "energy_seed_builder.py", "max_ht_seed_builder.py"):
        source = (REPO / "tools" / module).read_text(encoding="utf-8")
        assert "float(" not in source and "int(entity" not in source, module


def test_t24_05_equation_may_produce_only_a_declared_instance():
    from app.domain.equations.models import EquationDefinition
    from app.domain.variables.models import VariableInstanceDeclaration
    from app.domain.variables.registry import VariableDefinitionRegistry
    from app.engine.exceptions import EquationTargetScopeMismatchError
    from app.engine.registry_validator import RegistryIntegrityValidator

    registry = VariableDefinitionRegistry()
    registry.add(VariableDefinition(
        "VAR99101", "x", None, "-", "entrada", "diário", "linha", "L4_L7", "s", "ativo", "numerico",
        instances=tuple(VariableInstanceDeclaration(line, None, "s") for line in ("L4", "L5", "L6", "L7")),
    ))
    registry.add(VariableDefinition(
        "VAR99102", "y", None, "-", "calculado", "diário", "linha", "L1_L7", "s", "ativo", "numerico",
    ))
    validator = RegistryIntegrityValidator()

    def check(target, scope_value):
        validator.validate_definition(
            EquationDefinition("EQ99101", target, 1, "linha", scope_value, "1", "s", "PUBLISHED"),
            registry,
        )

    check("VAR99101", "L5")  # instância declarada
    with pytest.raises(EquationTargetScopeMismatchError):
        check("VAR99101", "L2")  # linha que a definição não declara
    with pytest.raises(EquationTargetScopeMismatchError):
        check("VAR99102", "L3")  # definição sem instâncias declaradas
